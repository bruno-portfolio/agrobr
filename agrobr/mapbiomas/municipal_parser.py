from __future__ import annotations

import hashlib
import io
import itertools
import json
import re
import xml.etree.ElementTree as ET
import zipfile
import zlib
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from typing import Any

import openpyxl
import pandas as pd
import pydantic
from lxml import etree
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from agrobr import constants
from agrobr.exceptions import InvalidParameterError, ParseError
from agrobr.utils import io as io_utils

from . import models

PARSER_VERSION = 2


def _error(reason: str) -> ParseError:
    return ParseError(source="mapbiomas", parser_version=PARSER_VERSION, reason=reason)


def _identity_fields(collection: int) -> tuple[str, ...]:
    return (
        constants.MAPBIOMAS_MUNICIPAL_IDENTITY_FIELDS_10
        if collection == 10
        else constants.MAPBIOMAS_MUNICIPAL_IDENTITY_FIELDS
    )


@dataclass(frozen=True)
class MunicipalLayout:
    sheet: str
    headers: tuple[str, ...]
    positions: dict[str, int]
    year_positions: dict[int, int]
    fingerprint: str


def _layout(header: tuple[Any, ...], collection: int, sheet: str) -> MunicipalLayout:
    names: list[str] = []
    years: dict[int, int] = {}
    for index, value in enumerate(header):
        if isinstance(value, bool) or not isinstance(value, (str, int)):
            raise _error(f"{sheet}: cabeçalho inválido na coluna {index + 1}")
        name = str(value).strip()
        match = re.fullmatch(r"y?([0-9]{4})", name)
        if match:
            year = int(match.group(1))
            years[year] = index
            name = f"y{year}"
        if not name or name in names:
            raise _error(f"{sheet}: cabeçalho vazio ou duplicado: {name!r}")
        names.append(name)
    expected_years = set(range(models.ANO_INICIO, models.ANOS_FINAIS[collection] + 1))
    if set(years) != expected_years:
        raise _error(
            f"{sheet}: anos incompatíveis com a coleção {collection}; "
            f"ausentes={sorted(expected_years - years.keys())}, "
            f"extras={sorted(years.keys() - expected_years)}"
        )
    identity = set(_identity_fields(collection))
    actual_identity = set(names) - {f"y{year}" for year in years}
    if actual_identity != identity:
        raise _error(
            f"{sheet}: dimensão municipal/layout incompatível; "
            f"ausentes={sorted(identity - actual_identity)}, "
            f"extras={sorted(actual_identity - identity)}"
        )
    structure = {
        "colecao": collection,
        "sheet": sheet,
        "headers": names,
        "years": list(years),
        "parser_version": PARSER_VERSION,
    }
    fingerprint = hashlib.sha256(
        json.dumps(structure, ensure_ascii=False, sort_keys=True).encode("utf-8")
    ).hexdigest()
    return MunicipalLayout(
        sheet, tuple(names), dict(zip(names, range(len(names)))), years, fingerprint
    )


def _parse_row(
    values: tuple[Any, ...],
    layout: MunicipalLayout,
    collection: int,
    source_row: int,
) -> models.MunicipalCoverageRecord:
    if len(values) != len(layout.headers):
        raise _error(f"{layout.sheet}: linha {source_row} com largura incompatível")
    payload = {name: values[layout.positions[name]] for name in _identity_fields(collection)}
    payload["areas"] = tuple(values[index] for index in layout.year_positions.values())
    try:
        model = models.MunicipalCoverageRow10 if collection == 10 else models.MunicipalCoverageRow
        return model.model_validate(
            payload,
            context={"colecao": collection, "years": tuple(layout.year_positions)},
        )
    except pydantic.ValidationError as exc:
        issue = exc.errors(include_url=False, include_input=False)[0]
        location = ".".join(map(str, issue["loc"]))
        raise _error(
            f"{layout.sheet}: linha {source_row}, campo {location}: {issue['msg']}"
        ) from exc


@dataclass
class MunicipalPopulation:
    years: tuple[int, ...]
    physical_rows: int = 0
    empty_rows: int = 0
    validated_rows: int = 0
    selected_rows: int = 0
    keys: Counter[tuple[str, str, str, int]] = field(default_factory=Counter)
    source_ids: set[int] = field(default_factory=set)
    names: dict[tuple[str, str], str] = field(default_factory=dict)
    geocode_states: dict[str, set[str]] = field(default_factory=lambda: defaultdict(set))
    states: Counter[str] = field(default_factory=Counter)
    classes: Counter[int] = field(default_factory=Counter)
    minimum: list[float | None] = field(default_factory=list)
    maximum: list[float | None] = field(default_factory=list)
    zeros: list[int] = field(default_factory=list)

    def __post_init__(self) -> None:
        self.minimum = [None] * len(self.years)
        self.maximum = [None] * len(self.years)
        self.zeros = [0] * len(self.years)

    def observe(self, row: models.MunicipalCoverageRecord, source_row: int) -> None:
        uf = models.estado_para_uf(row.state)
        key = (row.biome, uf, row.geocode, row.class_id)
        if row.source_id in self.source_ids:
            raise _error(f"linha {source_row}: ID publicado duplicado {row.source_id}")
        self.source_ids.add(row.source_id)
        self.keys[key] += 1
        name_key = (uf, row.geocode)
        name = row.municipality.strip()
        previous = self.names.setdefault(name_key, name)
        if previous != name:
            raise _error(f"linha {source_row}: geocode/nome incompatíveis na mesma UF {name_key}")
        self.geocode_states[row.geocode].add(uf)
        self.states[uf] += 1
        self.classes[row.class_id] += 1
        self.validated_rows += 1
        for index, area in enumerate(row.areas):
            minimum, maximum = self.minimum[index], self.maximum[index]
            self.minimum[index] = area if minimum is None else min(minimum, area)
            self.maximum[index] = area if maximum is None else max(maximum, area)
            self.zeros[index] += area == 0

    def details(self, layout: MunicipalLayout, collection: int, output_rows: int) -> dict[str, Any]:
        multiple_states = [
            {"geocodigo": code, "estados": sorted(states)}
            for code, states in sorted(self.geocode_states.items())
            if len(states) > 1
        ]
        return {
            "sheet": layout.sheet,
            "collection": collection,
            "parser_version": PARSER_VERSION,
            "layout_fingerprint": layout.fingerprint,
            "layout_fingerprint_algorithm": "sha256",
            "coverage": {
                "physical_rows": self.physical_rows,
                "empty_rows": self.empty_rows,
                "validated_rows": self.validated_rows,
                "selected_rows": self.selected_rows,
                "annual_cells": self.validated_rows * len(self.years),
                "output_rows": output_rows,
                "years": list(self.years),
                "eof_reached": True,
                "scope": "published_workbook_population",
            },
            "annual_statistics": {
                str(year): {
                    "count": self.validated_rows,
                    "null_count": 0,
                    "zero_count": self.zeros[index],
                    "minimum_ha": self.minimum[index],
                    "maximum_ha": self.maximum[index],
                }
                for index, year in enumerate(self.years)
            },
            "states": dict(sorted(self.states.items())),
            "classes": {str(key): value for key, value in sorted(self.classes.items())},
            "geocodes_with_multiple_states": multiple_states,
            "territorial_keys": {
                "unique": len(self.keys),
                "repeated_groups": sum(count > 1 for count in self.keys.values()),
                "rows_in_repeated_groups": sum(count for count in self.keys.values() if count > 1),
                "additional_rows": self.validated_rows - len(self.keys),
            },
            "published_ids": {
                "unique": len(self.source_ids),
                "minimum": min(self.source_ids),
                "maximum": max(self.source_ids),
                "scope": "one_collection_and_resource",
            },
            "warnings": [],
        }


def _matches(row: models.MunicipalCoverageRecord, selectors: dict[str, Any]) -> bool:
    actual = {
        "bioma": row.biome,
        "uf": models.estado_para_uf(row.state),
        "geocodigo": row.geocode,
        "classe_id": row.class_id,
    }
    return all(value is None or actual[name] == value for name, value in selectors.items())


def _append(
    buffers: dict[int, dict[str, list[Any]]],
    row: models.MunicipalCoverageRecord,
    years: tuple[int, ...],
    collection: int,
) -> None:
    shared = {
        "bioma": row.biome,
        "uf": models.estado_para_uf(row.state),
        "municipio": row.municipality.strip(),
        "classe_id": row.class_id,
        "classe": models.classe_para_nome(row.class_id, collection),
        "nivel_0": row.class_level_0,
        "geocodigo": row.geocode,
        "id_registro": row.source_id,
    }
    for year, area in zip(years, row.areas, strict=True):
        if year not in buffers:
            continue
        columns = buffers[year]
        for name, value in shared.items():
            columns[name].append(value)
        columns["ano"].append(year)
        columns["area_ha"].append(area)


def _build_frame(buffers: dict[int, dict[str, list[Any]]]) -> pd.DataFrame:
    columns: dict[str, pd.Series[Any]] = {}
    for name in models.COLUNAS_SAIDA_COBERTURA_MUNICIPAL_V2:
        dtype = (
            "Int64"
            if name in {"ano", "classe_id", "id_registro"}
            else "float64"
            if name == "area_ha"
            else pd.Series([""]).dtype
        )
        values = list(itertools.chain.from_iterable(year[name] for year in buffers.values()))
        columns[name] = pd.Series(values, dtype=dtype)
    return pd.DataFrame(columns, copy=False)


def _consume(
    worksheet: ReadOnlyWorksheet,
    collection: int,
    selectors: dict[str, Any],
    year_filter: int | None,
) -> tuple[pd.DataFrame, dict[str, Any]]:
    worksheet.reset_dimensions()
    rows = worksheet.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        raise _error(f"{worksheet.title}: cabeçalho municipal ausente")
    layout = _layout(header, collection, worksheet.title)
    years = tuple(layout.year_positions)
    population = MunicipalPopulation(years)
    buffers: dict[int, dict[str, list[Any]]] = {
        year: {name: [] for name in models.COLUNAS_SAIDA_COBERTURA_MUNICIPAL_V2}
        for year in years
        if year_filter is None or year == year_filter
    }
    for source_row, values in enumerate(rows, start=2):
        population.physical_rows += 1
        if all(value is None for value in values):
            population.empty_rows += 1
            continue
        row = _parse_row(values, layout, collection, source_row)
        population.observe(row, source_row)
        if _matches(row, selectors):
            population.selected_rows += 1
            _append(buffers, row, years, collection)
    if population.validated_rows == 0:
        raise _error(f"{worksheet.title}: população municipal sem linhas identificadas")
    geocodigo = selectors.get("geocodigo")
    if geocodigo is not None and geocodigo not in population.geocode_states:
        raise InvalidParameterError(
            f"municipio={geocodigo!r}: geocódigo ausente do recurso municipal da coleção "
            f"{collection}"
        )
    frame = _build_frame(buffers)
    details = population.details(layout, collection, len(frame))
    details["filters"] = {**selectors, "ano": year_filter}
    return frame, details


def parse_cobertura_municipal(
    data: bytes,
    *,
    colecao: int,
    bioma: str | None = None,
    uf: str | None = None,
    geocodigo: str | None = None,
    ano: int | None = None,
    classe_id: int | None = None,
) -> tuple[pd.DataFrame, dict[str, Any]]:
    if (
        isinstance(colecao, bool)
        or not isinstance(colecao, int)
        or colecao not in models.ANOS_FINAIS
    ):
        raise _error("coleção municipal não suportada")
    selectors = {
        "bioma": bioma,
        "uf": uf,
        "geocodigo": geocodigo,
        "classe_id": classe_id,
    }
    io_utils.check_xlsx_expansion(data, source="mapbiomas")
    workbook = None
    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(data),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
        sheet = f"COVERAGE_{colecao}"
        if sheet not in workbook.sheetnames:
            raise _error(f"aba municipal {sheet} ausente")
        return _consume(workbook[sheet], colecao, selectors, ano)
    except InvalidParameterError:
        raise
    except (
        zipfile.BadZipFile,
        zlib.error,
        openpyxl.utils.exceptions.InvalidFileException,
        ET.ParseError,
        etree.XMLSyntaxError,
        OSError,
        EOFError,
        KeyError,
        IndexError,
        ValueError,
    ) as exc:
        raise _error(f"workbook municipal inválido ou incompleto ({type(exc).__name__})") from exc
    finally:
        if workbook is not None:
            workbook.close()
