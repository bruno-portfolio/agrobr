from __future__ import annotations

import argparse
import asyncio
import csv
import hashlib
import io
import json
import re
import sys
import unicodedata
from collections.abc import Iterable
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import httpx
import openpyxl
from openpyxl.utils import get_column_letter

from agrobr import constants
from agrobr.comexstat import _tls
from agrobr.comexstat import client as comexstat_client
from agrobr.comexstat import models as comexstat_models
from agrobr.comtrade import models as comtrade_models
from agrobr.http.user_agents import UserAgentRotator

ROOT = Path(__file__).resolve().parents[1]
MANIFEST = ROOT / "tests/golden_data/reconciliacao_comercio_exterior_20260918/manifest.json"
PRODUTOS_COMEX = ["soja", "milho", "cafe", "algodao", "acucar", "farelo_soja", "oleo_soja"]
NCM_DICTIONARY_URL = constants.COMEXSTAT_DICTIONARY_URLS["unidades"].replace(
    "NCM_UNIDADE.csv", "NCM.csv"
)
COMTRADE_COUNT_URL = (
    constants.URLS[constants.Fonte.COMTRADE]["guest"]
    + "/C/A/HS?reporterCode=76&flowCode=X&cmdCode=1201&period=2023&includeDesc=True"
    "&partner2Code=0&motCode=0&customsCode=C00&maxRecords=500&countOnly=true"
)
COMTRADE_HS_REFERENCE_URL = "https://comtradeapi.un.org/files/v1/app/reference/HS.json"
NCM_REGRAS: dict[str, dict[str, Any]] = {
    "soja": {
        "posicoes": ("1201",),
        "exige": ("SOJA",),
        "alternativas": (),
        "exclui": ("OLEO", "FARINHA", "PELLETS", "BAGACO", "TORTA", "FARELO"),
    },
    "milho": {
        "posicoes": ("1005",),
        "exige": ("MILHO",),
        "alternativas": (),
        "exclui": ("OLEO", "FARINHA", "AMIDO", "GLUTEN", "FARELO"),
    },
    "cafe": {
        "posicoes": ("0901",),
        "exige": ("CAFE",),
        "alternativas": (),
        "exclui": ("SOLUVEL", "EXTRATO", "SUCEDANEO", "CASCAS"),
    },
    "algodao": {
        "posicoes": ("5201", "5203"),
        "exige": ("ALGODAO",),
        "alternativas": (),
        "exclui": ("LINTERS", "DESPERDICIOS", "FIOS", "TECIDO"),
    },
    "acucar": {
        "posicoes": ("1701",),
        "exige": ("ACUCAR",),
        "alternativas": (),
        "exclui": ("MELACO", "XAROPE", "GLICOSE", "LACTOSE"),
    },
    "farelo_soja": {
        "posicoes": ("2304",),
        "exige": ("SOJA",),
        "alternativas": ("FARINHA", "PELLETS", "FARELO", "TORTA", "BAGACO", "RESIDUOS"),
        "exclui": (),
    },
    "oleo_soja": {
        "posicoes": ("1507",),
        "exige": ("OLEO", "SOJA"),
        "alternativas": (),
        "exclui": ("FARINHA", "PELLETS", "FARELO", "TORTA", "SEMEADURA"),
    },
}
HS_KEYWORDS: dict[str, tuple[str, ...]] = {
    "soja": ("SOYA BEANS",),
    "complexo_soja": ("SOYA",),
    "farelo_soja": ("OIL-CAKE",),
    "oleo_soja": ("SOYA-BEAN OIL",),
    "milho": ("MAIZE",),
    "arroz": ("RICE",),
    "trigo": ("WHEAT",),
    "cafe": ("COFFEE",),
    "acucar": ("SUGAR",),
    "etanol": ("ETHYL ALCOHOL",),
    "algodao": ("COTTON",),
    "carne_bovina": ("BOVINE",),
    "carne_frango": ("GALLUS DOMESTICUS",),
    "carne_suina": ("SWINE",),
    "celulose": ("WOOD PULP",),
    "tabaco": ("TOBACCO",),
    "suco_laranja": ("ORANGE",),
}
ABIOVE_BLOCK = re.compile(r"^(\d+(?:\.\d+)*)\.?\s+(\S.*)$")
ANO_TOKEN = re.compile(r"(?<![0-9])(?:19|20)[0-9]{2}(?![0-9])")
MES_TOKEN = re.compile(
    r"\b(?:JANEIRO|FEVEREIRO|MARCO|ABRIL|MAIO|JUNHO|JULHO|AGOSTO|SETEMBRO|OUTUBRO|NOVEMBRO"
    r"|DEZEMBRO|JAN|FEV|MAR|ABR|MAI|JUN|JUL|AGO|SET|OUT|NOV|DEZ)\b"
)
NORMALIZACAO_TITULO = (
    "maiúsculas sem acento; ano de 4 dígitos -> <ANO>; nome/abreviação de mês -> <MES>; "
    "pontuação -> espaço; espaços colapsados"
)


def normalizar(texto: str) -> str:
    decomposed = unicodedata.normalize("NFD", texto.upper())
    return "".join(char for char in decomposed if unicodedata.category(char) != "Mn")


def normalizar_titulo(texto: str) -> str:
    """Título comparável entre publicações: a atualização normal de ano e mês vira token, para
    que só mudança de mercadoria ou de recorte apareça como divergência."""
    texto = MES_TOKEN.sub("<MES>", ANO_TOKEN.sub("<ANO>", normalizar(texto)))
    return " ".join(re.sub(r"[^A-Z0-9<> ]+", " ", texto).split())


def unit_dictionary(csv_text: str) -> dict[str, dict[str, str]]:
    reader = csv.DictReader(io.StringIO(csv_text), delimiter=";")
    return {row["CO_UNID"]: row for row in reader}


def ncm_dictionary(csv_text: str) -> dict[str, str]:
    reader = csv.DictReader(io.StringIO(csv_text), delimiter=";")
    return {
        row["CO_NCM"]: row["NO_NCM_POR"]
        for row in reader
        if len(row["CO_NCM"]) == 8 and row["CO_NCM"].isdigit()
    }


def hs_reference(payload: bytes) -> dict[str, str]:
    results = json.loads(payload.decode("utf-8"))["results"]
    reference = {}
    for entry in results:
        code = str(entry["id"])
        if len(code) not in (4, 6) or not code.isdigit():
            continue
        text = str(entry["text"])
        prefix = f"{code} - "
        reference[code] = text[len(prefix) :] if text.startswith(prefix) else text
    return reference


def compare_unit_dictionary(dictionary: dict[str, dict[str, str]]) -> dict[str, Any]:
    problems = []
    kg = dictionary.get("10")
    if kg is None:
        problems.append("unidade 10 (quilograma líquido) ausente do dicionário oficial")
    elif "QUILOGRAMA" not in kg.get("NO_UNID", "").upper():
        problems.append(f"unidade 10 mudou de nome: {kg.get('NO_UNID')!r}")
    return {
        "status": "mismatch" if problems else "ok",
        "problems": problems,
        "unidades": len(dictionary),
    }


def _divergencia_ncm(descricao: str, regra: dict[str, Any]) -> list[str]:
    motivos = []
    faltando = [word for word in regra["exige"] if word not in descricao]
    if faltando:
        motivos.append(f"sem {faltando}")
    alternativas = regra["alternativas"]
    if alternativas and not any(word in descricao for word in alternativas):
        motivos.append(f"sem nenhuma de {list(alternativas)}")
    proibidas = [word for word in regra["exclui"] if word in descricao]
    if proibidas:
        motivos.append(f"contém {proibidas}")
    return motivos


def compare_ncm_map(dictionary: dict[str, str]) -> dict[str, Any]:
    problems = []
    entries = []
    for produto in PRODUTOS_COMEX:
        selecao = comexstat_models.resolve_ncm(produto)
        regra = NCM_REGRAS.get(produto)
        if regra is None:
            problems.append(f"{produto}: sem regra declarada em NCM_REGRAS")
            continue
        posicoes = regra["posicoes"]
        for code in selecao.prefixos:
            if not code.isdigit() or not 4 <= len(code) <= 8:
                problems.append(f"{produto}: código NCM fora do padrão {code!r}")
            if code[:4] not in posicoes:
                problems.append(
                    f"{produto}: código {code} está na posição {code[:4]}, "
                    f"não nas posições declaradas {list(posicoes)}"
                )
            if not any(item.startswith(code) for item in dictionary):
                problems.append(f"{produto}: nenhum NCM de 8 dígitos começa por {code!r}")
        codes = sorted(item for item in dictionary if selecao.seleciona(item))
        divergentes = []
        for item in codes:
            motivos = (
                [] if item[:4] in posicoes else [f"posição {item[:4]} fora de {list(posicoes)}"]
            )
            motivos.extend(_divergencia_ncm(normalizar(dictionary[item]), regra))
            if motivos:
                divergentes.append(item)
                problems.append(f"{produto}: NCM {item} {dictionary[item]!r}: {'; '.join(motivos)}")
        entries.append(
            {
                "produto": produto,
                "codes": list(selecao.prefixos),
                "posicoes": list(posicoes),
                "match": "exato" if selecao.codigo_unico else "prefixo",
                "exige": list(regra["exige"]),
                "alternativas": list(regra["alternativas"]),
                "exclui": list(regra["exclui"]),
                "ncms": len(codes),
                "conferidos": codes,
                "divergentes": divergentes,
            }
        )
    return {"status": "mismatch" if problems else "ok", "problems": problems, "entries": entries}


def compare_hs_map(
    descriptions: dict[str, str], produtos: Iterable[str] | None = None
) -> dict[str, Any]:
    problems = []
    entries = []
    alvo = list(comtrade_models.HS_PRODUTOS_AGRO) if produtos is None else list(produtos)
    for produto in alvo:
        codes = comtrade_models.HS_PRODUTOS_AGRO[produto]
        keywords = HS_KEYWORDS.get(produto)
        if keywords is None:
            problems.append(f"{produto}: sem palavra-chave declarada em HS_KEYWORDS")
            continue
        conferidos = []
        for code in codes:
            if len(code) not in (4, 6) or not code.isdigit():
                problems.append(f"{produto}: HS fora do padrão de 4 ou 6 dígitos {code!r}")
                continue
            description = descriptions.get(code)
            if description is None:
                problems.append(f"{produto}: HS {code} ausente da descrição oficial capturada")
                continue
            if not any(word in normalizar(description) for word in keywords):
                problems.append(
                    f"{produto}: HS {code} descrito como {description!r}; "
                    f"nenhuma palavra-chave {keywords} confere"
                )
                continue
            conferidos.append(code)
        entries.append(
            {
                "produto": produto,
                "codes": list(codes),
                "keywords": list(keywords),
                "conferidos": conferidos,
            }
        )
    return {"status": "mismatch" if problems else "ok", "problems": problems, "entries": entries}


def compare_comtrade_count(count_body: bytes, expected_rows: int) -> dict[str, Any]:
    payload = json.loads(count_body.decode("utf-8"))
    count = payload.get("count")
    problems = []
    if not isinstance(count, int):
        problems.append(f"count ausente ou inválido: {count!r}")
    elif count != expected_rows:
        problems.append(f"count live {count} != {expected_rows} do manifesto (revisão da fonte?)")
    return {"status": "mismatch" if problems else "ok", "problems": problems, "count": count}


def abiove_sheet_inventory(sheet: Any) -> dict[str, Any]:
    blocos = []
    ultima: dict[str, Any] = {"cell": None, "value": None}
    ultima_coluna = 0
    ultima_linha = 0
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            text = " ".join(str(cell.value).split())
            if not text:
                continue
            address = f"{get_column_letter(cell.column)}{cell.row}"
            ultima = {"cell": address, "value": text}
            ultima_coluna = max(ultima_coluna, cell.column)
            ultima_linha = max(ultima_linha, cell.row)
            match = ABIOVE_BLOCK.match(text) if isinstance(cell.value, str) else None
            if match is not None:
                blocos.append(
                    {
                        "bloco": match.group(1),
                        "cell": address,
                        "titulo": match.group(2),
                        "titulo_norm": normalizar_titulo(match.group(2)),
                    }
                )
    return {
        "sheet": sheet.title,
        "sheet_norm": normalizar_titulo(sheet.title),
        "blocos": blocos,
        "ultima_celula": ultima,
        "ultima_coluna": get_column_letter(ultima_coluna) if ultima_coluna else None,
        "ultima_linha": ultima_linha,
    }


def abiove_inventory(data: bytes) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    try:
        return {"abas": [abiove_sheet_inventory(sheet) for sheet in workbook.worksheets]}
    finally:
        workbook.close()


def _decisao_incompleta(item: dict[str, Any]) -> str | None:
    if item["estado"] == "mapeada" and not item.get("campo"):
        return "mapeada sem destino declarado"
    if item["estado"] == "ignorada" and not item.get("motivo"):
        return "ignorada sem motivo nominal"
    if item["estado"] not in ("mapeada", "ignorada"):
        return f"estado {item['estado']!r} não decidido"
    return None


def compare_abiove_inventory(
    inventory: dict[str, Any], structure: list[dict[str, Any]]
) -> dict[str, Any]:
    problems = []
    publicadas = {aba["sheet_norm"]: aba for aba in inventory["abas"]}
    previstas = {}
    previstos: dict[tuple[str, str], dict[str, Any]] = {}
    extents = {}
    for item in structure:
        falha = _decisao_incompleta(item)
        locator = item["locator"]
        if item["kind"] == "sheet":
            previstas[locator["sheet_norm"]] = item
            if falha:
                problems.append(f"aba {locator['sheet']!r}: {falha}")
        elif item["kind"] == "sheet_table":
            previstos[(locator["sheet_norm"], locator["bloco"])] = item
            if falha:
                problems.append(f"bloco {locator['bloco']}: {falha}")
        elif item["kind"] == "extent":
            extents[locator["sheet_norm"]] = item
    for nome in sorted(set(publicadas) - set(previstas)):
        problems.append(f"aba {publicadas[nome]['sheet']!r} publicada e não classificada")
    for nome in sorted(set(previstas) - set(publicadas)):
        problems.append(f"aba {previstas[nome]['locator']['sheet']!r} prevista e ausente")
    publicados = {
        (aba["sheet_norm"], bloco["bloco"]): bloco
        for aba in inventory["abas"]
        for bloco in aba["blocos"]
    }
    for aba, bloco in sorted(set(publicados) - set(previstos)):
        problems.append(f"bloco {bloco} da aba {aba} publicado e não previsto no manifesto")
    for aba, bloco in sorted(set(previstos) - set(publicados)):
        problems.append(f"bloco {bloco} da aba {aba} previsto no manifesto e ausente")
    for chave in sorted(set(publicados) & set(previstos)):
        esperado = previstos[chave]["locator"]["titulo_norm"]
        publicado = publicados[chave]["titulo_norm"]
        if publicado != esperado:
            problems.append(
                f"bloco {chave[1]}: título publicado {publicado!r} != {esperado!r} do manifesto"
            )
    for nome, aba in sorted(publicadas.items()):
        extent = extents.get(nome)
        if extent is None:
            problems.append(f"aba {aba['sheet']!r} sem item de extensão lateral no manifesto")
        elif extent["locator"]["ultima_coluna"] != aba["ultima_coluna"]:
            problems.append(
                f"aba {aba['sheet']!r}: extensão lateral do manifesto "
                f"{extent['locator']['ultima_coluna']} x publicação {aba['ultima_coluna']}"
            )
    return {
        "status": "mismatch" if problems else "ok",
        "problems": problems,
        "normalizacao": NORMALIZACAO_TITULO,
        "abas_publicadas": sorted(aba["sheet"] for aba in inventory["abas"]),
        "abas_previstas": len(previstas),
        "blocos_publicados": len(publicados),
        "blocos_previstos": len(previstos),
        "titulos_conferidos": len(set(publicados) & set(previstos)),
        "extremos": {
            aba["sheet"]: {
                "ultima_coluna": aba["ultima_coluna"],
                "ultima_celula": aba["ultima_celula"],
            }
            for aba in inventory["abas"]
        },
    }


async def fetch_comexstat_dictionary(url: str) -> bytes:
    """Aquisição de conferência isolada: mesmo host, cabeçalhos, TLS e orçamento do cliente Comex
    Stat, sem acrescentar a tabela ao catálogo público `COMEXSTAT_DICTIONARY_URLS`."""
    if not url.startswith(f"https://{constants.COMEXSTAT_DOWNLOAD_HOST}/balanca/bd/tabelas/"):
        raise ValueError("Comex Stat: dicionário de conferência fora da pasta oficial de tabelas")
    headers = httpx.Headers(UserAgentRotator.get_headers(source="comexstat"))
    headers["Accept-Encoding"] = "identity"
    async with httpx.AsyncClient(
        timeout=comexstat_client.TIMEOUT,
        headers=headers,
        verify=_tls.build_context(),
        follow_redirects=False,
    ) as http:
        response = await http.get(url)
        response.raise_for_status()
        if len(response.content) > constants.COMEXSTAT_MAX_DICTIONARY_BYTES:
            raise ValueError("Comex Stat: dicionário de conferência excede o orçamento de bytes")
        return response.content


def registro(content: bytes, source: str, url: str) -> dict[str, Any]:
    return {
        "source": source,
        "url": url,
        "sha256": hashlib.sha256(content).hexdigest(),
        "bytes": len(content),
    }


async def run(output: Path) -> int:
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    expected_rows = next(
        case["period"]["rows"]
        for case in manifest["cases"]
        if case["id"] == "comtrade_soja_br_x_2023"
    )
    abiove_case = next(case for case in manifest["cases"] if case["source"] == "abiove")
    report: dict[str, Any] = {
        "fetched_at": datetime.now(UTC).isoformat(),
        "structure": [],
        "values": [],
    }
    async with comexstat_client.open_dictionary(tabela="unidades") as bundle:
        bundle.file.seek(0)
        content = bundle.file.read()
    report["structure"].append(
        {
            "case": "comexstat_unit_dictionary",
            **compare_unit_dictionary(unit_dictionary(content.decode("latin-1"))),
        }
    )
    report["values"].append(
        registro(
            content,
            "comexstat_dicionario_unidades",
            constants.COMEXSTAT_DICTIONARY_URLS["unidades"],
        )
    )
    ncm_content = await fetch_comexstat_dictionary(NCM_DICTIONARY_URL)
    report["structure"].append(
        {
            "case": "comexstat_ncm_map",
            **compare_ncm_map(ncm_dictionary(ncm_content.decode("latin-1"))),
        }
    )
    report["values"].append(registro(ncm_content, "comexstat_dicionario_ncm", NCM_DICTIONARY_URL))
    async with httpx.AsyncClient(timeout=120, follow_redirects=True) as http:
        count_response = await http.get(
            COMTRADE_COUNT_URL, headers=UserAgentRotator.get_headers(source="comtrade")
        )
        count_response.raise_for_status()
        report["structure"].append(
            {
                "case": "comtrade_soja_count",
                **compare_comtrade_count(count_response.content, expected_rows),
            }
        )
        year = datetime.now(UTC).year
        base = constants.URLS[constants.Fonte.ABIOVE]["exportacao"]
        latest = None
        for month in range(12, 0, -1):
            head = await http.head(f"{base}/exp_{year:04d}{month:02d}.xlsx")
            if head.status_code == 200:
                latest = {"url": str(head.url), "bytes": head.headers.get("content-length")}
                break
        report["structure"].append(
            {
                "case": "abiove_disponibilidade",
                "status": "ok" if latest else "mismatch",
                "problems": [] if latest else [f"nenhum exp_{year}MM.xlsx respondeu 200"],
                "latest": latest,
            }
        )
        if latest is None:
            report["structure"].append(
                {
                    "case": "abiove_inventario",
                    "status": "mismatch",
                    "problems": ["workbook mais recente indisponível: inventário não executado"],
                }
            )
        else:
            workbook_response = await http.get(
                latest["url"], headers=UserAgentRotator.get_headers(source="abiove")
            )
            workbook_response.raise_for_status()
            report["structure"].append(
                {
                    "case": "abiove_inventario",
                    "workbook": latest["url"],
                    **compare_abiove_inventory(
                        abiove_inventory(workbook_response.content), abiove_case["structure"]
                    ),
                }
            )
            report["values"].append(
                registro(workbook_response.content, "abiove_workbook", latest["url"])
            )
        reference_response = await http.get(
            COMTRADE_HS_REFERENCE_URL, headers=UserAgentRotator.get_headers(source="comtrade")
        )
        reference_response.raise_for_status()
    report["structure"].append(
        {"case": "comtrade_hs_map", **compare_hs_map(hs_reference(reference_response.content))}
    )
    report["values"].append(
        registro(reference_response.content, "comtrade_referencia_hs", COMTRADE_HS_REFERENCE_URL)
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, indent=1, ensure_ascii=False), encoding="utf-8")
    mismatches = [entry for entry in report["structure"] if entry["status"] != "ok"]
    for entry in report["structure"]:
        print(entry["status"], entry["case"], "; ".join(entry["problems"]))
    print(
        f"{len(report['structure']) - len(mismatches)} ok / {len(mismatches)} mismatch -> {output}"
    )
    return 1 if mismatches else 0


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Inventário N1 live do comércio exterior (ComexStat, ABIOVE, Comtrade)"
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=ROOT / f"reports/reconciliacao_comercio_{datetime.now(UTC):%Y%m%d}.json",
    )
    arguments = parser.parse_args()
    return asyncio.run(run(arguments.output))


if __name__ == "__main__":
    sys.exit(main())
