from __future__ import annotations

import argparse
import asyncio
import hashlib
import io
import json
import re
import unicodedata
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import openpyxl
import pdfplumber
import xlrd

from agrobr.anda import client as anda_client
from agrobr.anec import client as anec_client
from agrobr.deral import client as deral_client

GOLDEN = Path(__file__).resolve().parents[1] / "tests/golden_data/reconciliacao_r7_20260918"
PROFILES = GOLDEN / "structure_profiles.json"


def _label(value: str) -> str:
    value = unicodedata.normalize("NFKD", value)
    value = "".join(char for char in value if not unicodedata.combining(char)).casefold()
    if value.strip() == "%":
        return "%"
    value = re.sub(r"[+-]?\d[\d.,]*(?:st|nd|rd|th)?\s*%?", " ", value)
    value = re.sub(r"[^a-z%/()]+", " ", value)
    return " ".join(value.split()) if re.search(r"[a-z]", value) else ""


def _page_labels(page: Any) -> list[dict[str, Any]]:
    positions: dict[float, list[dict[str, Any]]] = {}
    for char in page.chars:
        positions.setdefault(round(char["top"], 1), []).append(char)
    labels = []
    for top, chars in sorted(positions.items()):
        text = "".join(char["text"] for char in sorted(chars, key=lambda char: char["x0"]))
        label = _label(text)
        if label:
            labels.append({"top": top, "label": label})
    return labels


def inventory_pdf(raw: bytes) -> dict[str, Any]:
    with pdfplumber.open(io.BytesIO(raw)) as document:
        pages = [
            {
                "page": ordinal,
                "size": [round(page.width, 1), round(page.height, 1)],
                "images": len(page.images),
                "labels": _page_labels(page),
            }
            for ordinal, page in enumerate(document.pages, 1)
        ]
    return {"format": "pdf", "pages": pages}


def _xls_sheet(sheet: Any, book: Any) -> dict[str, Any]:
    cells = []
    for row in range(sheet.nrows):
        for column in range(sheet.ncols):
            item = sheet.cell(row, column)
            if item.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                continue
            location = f"{openpyxl.utils.get_column_letter(column + 1)}{row + 1}"
            style = book.xf_list[sheet.cell_xf_index(row, column)]
            number_format = book.format_map[style.format_key].format_str
            cells.append(
                {
                    "cell": location,
                    "type": item.ctype,
                    "label": _label(str(item.value)) if item.ctype == xlrd.XL_CELL_TEXT else "",
                    "percent_format": "%" in number_format,
                }
            )
    return {"name": sheet.name, "cells": cells, "merged": [list(r) for r in sheet.merged_cells]}


def inventory_workbook(raw: bytes) -> dict[str, Any]:
    if raw.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
        book = xlrd.open_workbook(file_contents=raw, formatting_info=True)
        try:
            return {"format": "xls", "sheets": [_xls_sheet(sheet, book) for sheet in book.sheets()]}
        finally:
            book.release_resources()
    workbook = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
    try:
        sheets = [
            {
                "name": sheet.title,
                "cells": [
                    {
                        "cell": item.coordinate,
                        "type": item.data_type,
                        "label": _label(str(item.value)) if item.data_type == "s" else "",
                        "percent_format": "%" in item.number_format,
                    }
                    for row in sheet
                    for item in row
                    if item.value is not None
                ],
                "merged": [str(region) for region in sheet.merged_cells.ranges],
            }
            for sheet in workbook
        ]
    finally:
        workbook.close()
    return {"format": "xlsx", "sheets": sheets}


def inventory(raw: bytes) -> dict[str, Any]:
    return inventory_pdf(raw) if raw.startswith(b"%PDF-") else inventory_workbook(raw)


def compare_inventory(actual: dict[str, Any], expected: dict[str, Any]) -> dict[str, Any]:
    problems = []
    if actual.get("format") != expected.get("format"):
        problems.append("formato mudou")
    key = "pages" if expected.get("format") == "pdf" else "sheets"
    current, previous = actual.get(key, []), expected.get(key, [])
    if not current or not previous:
        problems.append("inventário vazio não comprova estrutura")
    if len(current) != len(previous):
        problems.append(f"{key}: quantidade mudou de {len(previous)} para {len(current)}")
    for ordinal, (old, new) in enumerate(zip(previous, current, strict=False), 1):
        if old != new:
            changes = sorted(
                name for name in old.keys() | new.keys() if old.get(name) != new.get(name)
            )
            problems.append(f"{key} {ordinal}: campos alterados {changes}")
    return {"status": "mismatch" if problems else "ok", "problems": problems}


async def _live_body(source: str, reference: dict[str, Any]) -> tuple[bytes, str]:
    if source == "anda":
        html = await anda_client.fetch_estatisticas_page()
        links = anda_client.parse_links_from_html(html, pattern=r"\.pdf")
        target = anda_client._select_pdf_target(links, str(reference["year"]))
        if target is None:
            raise ValueError(f"ANDA sem PDF para {reference['year']}")
        return await anda_client.download_file(target["url"]), target["url"]
    if source == "deral":
        return await deral_client.fetch_pc_xls(), reference["url"]
    body, url, _article = await anec_client.fetch_latest_pdf(
        year=reference["year"], use_cache=False
    )
    return body, url


async def sweep(
    profiles: dict[str, Any], *, live: bool = False, source: str | None = None
) -> list[dict[str, Any]]:
    results = []
    for reference in profiles["cases"]:
        if source and reference["source"] != source:
            continue
        if live and not reference["latest"]:
            continue
        body, url = (
            await _live_body(reference["source"], reference)
            if live
            else ((GOLDEN / reference["file"]).read_bytes(), reference["url"])
        )
        observed = inventory(body)
        results.append(
            {
                "id": reference["id"],
                "url": url,
                "observed_at": datetime.now(UTC).isoformat(),
                "sha256": hashlib.sha256(body).hexdigest(),
                "bytes": len(body),
                **compare_inventory(observed, reference["inventory"]),
            }
        )
    return results


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--live", action="store_true")
    parser.add_argument("--source", choices=("anda", "deral", "anec"))
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    profiles = json.loads(PROFILES.read_text(encoding="utf-8"))
    results = asyncio.run(sweep(profiles, live=args.live, source=args.source))
    report = {
        "scope": profiles["scope"],
        "mode": "live" if args.live else "preserved_fixtures",
        "results": results,
        "limits": profiles["limits"],
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return int(not results or any(item["status"] != "ok" for item in results))


if __name__ == "__main__":
    raise SystemExit(main())
