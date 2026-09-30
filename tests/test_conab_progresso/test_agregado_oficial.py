from __future__ import annotations

import io
import json
import re
import unicodedata
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from agrobr import datasets
from agrobr.conab.progresso import api, parser
from agrobr.deral import parser as deral_parser
from agrobr.exceptions import InvalidParameterError
from tests import helpers
from tests.helpers import levanta_exatamente

GOLDEN = Path(__file__).resolve().parents[1] / "golden_data"
OFICIAL = GOLDEN / "conab_progresso" / "oficial_r26"
MANIFESTO = json.loads((OFICIAL / "manifest.json").read_text(encoding="utf-8"))
URLS = {item["arquivo"]: item["url"] for item in MANIFESTO["arquivos"]}
PLANILHAS = [
    OFICIAL / "progresso_20260222.xlsx",
    OFICIAL / "progresso_20260920.xlsx",
    GOLDEN / "conab_progresso" / "progresso_20260828.xlsx",
]
PUBLICADAS = [
    *PLANILHAS,
    GOLDEN / "conab_progresso" / "historico_20250927" / "response.xlsx",
    GOLDEN / "conab_progresso" / "historico_20250927" / "latest.xlsx",
]
COLUNAS_PUBLICADAS = [
    "cultura",
    "operacao",
    "semana_atual",
    "pct_ano_anterior",
    "pct_semana_anterior",
    "pct_semana_atual",
    "pct_media_5_anos",
]
NOTA = re.compile(r"\(Esses (\d+) estados correspondem a (\d+(?:,\d+)?)% da área cultivada\)")
DOC = Path(__file__).resolve().parents[2] / "docs" / "api" / "conab_progresso.md"
DOC_EN = DOC.with_name("conab_progresso.en.md")
NOTA_DA_DOC = re.compile(r"(\d+) (?:estados|states) \((\d+(?:[.,]\d+)?)% ")


def _agregados(bruto: bytes) -> dict[tuple[str, str, str], dict]:
    folha = openpyxl.load_workbook(io.BytesIO(bruto), data_only=True).worksheets[0]
    saida: dict[tuple[str, str, str], dict] = {}
    cultura = nota = operacao = semana = None
    for linha in range(1, folha.max_row + 1):
        rotulo = folha.cell(linha, 2).value
        if isinstance(folha.cell(linha, 5).value, datetime):
            semana = folha.cell(linha, 5).value.strftime("%Y-%m-%d")
        if rotulo is None:
            continue
        texto = str(rotulo).strip()
        if " - Safra " in texto:
            cultura, nota, operacao = texto.split(" - Safra ")[0].strip(), None, None
        elif achado := NOTA.fullmatch(texto):
            nota = (int(achado[1]), float(achado[2].replace(",", ".")) / 100)
        elif achado := re.match(r"(Semeadura|Colheita)", texto):
            operacao = achado[1]
        elif achado := re.fullmatch(r"(\d+) estados", texto):
            saida[(cultura, operacao, semana)] = {
                "n": int(achado[1]),
                "nota": nota,
                "valores": [folha.cell(linha, coluna).value for coluna in range(3, 7)],
            }
    return saida


@pytest.mark.parametrize("planilha", PLANILHAS, ids=lambda caminho: caminho.name)
def test_agregado_de_estados_sai_como_media_com_a_nota_publicada(planilha):
    esperado = _agregados(planilha.read_bytes())
    assert esperado

    frame = parser.parse_progresso_xlsx(planilha.read_bytes())

    media = frame[frame["uf"] == "MEDIA_ESTADOS"]
    obtido = {
        (linha.cultura, linha.operacao, linha.semana_atual): linha for linha in media.itertuples()
    }
    assert set(obtido) == set(esperado)
    for chave, celulas in esperado.items():
        linha = obtido[chave]
        assert linha.n_estados == celulas["n"] == celulas["nota"][0], chave
        assert linha.cobertura_area_pct == pytest.approx(celulas["nota"][1], rel=1e-12), chave
        publicados = [
            linha.pct_ano_anterior,
            linha.pct_semana_anterior,
            linha.pct_semana_atual,
            linha.pct_media_5_anos,
        ]
        assert publicados == pytest.approx(celulas["valores"], rel=1e-12), chave
    assert "BR" not in set(frame["uf"])
    ufs = frame[frame["uf"] != "MEDIA_ESTADOS"]
    assert ufs["n_estados"].isna().all() and ufs["cobertura_area_pct"].isna().all()


def test_linha_brasil_publicada_sai_como_br_sem_nota():
    livro = openpyxl.load_workbook(OFICIAL / "progresso_20260920.xlsx")
    folha = livro.worksheets[0]
    celula = next(
        folha.cell(linha, 2)
        for linha in range(1, folha.max_row + 1)
        if folha.cell(linha, 2).value == "12 estados"
    )
    celula.value = "Brasil"
    buffer = io.BytesIO()
    livro.save(buffer)

    frame = parser.parse_progresso_xlsx(buffer.getvalue())

    brasil = frame[frame["uf"] == "BR"]
    assert brasil[["cultura", "operacao"]].values.tolist() == [["Soja", "Semeadura"]]
    assert brasil["n_estados"].isna().all() and brasil["cobertura_area_pct"].isna().all()


def _servir_boletim_de_20260920(monkeypatch) -> tuple[str, bytes]:
    recurso = URLS["progresso_20260920.xlsx"]
    url = recurso.rsplit("/", 1)[0]
    bruto = (OFICIAL / "progresso_20260920.xlsx").read_bytes()
    pagina = f'<div id="content-core"><a href="{recurso}">Plantio e colheita</a></div>'
    helpers.mock_progresso_http(monkeypatch, {url: pagina.encode(), recurso: bruto})
    return url, bruto


async def test_estado_br_recusado_com_a_cobertura_publicada(monkeypatch):
    url, _ = _servir_boletim_de_20260920(monkeypatch)

    with levanta_exatamente(InvalidParameterError, match="não publica Brasil") as recusa:
        await api.progresso_safra(produto="soja", uf="BR", semana_url=url)
    assert "12 estados, 96.0% da área" in str(recusa.value)
    assert "MEDIA_ESTADOS" in str(recusa.value)

    with levanta_exatamente(InvalidParameterError, match="não publica Brasil"):
        await datasets.progresso_safra("soja", uf="BR", semana_url=url)

    media, meta = await datasets.progresso_safra(
        "soja", uf="MEDIA_ESTADOS", semana_url=url, return_meta=True
    )
    assert media[["uf", "n_estados", "cobertura_area_pct"]].values.tolist() == [
        ["MEDIA_ESTADOS", 12, 0.96]
    ]
    assert meta.contract_version == meta.schema_version == "2.0"


def test_pr_da_conab_repete_o_levantamento_do_deral():
    conab = parser.parse_progresso_xlsx((OFICIAL / "progresso_20260920.xlsx").read_bytes())
    deral = deral_parser.parse_pc_xls(
        (GOLDEN / "deral" / "pc_20260915" / "response.xls").read_bytes()
    )
    deral = deral.drop_duplicates(["produto", "data"]).set_index(["produto", "data"])
    pares = {
        ("Feijão 1ª", "Semeadura"): ("feijao_1", "plantio_pct"),
        ("Milho 1ª", "Semeadura"): ("milho_1", "plantio_pct"),
        ("Milho 2ª", "Colheita"): ("milho_2", "colheita_pct"),
        ("Soja", "Semeadura"): ("soja", "plantio_pct"),
        ("Trigo", "Colheita"): ("trigo", "colheita_pct"),
    }
    pr = conab[conab["uf"] == "PR"].set_index(["cultura", "operacao"])
    for (cultura, operacao), (produto, coluna) in pares.items():
        linha = pr.loc[(cultura, operacao)]
        assert linha["semana_atual"] == "2026-09-18"
        assert linha["pct_semana_atual"] == pytest.approx(
            deral.loc[(produto, "14/09/2026"), coluna] / 100
        )
        assert linha["pct_semana_anterior"] == pytest.approx(
            deral.loc[(produto, "08/09/2026"), coluna] / 100
        )


def test_tabela_de_culturas_e_operacoes_da_doc_confere_as_planilhas():
    publicadas: dict[str, set[str]] = {}
    for planilha in [
        *PLANILHAS,
        GOLDEN / "conab_progresso" / "historico_20250927" / "response.xlsx",
    ]:
        frame = parser.parse_progresso_xlsx(planilha.read_bytes())
        for cultura, operacao in (
            frame[["cultura", "operacao"]].drop_duplicates().itertuples(index=False)
        ):
            publicadas.setdefault(cultura, set()).add(operacao)
    notas = {
        cultura: celulas["nota"]
        for (cultura, _, _), celulas in _agregados(
            (OFICIAL / "progresso_20260920.xlsx").read_bytes()
        ).items()
    }
    for doc, titulo in [(DOC, "### Culturas Disponiveis"), (DOC_EN, "### Available Crops")]:
        tabela = doc.read_text(encoding="utf-8").split(titulo, 1)[1].split("\n\n", 2)[1]
        linhas = [linha.split("|") for linha in tabela.splitlines()[2:]]
        operacoes = {
            celulas[1].strip(): {operacao.strip() for operacao in celulas[3].split(",")}
            for celulas in linhas
        }
        cobertura = {
            celulas[1].strip(): (int(achado[1]), float(achado[2].replace(",", ".")) / 100)
            for celulas in linhas
            if (achado := NOTA_DA_DOC.match(celulas[2].strip()))
        }
        assert operacoes == publicadas, doc.name
        assert cobertura == notas, doc.name


def _sem_acento(texto: str) -> str:
    return "".join(
        caractere
        for caractere in unicodedata.normalize("NFKD", texto)
        if not unicodedata.combining(caractere)
    )


async def test_culturas_da_doc_filtram_a_cultura_publicada(monkeypatch):
    url, bruto = _servir_boletim_de_20260920(monkeypatch)
    parametro = next(
        linha
        for linha in DOC.read_text(encoding="utf-8").splitlines()
        if linha.startswith("| `produto` |")
    )
    documentadas = re.findall(r'"([^"]+)"', parametro)
    publicado = parser.parse_progresso_xlsx(bruto)
    assert sorted(documentadas) == sorted(set(publicado["cultura"].map(_sem_acento)))

    for nome in documentadas:
        frame = await api.progresso_safra(produto=nome, semana_url=url)
        esperado = publicado[publicado["cultura"].map(_sem_acento) == nome]
        assert not esperado.empty and frame.equals(esperado.reset_index(drop=True)), nome


def _fracao(valor: object) -> float | None:
    if isinstance(valor, str):
        return float(valor.replace("*", "").strip().removesuffix("%")) / 100
    return None if valor is None else float(valor)


def _linhas_publicadas(bruto: bytes) -> list[tuple]:
    folha = openpyxl.load_workbook(io.BytesIO(bruto), data_only=True).worksheets[0]
    saida = []
    cultura = operacao = semana = None
    for linha in range(1, folha.max_row + 1):
        texto = str(folha.cell(linha, 2).value or "").strip()
        valores = [folha.cell(linha, coluna).value for coluna in range(3, 7)]
        if isinstance(valores[2], datetime):
            semana = valores[2].strftime("%Y-%m-%d")
        if " - Safra " in texto:
            cultura = texto.split(" - Safra ")[0].strip()
        elif achado := re.match(r"(Semeadura|Colheita)", texto):
            operacao = achado[1]
        elif texto and any(
            isinstance(valor, int | float) or "%" in str(valor or "") for valor in valores
        ):
            saida.append((cultura, operacao, semana, *map(_fracao, valores)))
    return saida


@pytest.mark.parametrize("planilha", PUBLICADAS, ids=lambda caminho: caminho.name)
def test_todas_as_linhas_publicadas_saem_com_os_percentuais_das_celulas(planilha):
    bruto = planilha.read_bytes()
    esperado = _linhas_publicadas(bruto)
    assert esperado

    frame = parser.parse_progresso_xlsx(bruto)

    obtido = [
        tuple(None if pd.isna(valor) else valor for valor in linha)
        for linha in frame[COLUNAS_PUBLICADAS].itertuples(index=False)
    ]
    assert sorted(obtido, key=repr) == sorted(esperado, key=repr)


async def test_trecho_do_nome_traz_as_culturas_que_o_contem(monkeypatch):
    url, bruto = _servir_boletim_de_20260920(monkeypatch)
    publicado = parser.parse_progresso_xlsx(bruto)

    frame = await api.progresso_safra(produto="milho", semana_url=url)

    milho = publicado[publicado["cultura"].isin(["Milho 1ª", "Milho 2ª"])]
    assert not milho.empty and frame.equals(milho.reset_index(drop=True))


@pytest.mark.parametrize(
    ("produto", "culturas", "linhas"),
    [
        ("milho", {"Milho 1ª", "Milho 2ª"}, 20),
        ("Milho 1ª", {"Milho 1ª"}, 10),
        ("milho_1", {"Milho 1ª"}, 10),
        ("MILHO 1A", {"Milho 1ª"}, 10),
    ],
)
async def test_produto_casa_nome_publicado_ou_alias_da_doc(monkeypatch, produto, culturas, linhas):
    url, _ = _servir_boletim_de_20260920(monkeypatch)
    frame = await api.progresso_safra(produto=produto, semana_url=url)
    assert set(frame["cultura"]) == culturas
    assert len(frame) == linhas
