from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from agrobr.conab.progresso.parser import parse_progresso_xlsx
from agrobr.exceptions import ParseError

GOLDEN_DIR = Path(__file__).resolve().parent.parent / "golden_data" / "conab_progresso"


@pytest.fixture()
def golden_xlsx() -> bytes:
    return (GOLDEN_DIR / "progresso_sample" / "response.xlsx").read_bytes()


@pytest.fixture()
def expected() -> dict:
    return json.loads(
        (GOLDEN_DIR / "progresso_sample" / "expected.json").read_text(encoding="utf-8")
    )


class TestParseEdgeCases:
    def test_empty_xlsx_raises(self) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Progresso de safra"
        import io

        buf = io.BytesIO()
        wb.save(buf)
        with pytest.raises(ParseError, match="Sheet vazia"):
            parse_progresso_xlsx(buf.getvalue())

    def test_no_records_raises(self) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Progresso de safra"
        ws["B1"] = "Header only"
        ws["B2"] = "No data here"
        import io

        buf = io.BytesIO()
        wb.save(buf)
        with pytest.raises(ParseError, match="Nenhum registro"):
            parse_progresso_xlsx(buf.getvalue())
