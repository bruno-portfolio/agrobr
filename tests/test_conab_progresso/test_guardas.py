from __future__ import annotations

import io
from collections.abc import Callable
from datetime import datetime
from pathlib import Path
from unittest.mock import AsyncMock

import openpyxl
import pytest

from agrobr import datasets
from agrobr.conab.progresso import api, client, parser
from agrobr.exceptions import InvalidParameterError, ParseError, SourceUnavailableError
from tests import helpers
from tests.helpers import levanta_exatamente, sem_excecao

OFICIAL = (
    Path(__file__).resolve().parents[1] / "golden_data" / "conab_progresso" / "oficial_20260925"
)


def _linha(folha, rotulo: str) -> int:
    return next(r for r in range(1, folha.max_row + 1) if folha.cell(r, 2).value == rotulo)


def _alterada(alterar: Callable, planilha: str = "progresso_20260920.xlsx") -> bytes:
    livro = openpyxl.load_workbook(OFICIAL / planilha)
    alterar(livro)
    buffer = io.BytesIO()
    livro.save(buffer)
    return buffer.getvalue()


def _duas_abas_de_progresso(livro) -> None:
    livro.copy_worksheet(livro.worksheets[0]).title = "Progresso de safra 2"


def _duas_abas_sem_progresso(livro) -> None:
    livro.worksheets[0].title = "Plantio"
    livro.create_sheet("Notas")


def _aba_de_notas(livro) -> None:
    livro.create_sheet("Notas")


def _bloco_sem_operacao(livro) -> None:
    folha = livro.worksheets[0]
    folha.cell(_linha(folha, "Semeadura"), 2).value = None


def _datas_fora_de_ordem(livro) -> None:
    folha = livro.worksheets[0]
    linha = next(
        r for r in range(1, folha.max_row + 1) if isinstance(folha.cell(r, 5).value, datetime)
    )
    folha.cell(linha, 4).value, folha.cell(linha, 5).value = (
        folha.cell(linha, 5).value,
        folha.cell(linha, 4).value,
    )


def _semeadura_do_arroz_sem_observacoes(livro) -> None:
    folha = livro.worksheets[0]
    arroz = next(
        r
        for r in range(1, folha.max_row + 1)
        if str(folha.cell(r, 2).value or "").startswith("Arroz - Safra")
    )
    colheita = next(
        r
        for r in range(arroz, folha.max_row + 1)
        if str(folha.cell(r, 2).value or "").startswith("Colheita")
    )
    for linha in range(arroz, colheita):
        rotulo = folha.cell(linha, 2)
        if rotulo.value is not None and any(
            isinstance(folha.cell(linha, coluna).value, int | float) for coluna in range(3, 7)
        ):
            rotulo.value = None


def _nota_de_outro_numero(livro) -> None:
    folha = livro.worksheets[0]
    nota = _linha(folha, "(Esses 12 estados correspondem a 96% da área cultivada)")
    folha.cell(nota, 2).value = "(Esses 11 estados correspondem a 96% da área cultivada)"


def _uf_repetida(livro) -> None:
    folha = livro.worksheets[0]
    primeira = _linha(folha, "Mato Grosso")
    segunda = next(
        r
        for r in range(primeira + 1, folha.max_row + 1)
        if folha.cell(r, 2).value not in (None, "Mato Grosso")
        and folha.cell(r, 3).value is not None
    )
    folha.cell(segunda, 2).value = "Mato Grosso"


@pytest.mark.parametrize(
    "alterar,motivo",
    [
        (_duas_abas_de_progresso, "Abas de progresso ambíguas"),
        (_duas_abas_sem_progresso, "Abas de progresso ambíguas"),
        (_bloco_sem_operacao, "Bloco sem operação ou observações"),
        (_datas_fora_de_ordem, "Colunas semanais incompletas ou fora de ordem"),
        (_nota_de_outro_numero, "Agregado de 12 estados sob a nota de 11 estados"),
        (_uf_repetida, "Observações de progresso duplicadas"),
    ],
)
def test_planilha_fora_do_layout_publicado_e_recusada(alterar, motivo):
    with levanta_exatamente(ParseError, match=motivo):
        parser.parse_progresso_xlsx(_alterada(alterar))


async def test_listagem_sem_semanas_levanta(monkeypatch):
    helpers.mock_progresso_http(monkeypatch, {client.BASE_URL: b"<html><body></body></html>"})
    with levanta_exatamente(SourceUnavailableError, match="Nenhuma semana encontrada"):
        await api.progresso_safra()


def test_operacao_sem_observacoes_antes_da_seguinte_e_recusada():
    bruto = _alterada(_semeadura_do_arroz_sem_observacoes, "progresso_20260222.xlsx")
    with levanta_exatamente(
        ParseError, match="Bloco sem operação ou observações: Arroz, Semeadura"
    ):
        parser.parse_progresso_xlsx(bruto)


def test_aba_extra_nao_desvia_da_aba_de_progresso():
    sem_notas = parser.parse_progresso_xlsx(_alterada(lambda _livro: None))
    with sem_excecao():
        com_notas = parser.parse_progresso_xlsx(_alterada(_aba_de_notas))
    assert com_notas.equals(sem_notas)


async def test_semana_so_com_pdf_levanta_sem_baixar(monkeypatch):
    url = f"{client.BASE_URL}/acompanhamento-so-com-pdf"
    pagina = b'<div id="content-core"><a href="plantio-e-colheita.pdf">Plantio e colheita</a></div>'
    chamadas = helpers.mock_progresso_http(monkeypatch, {url: pagina})
    with levanta_exatamente(SourceUnavailableError, match="Link plantio/colheita nao encontrado"):
        await api.progresso_safra(semana_url=url)
    assert chamadas == [url]


@pytest.mark.parametrize(
    ("filtro", "motivo"),
    [
        ({"produto": "sojaa"}, "Produto inválido.*Valores válidos: Algodão, Arroz"),
        ({"produto": " "}, "Produto inválido"),
        ({"produto": "1"}, "Produto inválido"),
        ({"produto": "o"}, "Produto inválido"),
        ({"uf": "XX"}, "UF inválida.*MEDIA_ESTADOS"),
        ({"uf": "Mato Grosso"}, "UF inválida"),
        ({"uf": 51}, "UF inválida"),
        ({"operacao": "plantio"}, "Operação inválida.*Semeadura, Colheita"),
    ],
)
async def test_filtro_fora_do_publicado_recusado_antes_da_rede(monkeypatch, filtro, motivo):
    baixar = AsyncMock()
    monkeypatch.setattr(client, "fetch_latest", baixar)
    with levanta_exatamente(InvalidParameterError, match=motivo):
        await api.progresso_safra(**filtro)
    baixar.assert_not_awaited()


@pytest.mark.parametrize(
    ("filtro", "motivo"),
    [({"uf": "XX"}, "UF inválida"), ({"operacao": "plantio"}, "Operação inválida")],
)
async def test_dataset_recusa_filtro_fora_do_publicado_antes_da_rede(monkeypatch, filtro, motivo):
    baixar = AsyncMock()
    monkeypatch.setattr(client, "fetch_latest", baixar)
    with levanta_exatamente(InvalidParameterError, match=motivo):
        await datasets.progresso_safra("soja", **filtro)
    baixar.assert_not_awaited()


@pytest.mark.parametrize("max_pages", [0, -1, True])
async def test_max_pages_nao_positivo_recusado_sem_requisicao(monkeypatch, max_pages):
    monkeypatch.setattr(client, "httpx", None)
    with levanta_exatamente(InvalidParameterError, match="max_pages deve ser inteiro positivo"):
        await api.semanas_disponiveis(max_pages=max_pages)
