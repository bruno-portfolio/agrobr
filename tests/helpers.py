from __future__ import annotations

import asyncio
import copy
import csv
import hashlib
import inspect
import io
import json
import math
import os
import re
import socket
import time
from collections.abc import AsyncIterator, Callable, Coroutine, Iterator
from contextlib import (
    AbstractContextManager,
    ExitStack,
    asynccontextmanager,
    contextmanager,
    suppress,
)
from datetime import UTC, date, datetime
from numbers import Real
from pathlib import Path
from types import SimpleNamespace
from typing import Any
from unittest.mock import AsyncMock, MagicMock
from urllib.parse import parse_qsl, unquote, urlsplit

import duckdb
import httpx
import openpyxl
import pandas as pd
import pytest
import structlog

from agrobr import _log, config, constants, contracts, datasets
from agrobr.alt.anp_diesel import client as anp_client
from agrobr.alt.antt_pedagio import client as antt_client
from agrobr.cache import duckdb_store
from agrobr.comexstat import client as comexstat_client
from agrobr.conab._custo_producao import _acquisition as custos_acquisition
from agrobr.conab._custo_producao import _context as custos_context
from agrobr.conab._custo_producao import _workbook as custos_workbook
from agrobr.conab._custo_producao import models as custos_models
from agrobr.conab._serie_historica import client as serie_historica_client
from agrobr.conab.progresso import client as progresso_client
from agrobr.contracts import desmatamento as desmatamento_contracts
from agrobr.http import wfs_transport
from agrobr.http.rate_limiter import RateLimiter
from agrobr.mapbiomas import resources as mapbiomas_resources
from agrobr.models import MetaInfo
from agrobr.rnc import acquisition as rnc_acquisition
from agrobr.snapshots import SnapshotManifest
from agrobr.utils import warnings as source_warnings

RECONCILIACAO_CONAB_GOLDEN = Path(__file__).parent / "golden_data/reconciliacao_conab_20260918"

RETRY_SLEEP = "agrobr.http.retry.asyncio.sleep"
SERIE_HISTORICA_GOLDEN = Path(__file__).parent / "golden_data/conab/serie_historica_20260917"
SOCIOBIO_GOLDEN = Path(__file__).parent / "golden_data/conab/sociobio_20260916"
PROGRESSO_HISTORICAL_GOLDEN = (
    Path(__file__).parent / "golden_data/conab_progresso/historico_20250927"
)


@contextmanager
def collect_failures() -> Iterator[Callable[[object], AbstractContextManager[None]]]:
    failures: list[str] = []
    exceptions: list[BaseException] = []

    @contextmanager
    def check(case: object) -> Iterator[None]:
        try:
            yield
        except (Exception, pytest.fail.Exception) as error:
            failures.append(f"{case}: {type(error).__name__}: {error}")
            exceptions.append(error)

    yield check
    if failures:
        failure = pytest.fail.Exception("\n".join(failures), pytrace=False)
        failure._agrobr_case_failures = tuple(exceptions)
        raise failure


@contextmanager
def levanta_exatamente(
    esperada: type[BaseException] | tuple[type[BaseException], ...], match: str | None = None
) -> Iterator[SimpleNamespace]:
    capturada = SimpleNamespace(value=None)
    try:
        yield capturada
    except Exception as erro:
        capturada.value = erro
        assert isinstance(erro, esperada), (
            f"esperado {esperada}; saiu {type(erro).__name__}: {erro}"
        )
        if match is not None:
            assert re.search(match, str(erro)), f"{match!r} não casa com {erro}"
    else:
        raise AssertionError(f"{esperada} não foi levantado")


@contextmanager
def sem_excecao() -> Iterator[None]:
    try:
        yield
    except Exception as erro:
        raise AssertionError(f"caminho válido levantou {type(erro).__name__}: {erro}") from erro


@contextmanager
def capturar_logs() -> Iterator[list[structlog.typing.EventDict]]:
    """``structlog.testing.capture_logs`` dos loggers do agrobr, que não usam os processadores globais."""
    captura = structlog.testing.LogCapture()
    originais = list(_log.PROCESSADORES)
    _log.PROCESSADORES[:] = [captura]
    try:
        yield captura.entries
    finally:
        _log.PROCESSADORES[:] = originais


def conferir_corpo(meta: MetaInfo, corpo: bytes) -> None:
    esperado = hashlib.sha256(corpo).hexdigest()
    assert meta.raw_content_hash == esperado
    assert meta.raw_content_size == len(corpo)
    assert meta.fetch_timestamp is not None
    assert meta.fetch_timestamp.tzinfo is UTC


def isolate_dataset_info(monkeypatch: pytest.MonkeyPatch) -> None:
    for name in datasets.list_datasets():
        dataset = datasets.get_dataset(name)
        monkeypatch.setattr(type(dataset), "info", copy.deepcopy(dataset.info))
        if "info" in vars(dataset):
            monkeypatch.setattr(dataset, "info", copy.deepcopy(dataset.info))


@contextmanager
def isolated_dataset_case(case: object) -> Iterator[pytest.MonkeyPatch]:
    with pytest.MonkeyPatch.context() as monkeypatch:
        cache_root = os.environ.get("AGROBR_CACHE_CACHE_DIR")
        if cache_root:
            key = hashlib.sha256(str(case).encode()).hexdigest()[:16]
            monkeypatch.setenv("AGROBR_CACHE_CACHE_DIR", str(Path(cache_root) / key))
        monkeypatch.setattr(duckdb_store, "_store", None)
        isolate_dataset_info(monkeypatch)
        try:
            yield monkeypatch
        finally:
            if duckdb_store._store is not None:
                duckdb_store._store.close()
            config.reset_config()
            RateLimiter.reset()
            source_warnings.warn_once_reset()
            structlog.reset_defaults()


@contextmanager
def fixture_instance(factory: Callable[..., Any], **kwargs: Any) -> Iterator[Any]:
    build = inspect.unwrap(factory)
    if inspect.isgeneratorfunction(build):
        with contextmanager(build)(**kwargs) as value:
            yield value
    else:
        yield build(**kwargs)


def assert_balance_dtypes(frame: pd.DataFrame, *, dataset: bool) -> None:
    numeric = {
        "estoque_inicial",
        "producao",
        "importacao",
        "suprimento",
        "consumo",
        "exportacao",
        "demanda_total",
        "estoque_final",
    }
    textual = {"produto", "safra", "levantamento", "unidade"}
    if dataset:
        textual.add("fonte")
    assert set(frame.columns) == numeric | textual
    for name in numeric:
        assert str(frame[name].dtype) == "float64", (name, frame[name].dtype)
    for name in textual:
        assert pd.api.types.is_string_dtype(frame[name].dtype), (name, frame[name].dtype)
        assert all(isinstance(value, str) for value in frame[name].dropna())


def load_reconciliacao_conab_manifest() -> dict[str, Any]:
    manifest: dict[str, Any] = json.loads(
        (RECONCILIACAO_CONAB_GOLDEN / "manifest.json").read_text(encoding="utf-8")
    )
    assert manifest["format_version"] == 2
    return manifest


def _reconciliation_mask(frame: pd.DataFrame, key: dict[str, Any]) -> pd.Series:
    mask = pd.Series(True, index=frame.index)
    for field, value in key.items():
        mask &= frame[field].isna() if value is None else frame[field].eq(value)
    return mask


def assert_reconciliation_case(frame: pd.DataFrame, case: dict[str, Any]) -> None:
    assert len(frame) == len(case["expected_keys"]), case["id"]
    if frame.empty:
        assert case.get("empty_reason"), case["id"]
        return
    assert not frame.duplicated(case["key_columns"]).any(), case["id"]
    for key in case["expected_keys"]:
        assert _reconciliation_mask(frame, key).sum() == 1, key
    for item in case["samples"]:
        selected = frame.loc[_reconciliation_mask(frame, item["key"])]
        assert len(selected) == 1, item
        actual = selected.iloc[0][item["column"]]
        expected = item["value"]
        if expected is None:
            assert pd.isna(actual), item
        elif isinstance(expected, bool):
            assert not pd.isna(actual) and bool(actual) == expected, item
        elif isinstance(expected, int | float):
            assert float(actual) == pytest.approx(expected, rel=1e-12, abs=1e-9), item
        else:
            assert actual == expected, item
    for item in case["absent_rows"]:
        assert not _reconciliation_mask(frame, item["key"]).any(), item
    for field in case["null_columns"]:
        assert frame[field].isna().all(), field
    for excluded in case.get("excluded_columns", []):
        assert excluded["name"] not in frame.columns, excluded


def install_reconciliacao_conab_http(monkeypatch: Any, case: dict[str, Any]) -> list[str]:
    manifest = load_reconciliacao_conab_manifest()
    selected_files = set(case["http_files"])
    pages: dict[str, tuple[bytes, str]] = {}
    observations: list[dict[str, Any]] = []
    for item in manifest["files"]:
        if item["file"] not in selected_files:
            continue
        content = (RECONCILIACAO_CONAB_GOLDEN / item["file"]).read_bytes()
        assert hashlib.sha256(content).hexdigest() == item["sha256"]
        for url in (item["url"], *item.get("url_aliases", [])):
            pages[url] = content, item["content_type"]
        if case["id"].startswith("lspa_"):
            observations.extend(json.loads(content))
    requests: list[str] = []

    async def send(
        _client: httpx.AsyncClient, request: httpx.Request, **_kwargs: Any
    ) -> httpx.Response:
        assert request.method == "GET"
        url = str(request.url)
        requests.append(url)
        if case["id"].startswith("lspa_"):
            parts = request.url.path.strip("/").split("/")
            assert request.url.host == "apisidra.ibge.gov.br"
            assert parts[parts.index("t") + 1] == "6588"
            codes = set(parts[parts.index("c48") + 1].split(","))
            periods = set(parts[parts.index("p") + 1].split(","))
            assert codes <= set(case["components"])
            expected_level = "n3" if case["selection"].get("uf") else "n1"
            assert parts[parts.index(expected_level) + 1] == (
                "51" if expected_level == "n3" else "all"
            )
            selected = [
                row for row in observations if row["D3C"] in codes and row["D2C"] in periods
            ]
            assert selected, (url, case["id"])
            return httpx.Response(200, json=selected, request=request)
        assert url in pages, (url, sorted(pages))
        raw, content_type = pages[url]
        return httpx.Response(
            200, content=raw, headers={"Content-Type": content_type}, request=request
        )

    monkeypatch.setattr(httpx.AsyncClient, "send", send)
    return requests


def load_serie_historica_manifest() -> dict[str, Any]:
    manifest: dict[str, Any] = json.loads(
        (SERIE_HISTORICA_GOLDEN / "manifest.json").read_text(encoding="utf-8")
    )
    assert manifest["format_version"] == 1
    assert manifest["files"] and manifest["cases"]
    return manifest


def _serie_historica_row_mask(frame: pd.DataFrame, key: dict[str, str]) -> pd.Series:
    mask = pd.Series(True, index=frame.index)
    for field, value in key.items():
        mask &= frame[field].eq(value)
    return mask


def assert_serie_historica_case(frame: pd.DataFrame, case: dict[str, Any]) -> None:
    assert not frame.duplicated(["produto", "safra", "uf"]).any()
    assert case["samples"]
    for sample in case["samples"]:
        rows = frame.loc[_serie_historica_row_mask(frame, sample["key"])]
        assert len(rows) == 1, sample
        actual = rows.iloc[0][sample["column"]]
        expected = sample["value"]
        if expected is None:
            assert pd.isna(actual), sample
        elif isinstance(expected, (int, float)):
            assert actual == pytest.approx(expected, rel=1e-10, abs=1e-9), sample
        else:
            assert actual == expected, sample
    for absent in case["absent_rows"]:
        assert not _serie_historica_row_mask(frame, absent["key"]).any(), absent
    for column in case["null_columns"]:
        assert frame[column].isna().all(), column


def install_serie_historica_http(monkeypatch: Any, case: dict[str, Any]) -> list[str]:
    manifest = load_serie_historica_manifest()
    receipt = next(item for item in manifest["files"] if item["file"] == case["file"])
    raw = (SERIE_HISTORICA_GOLDEN / case["file"]).read_bytes()
    original_client = httpx.AsyncClient
    requests: list[str] = []

    def respond(request: httpx.Request) -> httpx.Response:
        assert request.method == "GET"
        assert str(request.url) == receipt["url"]
        requests.append(str(request.url))
        return httpx.Response(
            200, content=raw, headers={"Content-Type": "application/vnd.ms-excel"}, request=request
        )

    def factory(**kwargs: Any) -> httpx.AsyncClient:
        return original_client(transport=httpx.MockTransport(respond), **kwargs)

    monkeypatch.setattr(serie_historica_client.httpx, "AsyncClient", factory)
    return requests


def progresso_xlsx_cells(cells: dict[str, Any]) -> bytes:
    book = openpyxl.load_workbook(PROGRESSO_HISTORICAL_GOLDEN / "response.xlsx")
    sheet = book["Progresso de safra"]
    for coordinate, value in cells.items():
        sheet[coordinate] = value
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def mock_progresso_http(
    monkeypatch: Any,
    pages: dict[str, bytes],
    *,
    statuses: dict[str, int] | None = None,
    headers: dict[str, str] | None = None,
) -> list[str]:
    calls: list[str] = []
    factory = httpx.AsyncClient

    def respond(request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        calls.append(url)
        return httpx.Response(
            (statuses or {}).get(url, 200), content=pages[url], headers=headers or {}
        )

    namespace = SimpleNamespace(**vars(httpx))
    namespace.AsyncClient = lambda **kwargs: factory(
        **kwargs, transport=httpx.MockTransport(respond)
    )
    monkeypatch.setattr(progresso_client, "httpx", namespace)
    return calls


def mock_sociobio_http(monkeypatch: Any, pages: dict[str, bytes] | None = None) -> list[str]:
    manifest = json.loads((SOCIOBIO_GOLDEN / "manifest.json").read_text(encoding="utf-8"))
    bodies = {
        item["receipt"]["url"]: (SOCIOBIO_GOLDEN / item["file"]).read_bytes()
        for item in manifest["captures"]
    }
    bodies.update(pages or {})
    calls: list[str] = []
    factory = httpx.AsyncClient

    def respond(request: httpx.Request) -> httpx.Response:
        calls.append(str(request.url))
        body = bodies[str(request.url)]
        return httpx.Response(
            200, headers={"content-length": str(len(body))}, stream=httpx.ByteStream(body)
        )

    namespace = SimpleNamespace(**vars(httpx))
    namespace.AsyncClient = lambda **kwargs: factory(
        **kwargs, transport=httpx.MockTransport(respond)
    )
    monkeypatch.setattr(custos_acquisition, "httpx", namespace)
    custos_acquisition.clear()
    return calls


class TrackedAsyncStream(httpx.AsyncByteStream):
    def __init__(self, chunks: list[bytes]) -> None:
        self.chunks = chunks
        self.received = 0
        self.closed = False

    async def __aiter__(self) -> AsyncIterator[bytes]:
        for chunk in self.chunks:
            self.received += 1
            yield chunk

    async def aclose(self) -> None:
        self.closed = True


def make_anp_precos_resource(
    *,
    nivel: str = "municipio",
    rows: list[dict[str, Any]] | None = None,
    url: str = "https://www.gov.br/anp/precos.xlsx",
) -> anp_client.PrecosResource:
    if rows is None:
        rows = [
            {
                "DATA INICIAL": "28/01/2024",
                "DATA FINAL": "03/02/2024",
                "PREÇO MÉDIO REVENDA": 6.4,
                "NÚMERO DE POSTOS PESQUISADOS": 8,
            },
            {
                "DATA INICIAL": "04/02/2024",
                "DATA FINAL": "10/02/2024",
                "PREÇO MÉDIO REVENDA": 6.8,
                "NÚMERO DE POSTOS PESQUISADOS": 9,
            },
        ]
    frame = pd.DataFrame(rows).copy()
    frame["PRODUTO"] = "OLEO DIESEL S10"
    frame["UNIDADE DE MEDIDA"] = "R$/l"
    if nivel in {"uf", "municipio"}:
        frame["ESTADO"] = "MATO GROSSO"
    if nivel == "municipio":
        frame["MUNICÍPIO"] = "CUIABA"
    buffer = io.BytesIO()
    frame.to_excel(buffer, index=False, engine="openpyxl")
    return anp_client.PrecosResource(
        content=buffer.getvalue(),
        requested_url=url,
        url=url,
        fetched_at=datetime(2026, 9, 8, 12, 0, tzinfo=UTC),
        etag='"fixture"',
        last_modified="Tue, 08 Sep 2026 11:00:00 GMT",
    )


def comexstat_csv(
    *, fluxo: str = "exportacao", ano: int = 2024, rows: list[list[str]] | None = None
) -> bytes:
    header = [
        "CO_ANO",
        "CO_MES",
        "CO_NCM",
        "CO_UNID",
        "CO_PAIS",
        "SG_UF_NCM",
        "CO_VIA",
        "CO_URF",
        "QT_ESTAT",
        "KG_LIQUIDO",
        "VL_FOB",
    ]
    if fluxo == "importacao":
        header += ["VL_FRETE", "VL_SEGURO"]
    if rows is None:
        rows = [
            [
                str(ano),
                "01",
                "12019000",
                "10",
                "160",
                "MT",
                "04",
                "0817800",
                "1000",
                "50000000",
                "20000000",
            ],
            [
                str(ano),
                "02",
                "12019000",
                "10",
                "160",
                "MT",
                "04",
                "0817800",
                "800",
                "40000000",
                "16000000",
            ],
            [
                str(ano),
                "03",
                "12019000",
                "21",
                "586",
                "PR",
                "07",
                "0917502",
                "500",
                "30000000",
                "12000000",
            ],
        ]
        if fluxo == "importacao":
            rows = [row + [str(10 + index), str(index)] for index, row in enumerate(rows)]
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, delimiter=";", lineterminator="\n")
    writer.writerow(header)
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8")


def install_comexstat_http(
    monkeypatch: Any, body: bytes, *, status: int = 200
) -> list[httpx.Request]:
    calls: list[httpx.Request] = []

    def respond(request: httpx.Request) -> httpx.Response:
        calls.append(request)
        return httpx.Response(
            status,
            headers={
                "Content-Length": str(len(body)),
                "Content-Type": "text/csv",
                "ETag": "synthetic",
            },
            stream=httpx.ByteStream(body),
            request=request,
        )

    namespace = SimpleNamespace(**vars(httpx))
    namespace.AsyncClient = lambda **kwargs: httpx.AsyncClient(
        transport=httpx.MockTransport(respond), **kwargs
    )
    monkeypatch.setattr(comexstat_client, "httpx", namespace)
    return calls


_SOCKET_TYPE = socket.socket
_SOCKET_CONNECT = socket.socket.connect
INCRA_EXPECTED_ALIASES = {
    "codigo": "cd_quilomb",
    "nome": "no_comunidade",
    "municipio": "no_municipio",
    "uf": "sg_uf",
    "area_ha": "nu_area_ha",
    "familias": "nu_familia",
    "fase": "ds_fase",
    "titulado": "st_titulad",
    "data_publicacao": "dt_publica",
    "data_titulo": "dt_titulo",
    "feature_id": None,
    "regional": "co_sr",
    "processo": "nu_processo",
    "data_publicacao_2": "dt_public1",
    "responsavel": "no_responsavel",
    "esfera": "no_esfera",
    "data_cadastro": "dt_cadastro",
    "codigo_sipra": "cd_sipra",
    "descricao": "ds_descricao",
    "data_decreto": "dt_decreto",
    "tipo_levantamento": "tp_levanta",
    "escala": "nr_escalao",
}


def desmatamento_features(sample: str = "prodes_amazonia") -> list[dict[str, Any]]:
    path = Path(__file__).parent / "golden_data/desmatamento/selecao_20260907"
    return json.loads((path / f"{sample}.json").read_bytes())["features"]


def desmatamento_frame(tipo: str = "prodes", **overrides: Any) -> pd.DataFrame:
    contract = (
        desmatamento_contracts.PRODES_FEICOES_V2
        if tipo == "prodes"
        else desmatamento_contracts.DETER_FEICOES_V2
    )
    row = {
        **dict.fromkeys(contract.list_columns()),
        "ano": 2023.0,
        "data": pd.Timestamp("2024-06-15"),
        "uf": "MT" if tipo == "prodes" else "PA",
        "classe": "DESFLORESTAMENTO" if tipo == "prodes" else "DESMATAMENTO_CR",
        "area_km2": 1234.56 if tipo == "prodes" else 5.42,
        "satelite": "LANDSAT",
        "sensor": "OLI",
        "bioma": "Cerrado" if tipo == "prodes" else "Amazônia",
        "feature_id": "synthetic.1",
        "municipio": "Altamira",
        "municipio_id": "1500602",
        **overrides,
    }
    return pd.DataFrame(
        {
            name: pd.Series([row[name]], dtype=dtype)
            for name, dtype in contract.empty_frame().dtypes.items()
        }
    )


def desmatamento_meta(rows: int, **coverage: Any) -> MetaInfo:
    return MetaInfo(
        source="desmatamento",
        source_url="https://example.test/wfs",
        source_method="synthetic",
        fetched_at=datetime(2026, 9, 7, tzinfo=UTC),
        schema_version="2.0",
        contract_version="2.0",
        attempted_sources=["terrabrasilis_prodes"],
        selected_source="terrabrasilis_prodes",
        source_details={
            "coverage": {
                "status": "reconciled",
                "expected_rows": rows,
                "accepted_rows": rows,
                "returned_rows": rows,
                "count_reconciled": True,
                "truncated": False,
                "all_logical_requests_succeeded": True,
                "transactional": False,
                "semantic_progress_proven": False,
                **coverage,
            }
        },
    )


def install_desmatamento_wfs(
    monkeypatch: Any,
    features: list[dict[str, Any]],
    transform: Callable[[httpx.Request, dict[str, Any]], None] | None = None,
) -> list[tuple[httpx.Request, bytes]]:
    calls: list[tuple[httpx.Request, bytes]] = []

    def respond(request: httpx.Request) -> httpx.Response:
        params = request.url.params
        offset = int(params.get("startIndex", "0"))
        count = int(params.get("count", "0"))
        selected = [] if params.get("resultType") == "hits" else features[offset : offset + count]
        payload = {
            "type": "FeatureCollection",
            "features": copy.deepcopy(selected),
            "numberMatched": len(features),
            "numberReturned": len(selected),
            "totalFeatures": len(features),
            "crs": {"type": "name", "properties": {"name": "urn:ogc:def:crs:EPSG::4326"}},
        }
        if transform is not None:
            transform(request, payload)
        body = json.dumps(payload, ensure_ascii=False, allow_nan=False).encode("utf-8")
        calls.append((request, body))
        return httpx.Response(200, content=body, headers={"etag": "synthetic"}, request=request)

    namespace = SimpleNamespace(**vars(httpx))
    namespace.AsyncClient = lambda **kwargs: httpx.AsyncClient(
        transport=httpx.MockTransport(respond), **kwargs
    )
    monkeypatch.setattr(wfs_transport, "httpx", namespace)
    return calls


def zarc_frame(**overrides: Any) -> pd.DataFrame:
    row = {
        "cultura": "soja",
        "safra": "2026/2027",
        "geocodigo": "5208707",
        "uf": "GO",
        "municipio": "Goiânia",
        "solo_codigo": 11,
        "ciclo_codigo": 20,
        "clima": "Não se aplica",
        "manejo": "Sequeiro",
        "portaria": "Portaria sintética para teste",
        **dict.fromkeys(constants.ZARC_RISK_COLUMNS, 20),
        "cultura_original": "Soja",
        "safra_inicio": "2026",
        "safra_fim": "2027",
        "cultura_codigo": "12016500000011",
        "clima_codigo": "0",
        "manejo_codigo": "1",
        "produtividade_texto": "",
        "nm_codigo": "",
        "municipio_sicor_codigo": "002345",
        "mesorregiao_codigo": "03",
        "microrregiao_codigo": "010",
        "registro_origem": 105,
        **overrides,
    }
    return pd.DataFrame(
        {
            name: pd.Series(
                [row[name]], dtype="Int64" if name in constants.ZARC_INTEGER_COLUMNS else "object"
            )
            for name in constants.ZARC_OUTPUT_COLUMNS
        }
    )


def zarc_csv(
    rows: list[dict[str, str]] | None = None, *, headers: list[str] | None = None
) -> bytes:
    columns = list(constants.ZARC_CSV_COLUMNS) if headers is None else headers
    base_row = {
        **dict.fromkeys(constants.ZARC_CSV_COLUMNS, ""),
        "Nome_cultura": "Soja",
        "SafraIni": "2026",
        "SafraFin": "2027",
        "Cod_Cultura": "123",
        "Cod_Ciclo": "20",
        "Cod_Solo": "1",
        "geocodigo": "5107925",
        "UF": "MT",
        "municipio": "Sorriso",
        "Cod_Clima": "0",
        "Nome_Clima": "Não se aplica",
        "Cod_Outros_Manejos": "1",
        "Nome_Outros_Manejos": "Sequeiro",
        "Cod_Munic": "001234",
        "Cod_Meso": "01",
        "Cod_Micro": "002",
        "Portaria": "Portaria sintética para teste",
        **dict.fromkeys(constants.ZARC_RISK_COLUMNS, "0"),
    }
    stream = io.StringIO(newline="")
    writer = csv.writer(stream, delimiter=";", lineterminator="\n")
    writer.writerow(columns)
    for overrides in [{}] if rows is None else rows:
        row = {**base_row, **overrides}
        writer.writerow([row.get(name, "") for name in columns])
    return stream.getvalue().encode("utf-8-sig")


def rnc_csv_acquisition(
    content: bytes,
    kind: rnc_acquisition.Family,
    *,
    reported_total: int | None = None,
    received_at: datetime | None = None,
) -> rnc_acquisition.CSVAcquisition:
    url = constants.RNC_PUBLIC_URLS[kind]
    request = httpx.Request("POST", url)
    search = rnc_acquisition.SearchResource.model_validate(
        {
            **rnc_acquisition.describe_response(
                httpx.Response(200, content=b"Synthetic public search", request=request)
            ).model_dump(),
            "reported_total": reported_total,
        }
    )
    resource = rnc_acquisition.describe_response(
        httpx.Response(200, content=content, request=request)
    )
    if received_at is not None:
        search = search.model_copy(update={"received_at": received_at})
        resource = resource.model_copy(update={"received_at": received_at})
    return rnc_acquisition.CSVAcquisition(
        kind=kind, resource=resource, search=search, content=content
    )


def rnc_saida_do_publicado(coluna: str, valor: Any) -> Any:
    """Célula publicada pelo CultivarWeb como o RNC/SNPC a entrega: nas colunas de espécie, brancos internos
    repetidos viram 1 espaço; nas outras, o texto fica como publicado."""
    if coluna in ("nome_cientifico", "nome_comum") and isinstance(valor, str):
        return " ".join(valor.split())
    return valor


def mapbiomas_workbook_bundle(content: bytes, url: str) -> mapbiomas_resources.WorkbookAcquisition:
    response = httpx.Response(200, content=content, request=httpx.Request("GET", url))
    return mapbiomas_resources.WorkbookAcquisition(
        content=content,
        source_url=url,
        resource=mapbiomas_resources.HTTPResource.from_response(response, url),
    )


@asynccontextmanager
async def binary_stream(content: bytes) -> AsyncIterator[io.BytesIO]:
    with io.BytesIO(content) as stream:
        yield stream


def local_socketpair(
    family: int | None = None, type: int = socket.SOCK_STREAM, proto: int = 0
) -> tuple[socket.socket, socket.socket]:
    if family is None:
        family = socket.AF_INET
    hosts = {socket.AF_INET: "127.0.0.1", socket.AF_INET6: "::1"}
    if family not in hosts or type != socket.SOCK_STREAM or proto != 0:
        raise ValueError("Only local TCP socket pairs are supported")
    with _SOCKET_TYPE(family, type, proto) as listener, ExitStack() as cleanup:
        listener.settimeout(5.0)
        listener.bind((hosts[family], 0))
        listener.listen(1)
        client = cleanup.enter_context(_SOCKET_TYPE(family, type, proto))
        client.setblocking(False)
        with suppress(BlockingIOError, InterruptedError):
            _SOCKET_CONNECT(client, listener.getsockname())
        client.setblocking(True)
        descriptor, _ = listener._accept()
        try:
            server = _SOCKET_TYPE(family, type, proto, fileno=descriptor)
        except BaseException:
            socket.close(descriptor)
            raise
        cleanup.enter_context(server)
        server.setblocking(True)
        if (
            server.getsockname() != client.getpeername()
            or client.getsockname() != server.getpeername()
        ):
            raise ConnectionError("Unexpected socket-pair endpoint")
        cleanup.pop_all()
        return server, client


def seed_cache_schema(conn: duckdb.DuckDBPyConnection, version: int) -> None:
    conn.execute(duckdb_store.SCHEMA_CACHE)
    conn.execute(duckdb_store.SCHEMA_HISTORY)
    conn.execute(duckdb_store.SCHEMA_INDICADORES)
    conn.execute(
        "CREATE TABLE schema_version (version INTEGER PRIMARY KEY, "
        "applied_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)"
    )
    conn.execute("INSERT INTO schema_version (version) VALUES (?)", [version])


def insert_cache_indicator(conn: duckdb.DuckDBPyConnection, **overrides: Any) -> None:
    record = {
        "produto": "acucar_refinado",
        "praca": "São Paulo/SP",
        "data": date(2000, 8, 1),
        "valor": "2.9367",
        "unidade": "BRL/sc50kg",
        "fonte": "noticias_agricolas",
        "metodologia": "original",
        "variacao_percentual": "-0.1250",
        "collected_at": datetime(2026, 9, 4, 18, 25, 1, 123456),
        "parser_version": 1,
        **overrides,
    }
    columns = ", ".join(record)
    placeholders = ", ".join("?" for _ in record)
    conn.execute(
        f"INSERT INTO indicadores ({columns}) VALUES ({placeholders})", list(record.values())
    )


async def make_snapshot_source(path: Path, manifest: SnapshotManifest) -> None:
    (path / "sample.parquet").touch()
    manifest.files[f"{path.name}/sample.parquet"] = {"rows": 1, "columns": ["valor"]}


def make_sleep_tracker() -> tuple[list[float], Callable[[float], Coroutine[Any, Any, None]]]:
    sleep_calls: list[float] = []

    async def track_sleep(delay: float) -> None:
        sleep_calls.append(delay)

    return sleep_calls, track_sleep


def make_mock_response(
    status_code: int = 200,
    *,
    text: str | None = None,
    content: bytes | None = None,
    json_data: dict | list | None = None,
    headers: dict | None = None,
    url: str = "https://test.agrobr.dev/mock",
    charset_encoding: str | None = None,
) -> MagicMock:
    resp = MagicMock(spec=httpx.Response)
    resp.status_code = status_code

    if text is not None:
        resp.text = text
    if content is not None:
        resp.content = content
    if json_data is not None:
        resp.json.return_value = json_data

    if headers is not None:
        resp.headers = headers
    elif charset_encoding is not None:
        resp.headers = {"content-type": f"text/html; charset={charset_encoding}"}
    else:
        resp.headers = {}

    resp.url = url
    resp.history = []
    resp.request = httpx.Request("GET", url)

    if charset_encoding is not None:
        resp.charset_encoding = charset_encoding

    resp.raise_for_status = MagicMock()
    if status_code >= 400:
        resp.raise_for_status.side_effect = httpx.HTTPStatusError(
            f"HTTP {status_code}", request=MagicMock(), response=resp
        )

    return resp


def make_mock_async_client() -> AsyncMock:
    client = AsyncMock()
    client.__aenter__ = AsyncMock(return_value=client)
    client.__aexit__ = AsyncMock(return_value=False)
    return client


def make_alert_settings(
    *,
    enabled: bool = True,
    slack_webhook: str | None = None,
    discord_webhook: str | None = None,
    sendgrid_api_key: str | None = None,
    email_from: str = "alerts@agrobr.dev",
    email_to: list[str] | None = None,
    discord_embed_char_limit: int = 3900,
) -> MagicMock:
    settings = MagicMock()
    settings.enabled = enabled
    settings.slack_webhook = slack_webhook
    settings.discord_webhook = discord_webhook
    settings.sendgrid_api_key = sendgrid_api_key
    settings.email_from = email_from
    settings.email_to = email_to if email_to is not None else []
    settings.discord_embed_char_limit = discord_embed_char_limit
    return settings


def embrapa_solos_features(
    product: str = "perfis", *, include_geometry: bool = False
) -> list[dict[str, Any]]:
    if product not in {"perfis", "mapa"}:
        raise ValueError("Unknown Embrapa test product")
    suffix = "geo" if include_geometry else "reference"
    path = (
        Path(__file__).parent
        / "golden_data/embrapa_solos/official_20260907"
        / f"{product}_{suffix}.json"
    )
    return json.loads(path.read_bytes())["features"]


def install_embrapa_solos_wfs(
    monkeypatch: Any,
    features: list[dict[str, Any]],
    transform: Callable[[httpx.Request, dict[str, Any]], None] | None = None,
) -> list[tuple[httpx.Request, bytes]]:
    calls: list[tuple[httpx.Request, bytes]] = []

    def respond(request: httpx.Request) -> httpx.Response:
        params = request.url.params
        offset = int(params.get("startIndex", "0"))
        count = int(params.get("count", "0"))
        selected = [] if params.get("resultType") == "hits" else features[offset : offset + count]
        payload = {
            "type": "FeatureCollection",
            "features": copy.deepcopy(selected),
            "numberMatched": len(features),
            "numberReturned": len(selected),
            "totalFeatures": len(features),
            "crs": {"type": "name", "properties": {"name": "urn:ogc:def:crs:EPSG::4326"}},
        }
        if transform is not None:
            transform(request, payload)
        body = json.dumps(payload, ensure_ascii=False, allow_nan=False).encode("utf-8")
        calls.append((request, body))
        return httpx.Response(200, content=body, headers={"etag": "synthetic"}, request=request)

    namespace = SimpleNamespace(**vars(httpx))
    namespace.AsyncClient = lambda **kwargs: httpx.AsyncClient(
        transport=httpx.MockTransport(respond), **kwargs
    )
    monkeypatch.setattr(wfs_transport, "httpx", namespace)
    return calls


def funai_features(*, include_geometry: bool = False) -> list[dict[str, Any]]:
    name = "geo.json" if include_geometry else "reference.json"
    path = Path(__file__).parent / "golden_data/funai/official_20260907" / name
    return json.loads(path.read_bytes())["features"]


def install_funai_wfs(
    monkeypatch: Any,
    features: list[dict[str, Any]],
    transform: Callable[[httpx.Request, dict[str, Any]], httpx.Response | None] | None = None,
    *,
    volatile_ids: bool = True,
) -> list[tuple[httpx.Request, bytes]]:
    calls: list[tuple[httpx.Request, bytes]] = []

    def respond(request: httpx.Request) -> httpx.Response:
        index = len(calls)
        calls.append((request, b""))
        params = request.url.params
        offset = int(params.get("startIndex", "0"))
        count = int(params.get("count", "0"))
        selected = (
            []
            if params.get("resultType") == "hits"
            else copy.deepcopy(features[offset : offset + count])
        )
        for ordinal, feature in enumerate(selected):
            if volatile_ids:
                feature["id"] = f"synthetic.funai.request{index}.row{ordinal}"
            if "the_geom" not in params.get("propertyName", "").split(","):
                feature["geometry"] = None
        payload = {
            "type": "FeatureCollection",
            "features": selected,
            "numberMatched": len(features),
            "numberReturned": len(selected),
            "totalFeatures": len(features),
            "timeStamp": "2026-09-07T00:00:00.000Z",
            "crs": {"type": "name", "properties": {"name": "urn:ogc:def:crs:EPSG::4326"}}
            if params.get("srsName") == "EPSG:4326"
            else None,
        }
        if count and len(selected) == count:
            payload["links"] = [
                {
                    "rel": "next",
                    "href": str(request.url.copy_set_param("startIndex", str(offset + count))),
                }
            ]
        override = transform(request, payload) if transform is not None else None
        body = json.dumps(payload, ensure_ascii=False, allow_nan=False).encode("utf-8")
        if override is not None:
            calls[index] = (request, override.content)
            return override
        calls[index] = (request, body)
        return httpx.Response(
            200,
            content=body,
            headers={"etag": '"synthetic.funai"', "content-type": "application/json"},
            request=request,
        )

    namespace = SimpleNamespace(**vars(httpx))
    namespace.AsyncClient = lambda **kwargs: httpx.AsyncClient(
        transport=httpx.MockTransport(respond), **kwargs
    )
    monkeypatch.setattr(wfs_transport, "httpx", namespace)
    return calls


def incra_features(*, include_geometry: bool = False) -> list[dict[str, Any]]:
    name = "geo.json" if include_geometry else "reference.json"
    path = Path(__file__).parent / "golden_data/incra/wfs_v2_20260908" / name
    return json.loads(path.read_bytes())["features"]


def _incra_expected_column(name: str, values: list[Any]) -> pd.Series:
    if name == "data_cadastro":
        return pd.Series(
            [pd.Timestamp(value).tz_convert("UTC") for value in values],
            dtype="datetime64[ns, UTC]",
        )
    if name.startswith("data_"):
        return pd.Series(
            [
                pd.Timestamp(value) if value and 1900 <= int(value[:4]) <= 2099 else None
                for value in values
            ],
            dtype="datetime64[ns]",
        )
    dtype = (
        "Int64"
        if name in {"codigo", "familias"}
        else "float64"
        if name == "area_ha"
        else pd.Series([""]).dtype
    )
    return pd.Series(values, dtype=dtype)


def incra_expected_frame(*, include_geometry: bool = False) -> pd.DataFrame:
    features = incra_features(include_geometry=include_geometry)
    return pd.DataFrame(
        {
            name: _incra_expected_column(
                name,
                [
                    feature["id"] if raw is None else feature["properties"][raw]
                    for feature in features
                ],
            )
            for name, raw in INCRA_EXPECTED_ALIASES.items()
        }
    )


def install_incra_wfs(
    monkeypatch: Any,
    features: list[dict[str, Any]],
    transform: Callable[[httpx.Request, dict[str, Any]], httpx.Response | None] | None = None,
    *,
    volatile_ids: bool = False,
) -> list[tuple[httpx.Request, bytes]]:
    calls: list[tuple[httpx.Request, bytes]] = []

    def respond(request: httpx.Request) -> httpx.Response:
        index = len(calls)
        calls.append((request, b""))
        params = request.url.params
        offset = int(params.get("startIndex", "0"))
        count = int(params.get("count", "0"))
        selected = copy.deepcopy(features[offset : offset + count])
        for ordinal, feature in enumerate(selected):
            if volatile_ids:
                feature["id"] = f"synthetic.incra.request{index}.row{ordinal}"
            if "geom" not in params.get("propertyName", "").split(","):
                feature["geometry"] = None
        payload = {
            "type": "FeatureCollection",
            "features": selected,
            "numberMatched": len(features),
            "numberReturned": len(selected),
            "totalFeatures": len(features),
            "timeStamp": "2026-09-08T00:00:00.000Z",
            "crs": {"type": "name", "properties": {"name": "urn:ogc:def:crs:EPSG::4326"}}
            if selected and params.get("srsName") == "EPSG:4326"
            else None,
        }
        if count and len(selected) == count:
            payload["links"] = [
                {
                    "rel": "next",
                    "href": str(request.url.copy_set_param("startIndex", str(offset + count))),
                }
            ]
        override = transform(request, payload) if transform is not None else None
        body = json.dumps(payload, ensure_ascii=False, allow_nan=False).encode("utf-8")
        if override is not None:
            calls[index] = (request, override.content)
            return override
        calls[index] = (request, body)
        return httpx.Response(
            200,
            content=body,
            headers={"etag": '"synthetic.incra"', "content-type": "application/json"},
            request=request,
        )

    namespace = SimpleNamespace(**vars(httpx))
    namespace.AsyncClient = lambda **kwargs: httpx.AsyncClient(
        transport=httpx.MockTransport(respond), **kwargs
    )
    monkeypatch.setattr(wfs_transport, "httpx", namespace)
    return calls


def install_anttpedagio_http(
    monkeypatch: Any,
    handler: Callable[[httpx.Request], httpx.Response],
) -> list[httpx.Request]:
    calls: list[httpx.Request] = []

    def respond(request: httpx.Request) -> httpx.Response:
        calls.append(request)
        return handler(request)

    def factory(**kwargs: Any) -> httpx.AsyncClient:
        return httpx.AsyncClient(transport=httpx.MockTransport(respond), **kwargs)

    namespace = SimpleNamespace(**vars(httpx))
    namespace.AsyncClient = factory
    monkeypatch.setattr(antt_client, "httpx", namespace)
    return calls


def install_anttpedagio_source(
    monkeypatch: Any,
    files: dict[str, bytes],
    *,
    plazas: bytes | None = None,
    faults: dict[str, Exception | int] | None = None,
) -> tuple[list[httpx.Request], dict[str, bytes]]:
    bodies = {f"https://dados.antt.gov.br/{name}": content for name, content in files.items()}
    groups = {"volume-trafego-praca-pedagio": list(bodies)}
    if plazas is not None:
        bodies["https://dados.antt.gov.br/pracas.csv"] = plazas
        groups["praca-de-pedagio"] = ["https://dados.antt.gov.br/pracas.csv"]
    for slug, urls in groups.items():
        envelope = {
            "success": True,
            "result": {
                "id": f"test-{slug}",
                "name": slug,
                "license_id": "cc-by",
                "metadata_modified": "2026-08-28T11:35:32.019937",
                "resources": [
                    {
                        "id": f"test-resource-{index}",
                        "name": url.rsplit("/", 1)[-1],
                        "url": url,
                        "format": "CSV",
                        "size": len(bodies[url]),
                        "last_modified": "2026-08-28T11:35:14.750047",
                    }
                    for index, url in enumerate(urls)
                ],
            },
        }
        url = f"https://dados.antt.gov.br/api/3/action/package_show?id={slug}"
        bodies[url] = json.dumps(envelope, ensure_ascii=False).encode()

    def respond(request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        fault = (faults or {}).get(url)
        if isinstance(fault, Exception):
            raise fault
        if isinstance(fault, int):
            return httpx.Response(fault, stream=httpx.ByteStream(b"failure"), request=request)
        body = bodies[url]
        return httpx.Response(
            200,
            stream=httpx.ByteStream(body),
            headers={"content-length": str(len(body)), "etag": '"synthetic-antt"'},
            request=request,
        )

    return install_anttpedagio_http(monkeypatch, respond), bodies


def anttpedagio_fluxo_frame() -> pd.DataFrame:
    row = {
        "data": pd.Timestamp("2025-01-15"),
        "concessionaria": "Concessionária",
        "praca": " Praça 1 ",
        "sentido": "Crescente",
        "n_eixos": 3,
        "tipo_veiculo": "Comercial",
        "volume": 9007199254740993,
        "rodovia": None,
        "uf": None,
        "municipio": None,
        "categoria_eixo": "Veículo Comercial 3 eixos",
        "tipo_cobranca": "N/I",
        "frequencia": "diaria",
    }
    return pd.DataFrame(
        {
            name: pd.Series(
                [value],
                dtype="datetime64[ns]"
                if name == "data"
                else "Int64"
                if name in ("volume", "n_eixos")
                else "string[python]",
            )
            for name, value in row.items()
        }
    )


async def assert_live_product_result(
    dataset_name: str,
    product: str,
    args: tuple[Any, ...],
    kwargs: dict[str, Any],
    pause: float,
    record_property: Callable[[str, Any], None],
) -> None:
    await asyncio.sleep(pause)
    started = time.perf_counter()
    try:
        frame, meta = await datasets.get_dataset(dataset_name).fetch(
            *args, **kwargs, return_meta=True
        )
    finally:
        record_property("seconds", round(time.perf_counter() - started, 3))
    record_property("records_count", len(frame))
    record_property("selected_source", meta.selected_source)
    record_property("attempted_sources", ",".join(meta.attempted_sources))
    assert len(frame) > 0, f"{dataset_name}/{product}: consulta de referência sem registros"
    record_property("live_matrix_status", "validated")


def sicar_feature_collection(
    features: list[dict[str, Any]], *, number_matched: int | str | None = None
) -> bytes:
    return json.dumps(
        {
            "type": "FeatureCollection",
            "features": features,
            "numberReturned": len(features),
            "numberMatched": len(features) if number_matched is None else number_matched,
        }
    ).encode()


def install_reconciliacao_custos_http(monkeypatch: Any) -> list[str]:
    golden = Path(__file__).parent / "golden_data/reconciliacao_custos_conab_20260918"
    receipts = json.loads((golden / "receipts.json").read_text(encoding="utf-8"))
    pages = {}
    for receipt in receipts:
        raw = (golden / receipt["file"]).read_bytes()
        assert hashlib.sha256(raw).hexdigest() == receipt["sha256"]
        assert receipt["status"] == 200
        pages[receipt["url"]] = raw, receipt["headers"]
    calls = []

    async def send(
        _client: httpx.AsyncClient, request: httpx.Request, **_kwargs: Any
    ) -> httpx.Response:
        assert request.method == "GET"
        url = str(request.url)
        calls.append(url)
        assert url in pages, url
        raw, headers = pages[url]
        return httpx.Response(200, headers=headers, stream=httpx.ByteStream(raw), request=request)

    monkeypatch.setattr(httpx.AsyncClient, "send", send)
    return calls


def load_custo_sheet(
    file: str, product: str, name: str
) -> tuple[custos_workbook.Aba, custos_models.ContextoCusto]:
    golden = Path(__file__).parent / "golden_data/reconciliacao_custos_conab_20260918"
    receipts = json.loads((golden / "receipts.json").read_text(encoding="utf-8"))
    receipt = next(item for item in receipts if item["file"] == file)
    raw = (golden / file).read_bytes()
    assert hashlib.sha256(raw).hexdigest() == receipt["sha256"]
    identifier = receipt["url"].split("/@@download")[0].rsplit("/", 1)[-1]
    resource = custos_models.RecursoCusto(
        cultura=product,
        planilha=identifier,
        titulo=identifier,
        pagina_url=receipt["url"],
    )
    workbook = custos_workbook.Workbook(raw)
    try:
        sheet = workbook.read(name)
        context = custos_context.context(sheet, resource, workbook.names.index(name))
        return sheet, context
    finally:
        workbook.close()


def assert_custo_meta(
    frame: pd.DataFrame,
    meta: MetaInfo,
    case: dict[str, Any],
    manifest: dict[str, Any],
    *,
    source: bool = False,
) -> None:
    sociobio = case["dataset"] == "custo_sociobiodiversidade"
    selected_source = "conab_sociobio" if sociobio else "conab_custo" if source else "conab"
    receipt = next(item for item in manifest["files"] if item["file"] == case["file"])
    assert meta.attempted_sources == [selected_source]
    assert meta.selected_source == selected_source
    assert meta.parser_version == (2 if sociobio else 5)
    assert meta.schema_version == ("1.0" if sociobio else "3.0")
    assert meta.fetch_timestamp is not None and meta.fetch_timestamp.tzinfo is not None
    assert meta.records_count == len(frame)
    details = meta.source_details
    assert details["resource"]["planilha"] == case["selection"]["planilha"]
    selection = details["selection"][0] if sociobio else details["selection"]
    assert selection["aba"] == case["selection"]["aba"]
    resources = details["manifest"]["acquisition"]["resources"]
    workbook = next(item for item in resources if item["role"] == "workbook")
    assert workbook["sha256"] == receipt["sha256"]
    assert workbook["url"] == receipt["url"]
    assert workbook["bytes"] == receipt["bytes"]
    assert workbook["eof"] is True and workbook["closed"] is True
    if sociobio:
        assert meta.raw_content_hash == receipt["sha256"]
    else:
        assert details["workbook_sha256"] == receipt["sha256"]
    for field in frame.columns:
        if field.startswith(("valor", "participacao")):
            assert str(frame[field].dtype) == "float64", field


def install_reconciliacao_censos_http(
    monkeypatch: Any, case: dict[str, Any], manifest: dict[str, Any]
) -> list[str]:
    golden = Path(__file__).parent / "golden_data/reconciliacao_censos_producao_ibge_conab_20260918"
    file = next(item for item in manifest["files"] if item["file"] == case["file"])
    raw = (golden / file["file"]).read_bytes()
    assert hashlib.sha256(raw).hexdigest() == file["sha256"]
    if file["file"].endswith(".csv"):
        rows = list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig"))))
    else:
        rows = json.loads(raw)
        if rows and rows[0].get("V") == "Valor":
            rows = rows[1:]
    query = case["transport"]
    selected = [
        row
        for row in rows
        if row["D4C"] == query["product_code"]
        and row["D2C"] == query["period"]
        and row["D3C"] in query["variables"]
    ]
    assert selected
    payload = json.dumps(selected, ensure_ascii=False).encode("utf-8")
    requests: list[str] = []

    async def send(
        _client: httpx.AsyncClient, request: httpx.Request, **_kwargs: Any
    ) -> httpx.Response:
        assert request.method == "GET"
        assert request.url.host == "apisidra.ibge.gov.br"
        parts = request.url.path.strip("/").split("/")
        assert parts[0] == "values"
        params = dict(zip(parts[1::2], parts[2::2], strict=True))
        assert params["t"] == query["table"]
        assert params["n" + query["level"]] == query["territory"]
        assert params["p"] == query["period"]
        assert set(params["v"].split(",")) == set(query["variables"])
        assert params["c" + query["classification"]] == query["product_code"]
        assert params["h"] == "n"
        requests.append(str(request.url))
        return httpx.Response(
            200, content=payload, headers={"Content-Type": "application/json"}, request=request
        )

    monkeypatch.setattr(httpx.AsyncClient, "send", send)
    return requests


def install_reconciliacao_censos_conab_http(
    monkeypatch: Any, conab_case_id: str, extra_file: dict[str, Any] | None = None
) -> list[str]:
    manifest = load_reconciliacao_conab_manifest()
    case = next(item for item in manifest["cases"] if item["id"] == conab_case_id)
    requests = install_reconciliacao_conab_http(monkeypatch, case)
    original = httpx.AsyncClient.send
    extra_raw = b""
    if extra_file is not None:
        golden = (
            Path(__file__).parent / "golden_data/reconciliacao_censos_producao_ibge_conab_20260918"
        )
        extra_raw = (golden / extra_file["file"]).read_bytes()
        assert hashlib.sha256(extra_raw).hexdigest() == extra_file["sha256"]

    async def send(
        client: httpx.AsyncClient, request: httpx.Request, **kwargs: Any
    ) -> httpx.Response:
        if request.url.host == "apisidra.ibge.gov.br":
            requests.append(str(request.url))
            return httpx.Response(
                403,
                text="Just a moment",
                headers={"cf-mitigated": "challenge"},
                request=request,
            )
        if extra_file is not None and str(request.url) == extra_file["url"]:
            assert request.method == "GET"
            requests.append(str(request.url))
            return httpx.Response(
                200,
                content=extra_raw,
                headers={"Content-Type": extra_file["content_type"]},
                request=request,
            )
        return await original(client, request, **kwargs)

    monkeypatch.setattr(httpx.AsyncClient, "send", send)
    return requests


def install_reconciliacao_censos_quarter_http(
    monkeypatch: Any,
    case: dict[str, Any],
    manifest: dict[str, Any],
    override: list[dict[str, str]] | None = None,
) -> list[str]:
    golden = Path(__file__).parent / "golden_data/reconciliacao_censos_producao_ibge_conab_20260918"
    receipt = next(file for file in manifest["files"] if file["file"] == case["file"])
    raw = (golden / case["file"]).read_bytes()
    assert hashlib.sha256(raw).hexdigest() == receipt["sha256"]
    rows = (
        list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig")))) if override is None else override
    )
    requests: list[str] = []

    async def send(
        _client: httpx.AsyncClient, request: httpx.Request, **_kwargs: Any
    ) -> httpx.Response:
        assert request.method == "GET"
        assert request.url.host == "apisidra.ibge.gov.br"
        parts = request.url.path.strip("/").split("/")
        assert parts[0] == "values"
        params = dict(zip(parts[1::2], parts[2::2], strict=True))
        assert params.keys() == case["transport"].keys()
        for key, value in case["transport"].items():
            assert set(params[key].split(",")) == set(value.split(","))
        requests.append(str(request.url))
        return httpx.Response(200, json=rows, request=request)

    monkeypatch.setattr(httpx.AsyncClient, "send", send)
    return requests


def install_reconciliacao_censos_agro_http(
    monkeypatch: Any,
    responses: dict[str, list[dict[str, str]]],
    variables: dict[str, str],
) -> list[str]:
    requests: list[str] = []

    async def send(
        _client: httpx.AsyncClient, request: httpx.Request, **_kwargs: Any
    ) -> httpx.Response:
        assert request.method == "GET"
        assert request.url.host == "apisidra.ibge.gov.br"
        parts = request.url.path.strip("/").split("/")
        assert parts[0] == "values"
        params = dict(zip(parts[1::2], parts[2::2], strict=True))
        table = params["t"]
        assert table in responses
        assert params["h"] == "n"
        assert params["n3"] == "all"
        assert set(params["v"].split(",")) == set(variables[table].split(","))
        requests.append(str(request.url))
        return httpx.Response(200, json=responses[table], request=request)

    monkeypatch.setattr(httpx.AsyncClient, "send", send)
    return requests


def install_reconciliacao_censos_legacy_http(
    monkeypatch: Any, files: list[dict[str, Any]]
) -> list[str]:
    golden = Path(__file__).parent / "golden_data"
    payloads = {}
    for file in files:
        raw = (golden / file["fixture"]).read_bytes()
        assert hashlib.sha256(raw).hexdigest() == file["sha256"]
        payloads[file["url"]] = raw
    requests: list[str] = []

    async def send(
        _client: httpx.AsyncClient, request: httpx.Request, **_kwargs: Any
    ) -> httpx.Response:
        assert request.method == "GET"
        url = str(request.url)
        assert url in payloads, url
        requests.append(url)
        return httpx.Response(200, content=payloads[url], request=request)

    monkeypatch.setattr(httpx.AsyncClient, "send", send)
    return requests


def assert_reconciliacao_censos_legacy_state(
    frame: pd.DataFrame, oracle: dict[str, Any], level: str
) -> None:
    metrics = oracle["metrics"]
    expected_rows = len(metrics)
    if level == "municipio":
        expected_rows *= oracle["source_counts"]["municipality_rows"]
    assert len(frame) == expected_rows
    assert frame["ano"].eq(1995).all()
    assert frame["uf"].eq(oracle["uf"]).all()
    assert frame["tema"].eq(oracle["theme"]).all()
    assert frame["categoria"].eq("Total").all()
    assert not frame.duplicated(["ano", "uf", "localidade", "tema", "categoria", "variavel"]).any()
    assert frame["valor"].dtype == "float64"
    if level == "municipio":
        assert frame["localidade_cod"].isna().all()
    else:
        assert frame["localidade_cod"].notna().all()
    for sample in oracle["samples"]:
        if (sample["selection"] == "uf_total") != (level == "uf"):
            continue
        selected = (
            frame if level == "uf" else frame.loc[frame["localidade"].eq(sample["label"].strip())]
        )
        assert len(selected) == len(metrics)
        for metric, cell in zip(metrics, sample["cells"], strict=True):
            row = selected.loc[selected["variavel"].eq(metric["variable"])]
            assert len(row) == 1
            assert row.iloc[0]["unidade"] == metric["unit"]
            if cell["expected"] is None:
                assert pd.isna(row.iloc[0]["valor"])
            else:
                assert row.iloc[0]["valor"] == pytest.approx(cell["expected"], rel=1e-12, abs=1e-9)


def assert_reconciliacao_censos_legacy_national(
    frame: pd.DataFrame, oracle: dict[str, Any]
) -> None:
    variables = {metric["variable"] for metric in oracle["metrics"]}
    selected = frame.loc[frame["variavel"].isin(variables)]
    expected_keys = {
        (category, variable) for category in oracle["categories"] for variable in variables
    }
    assert (
        set(selected[["categoria", "variavel"]].itertuples(index=False, name=None)) == expected_keys
    )
    assert len(selected) == len(expected_keys)
    assert selected["ano"].eq(1995).all()
    assert selected["localidade"].eq("Brasil").all()
    assert selected["localidade_cod"].eq(1).all()
    assert selected["uf"].isna().all()
    assert selected["valor"].dtype == "float64"
    for sample in oracle["samples"]:
        row = selected.loc[
            selected["categoria"].eq(sample["category"])
            & selected["variavel"].eq(sample["variable"])
        ]
        assert len(row) == 1
        assert row.iloc[0]["unidade"] == sample["unit"]
        assert row.iloc[0]["valor"] == pytest.approx(sample["value"], rel=1e-12, abs=1e-9)


REPLAY_PAGE_SIZE_PARAMS = frozenset({"$top", "$limit"})
REPLAY_CREDENTIAL_PARAMS = frozenset({"userid", "password"})
_REPLAY_FLOAT_EPSILON = 2.0**-52
_REPLAY_REAL_ASYNC_CLIENT = httpx.AsyncClient

ReplaySignature = tuple[str, tuple[tuple[str, str], ...], int]


def replay_signature(url: str) -> ReplaySignature:
    parts = urlsplit(unquote(url))
    params = dict(parse_qsl(parts.query, keep_blank_values=True))
    skip = int(params.pop("$skip", 0))
    for key in REPLAY_PAGE_SIZE_PARAMS | REPLAY_CREDENTIAL_PARAMS:
        params.pop(key, None)
    return f"{parts.scheme}://{parts.netloc}{parts.path}", tuple(sorted(params.items())), skip


def _replay_match_signature(match: dict[str, Any]) -> ReplaySignature:
    return match["path"], tuple(sorted(match["params"].items())), match["skip"]


def install_replay_http(
    monkeypatch: Any, case: dict[str, Any], golden_dir: Path
) -> dict[str, list[str]]:
    bodies = {
        (
            request["match"].get("method", "GET"),
            _replay_match_signature(request["match"]),
            request["match"].get("occurrence", 0),
        ): (
            (golden_dir / request["file"]).resolve().read_bytes(),
            request["content_type"],
            request.get("status", 200),
            request.get("response_headers", {}),
        )
        for request in case["requests"]
    }
    seen: dict[str, list[str]] = {"served": [], "unmatched": []}
    occurrences: dict[tuple[str, ReplaySignature], int] = {}

    def handler(request: httpx.Request) -> httpx.Response:
        key = (request.method, replay_signature(str(request.url)))
        occurrences[key] = occurrences.get(key, 0) + 1
        found = bodies.get((*key, occurrences[key]), bodies.get((*key, 0)))
        if found is None:
            seen["unmatched"].append(unquote(str(request.url)))
            return httpx.Response(404, content=b"{}", headers={"content-type": "application/json"})
        seen["served"].append(unquote(str(request.url)))
        body, content_type, status, headers = found
        return httpx.Response(
            status, content=body, headers={"content-type": content_type, **headers}
        )

    class ReplayClient(_REPLAY_REAL_ASYNC_CLIENT):
        def __init__(self, *args: Any, **kwargs: Any) -> None:
            kwargs["transport"] = httpx.MockTransport(handler)
            super().__init__(*args, **kwargs)

    monkeypatch.setattr(httpx, "AsyncClient", ReplayClient)
    return seen


def assert_replay_served(seen: dict[str, list[str]]) -> None:
    assert seen["served"], "nenhuma requisição"
    assert not seen["unmatched"], seen["unmatched"]


def _replay_is_null(value: Any) -> bool:
    return value is None or value is pd.NA or (isinstance(value, float) and math.isnan(value))


def replay_matches(got: Any, expected: Any, sample: dict[str, Any]) -> bool:
    if expected is None:
        return _replay_is_null(got)
    if sample.get("match") == "prefix":
        return str(got).startswith(str(expected))
    if isinstance(expected, bool) or not isinstance(expected, (int, float)):
        return str(got) == str(expected)
    if not isinstance(got, Real) or isinstance(got, bool) or _replay_is_null(got):
        return False
    if sample.get("tolerance") == "float_sum":
        bound = sample.get("terms", 300) * _REPLAY_FLOAT_EPSILON * abs(float(expected))
        return abs(float(got) - float(expected)) <= bound
    return float(got) == float(expected)


def _replay_key_mask(frame: pd.DataFrame, key: dict[str, Any]) -> pd.Series:
    mask = pd.Series(True, index=frame.index)
    for column, value in key.items():
        series = frame[column]
        if str(series.dtype).startswith("datetime"):
            mask &= series == pd.Timestamp(value)
        else:
            mask &= series.astype(str) == str(value)
    return mask


def assert_replay_samples(frame: pd.DataFrame, case: dict[str, Any]) -> None:
    for sample in case["samples"]:
        missing = [column for column in sample["key"] if column not in frame.columns]
        assert not missing, f"{case['id']}: colunas de chave ausentes {missing}"
        rows = frame[_replay_key_mask(frame, sample["key"])]
        assert len(rows) == 1, f"{case['id']}: chave {sample['key']} -> {len(rows)} linhas"
        assert sample["column"] in frame.columns, f"{case['id']}: coluna {sample['column']} ausente"
        got = rows.iloc[0][sample["column"]]
        if hasattr(got, "strftime") and sample["column"] in ("data", "mes"):
            got = got.strftime("%Y-%m-%d")
        assert replay_matches(got, sample["value"], sample), (
            f"{case['id']} {sample['key']} {sample['column']}: saída={got!r} "
            f"esperado={sample['value']!r}"
        )


def assert_replay_structure(frame: pd.DataFrame, case: dict[str, Any]) -> None:
    if "columns" in case:
        expected = set(case["columns"])
    else:
        expected = {
            item["destino"]
            for item in case["structure"]
            if item["estado"] in ("mapeada", "derivada", "ausente_na_fonte")
        }
    assert set(frame.columns) == expected, f"{case['id']}: colunas {sorted(frame.columns)}"
    for item in case["structure"]:
        assert item["estado"] != "desconhecida", f"{case['id']}: item desconhecido {item}"
        if item["estado"] == "ignorada":
            assert item["motivo"] and item["destino"] is None, item


async def replay_censo_legacy_tables(
    tema: str,
    uf: str | None,
    *,
    case: tuple[str, str],
    frames: list[pd.DataFrame],
    urls: list[str],
) -> tuple[list[pd.DataFrame], list[str]]:
    assert (uf, tema) == case
    return [frame.copy(deep=True) for frame in frames], list(urls)


def _assert_censo_legacy_values(actual: list[float], cells: list[dict[str, Any]]) -> None:
    assert len(actual) == len(cells)
    for value, cell in zip(actual, cells, strict=True):
        if cell["expected"] is None:
            assert pd.isna(value), cell["coordinate"]
        else:
            assert value == pytest.approx(cell["expected"], rel=1e-12), cell["coordinate"]


def assert_censo_legacy_matrix(frame: pd.DataFrame, oracle: dict[str, Any]) -> None:
    assert frame["uf"].eq(oracle["uf"]).all()
    assert len(frame) == oracle["source_counts"]["data_rows"] * len(oracle["metrics"])
    for level in ("uf", "municipio"):
        contracts.validate_dataset(
            frame.loc[frame["nivel_geo"].eq(level)].drop(columns="nivel_geo"),
            "censo_agropecuario_legado",
        )
    assert (
        frame.loc[frame["nivel_geo"].eq("municipio"), "localidade"].nunique()
        == oracle["source_counts"]["municipality_rows"]
    )
    assert frame["valor"].isna().sum() == oracle["source_counts"].get("missing", 0)
    assert frame["valor"].eq(0).sum() == oracle["source_counts"].get("zero", 0)
    for sample in oracle["samples"]:
        if sample["selection"] == "uf_total":
            selected = frame.loc[frame["nivel_geo"].eq("uf")]
            assert selected["localidade_cod"].notna().all()
        else:
            selected = frame.loc[
                frame["nivel_geo"].eq("municipio") & frame["localidade"].eq(sample["label"].strip())
            ]
            assert selected["localidade_cod"].isna().all()
        assert selected["variavel"].tolist() == [item["variable"] for item in oracle["metrics"]]
        assert selected["unidade"].tolist() == [item["unit"] for item in oracle["metrics"]]
        _assert_censo_legacy_values(selected["valor"].tolist(), sample["cells"])
    for example in oracle["marker_examples"].values():
        selected = frame.loc[
            frame["nivel_geo"].eq(example["level"])
            & frame["variavel"].eq(oracle["metrics"][example["column_index"]]["variable"])
        ]
        if example["level"] != "uf":
            selected = selected.loc[selected["localidade"].eq(example["label"].strip())]
        _assert_censo_legacy_values(selected["valor"].tolist(), [example["cell"]])
