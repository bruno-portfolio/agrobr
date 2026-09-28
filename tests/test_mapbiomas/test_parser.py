from __future__ import annotations

import pytest

from agrobr.exceptions import ParseError
from agrobr.mapbiomas.parser import (
    parse_cobertura_xlsx,
    parse_transicao_xlsx,
)


class TestParseCoberturaXlsx:
    def test_empty_xlsx_raises(self):
        from io import BytesIO

        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "COVERAGE_10"
        ws.append(["biome", "state", "class", "class_level_0"])
        buf = BytesIO()
        wb.save(buf)
        with pytest.raises(ParseError, match="COVERAGE vazia"):
            parse_cobertura_xlsx(colecao=10, data=buf.getvalue())

    def test_missing_columns_raises(self):
        from io import BytesIO

        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "COVERAGE_10"
        ws.append(["id", "nome", "valor"])
        ws.append([1, "teste", 100])
        buf = BytesIO()
        wb.save(buf)
        with pytest.raises(ParseError, match="Colunas obrigatorias ausentes"):
            parse_cobertura_xlsx(colecao=10, data=buf.getvalue())


class TestParseTransicaoXlsx:
    def test_missing_columns_raises(self):
        from io import BytesIO

        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "TRANSITION_10"
        ws.append(["id", "nome", "valor"])
        ws.append([1, "teste", 100])
        buf = BytesIO()
        wb.save(buf)
        with pytest.raises(ParseError, match="Colunas obrigatorias ausentes"):
            parse_transicao_xlsx(colecao=10, data=buf.getvalue())
