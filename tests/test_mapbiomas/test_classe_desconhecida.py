from __future__ import annotations

import io
import re
import warnings
from pathlib import Path
from unittest.mock import AsyncMock, patch

import openpyxl
import pytest

from agrobr import datasets
from agrobr.exceptions import ParseError
from agrobr.mapbiomas import client, parser
from tests.helpers import levanta_exatamente, mapbiomas_workbook_bundle, sem_excecao
from tests.test_mapbiomas.test_municipal_parser import _row, _xlsx

RECORTE = (
    Path(__file__).parents[1] / "golden_data/mapbiomas/legenda_20260918/col10_classe50_recorte.xlsx"
)
DESCONHECIDA = 999


def _recorte_com_classe_desconhecida() -> bytes:
    workbook = openpyxl.load_workbook(io.BytesIO(RECORTE.read_bytes()))
    for aba, coluna in (("COVERAGE_10", "class"), ("TRANSITION_10", "class_to")):
        sheet = workbook[aba]
        indice = [cell.value for cell in sheet[1]].index(coluna) + 1
        for linha in range(2, sheet.max_row + 1):
            sheet.cell(row=linha, column=indice, value=DESCONHECIDA)
    content = io.BytesIO()
    workbook.save(content)
    workbook.close()
    return content.getvalue()


def _recorte_com_celula(aba: str, coluna: str, valor: object) -> bytes:
    workbook = openpyxl.load_workbook(io.BytesIO(RECORTE.read_bytes()))
    sheet = workbook[aba]
    sheet.cell(row=2, column=[cell.value for cell in sheet[1]].index(coluna) + 1, value=valor)
    content = io.BytesIO()
    workbook.save(content)
    workbook.close()
    return content.getvalue()


async def _consultar(fetch: str, bundle, **kwargs):
    with (
        patch.object(client, fetch, new=AsyncMock(return_value=bundle)),
        warnings.catch_warnings(record=True) as emitidos,
        sem_excecao(),
    ):
        warnings.simplefilter("always")
        frame, meta = await datasets.uso_do_solo(return_meta=True, **kwargs)
    return frame, meta, [str(aviso.message) for aviso in emitidos]


def _conferir_aviso(meta, emitidos: list[str], colecao: int) -> None:
    esperado = f"MapBiomas coleção {colecao}: classes [{DESCONHECIDA}] fora da legenda conhecida"
    assert [aviso for aviso in meta.validation_warnings if aviso.startswith(esperado)]
    assert [aviso for aviso in emitidos if aviso.startswith(esperado)]


@pytest.mark.parametrize(
    ("tipo", "rotulo", "coluna_id"),
    [("cobertura", "classe", "classe_id"), ("transicao", "classe_para", "classe_para_id")],
)
async def test_classe_fora_da_legenda_no_estadual_sai_nula_com_aviso(tipo, rotulo, coluna_id):
    bundle = mapbiomas_workbook_bundle(
        _recorte_com_classe_desconhecida(), client._build_xlsx_url("BIOME_STATE", 10)
    )

    frame, meta, emitidos = await _consultar(
        "fetch_biome_state_bundle", bundle, tipo=tipo, colecao=10
    )

    assert len(frame) > 0
    assert frame[coluna_id].eq(DESCONHECIDA).all()
    assert frame[rotulo].isna().all()
    assert meta.contract_version == "2.0"
    _conferir_aviso(meta, emitidos, 10)


async def test_classe_fora_da_legenda_no_municipal_sai_nula_com_aviso():
    conteudo = _xlsx(
        [_row(), _row(ID=2, geocode="5300109", municipality="Outro", **{"class": DESCONHECIDA})]
    )
    bundle = mapbiomas_workbook_bundle(
        conteudo, client._build_xlsx_url("BIOME_STATE_MUNICIPALITY", 11)
    )

    frame, meta, emitidos = await _consultar(
        "fetch_biome_state_municipality_bundle", bundle, nivel="municipio", estado="DF", ano=1985
    )

    rotulos = dict(zip(frame["classe_id"], frame["classe"], strict=True))
    assert set(rotulos) == {3, DESCONHECIDA}
    assert rotulos[3] == "Formação Florestal"
    assert rotulos[DESCONHECIDA] is None
    _conferir_aviso(meta, emitidos, 11)


@pytest.mark.parametrize(
    ("aba", "coluna", "valor", "parse"),
    [
        ("COVERAGE_10", "class", "abc", parser.parse_cobertura_xlsx),
        ("TRANSITION_10", "class_from", "abc", parser.parse_transicao_xlsx),
        ("TRANSITION_10", "class_to", 3.5, parser.parse_transicao_xlsx),
    ],
)
def test_classe_sem_codigo_inteiro_no_estadual_recusa_a_planilha(aba, coluna, valor, parse):
    conteudo = _recorte_com_celula(aba, coluna, valor)
    with levanta_exatamente(
        ParseError, re.escape(f"{coluna} sem código inteiro na linha 2: {valor!r}")
    ):
        parse(conteudo, colecao=10)
