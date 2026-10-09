from __future__ import annotations

import functools
import io
import warnings
from types import SimpleNamespace
from urllib.parse import parse_qsl, urlsplit

import httpx
import openpyxl
import pandas as pd
import pytest

from agrobr import unica
from agrobr.exceptions import InvalidParameterError, ParseError, SourceUnavailableError
from agrobr.unica import api, client, parser
from tests import helpers
from tests.test_unica import oficial

PRODUTOS = ("cana", "acucar", "etanol_total", "etanol_anidro", "etanol_hidratado")
HISTORICOS = {
    ("acucar", "2018/2019", "2020/2021"): oficial.GOLDEN / "hist_acucar.xlsx",
    ("etanol_total", "2018/2019", "2020/2021"): oficial.GOLDEN / "hist_etanol_total.xlsx",
    ("cana", "1980/1981", "2020/2021"): oficial.OFICIAL / "historico_cana_1980_2021.xlsx",
}


def instalar(
    monkeypatch: pytest.MonkeyPatch,
    edicao: str = "01/07/2026",
    trocar: dict[str, bytes] | None = None,
) -> list[httpx.Request]:
    trocar = trocar or {}
    pedidos: list[httpx.Request] = []

    def responder(request: httpx.Request) -> httpx.Response:
        pedidos.append(request)
        caminho = request.url.path
        if caminho == "/listagem.php":
            corpo = trocar.get("listagem", oficial.corpo("listagem_idMn63.html"))
            return httpx.Response(200, content=corpo, request=request)
        if str(request.url) == oficial.PDF_URL:
            corpo = trocar.get("pdf", oficial.EDICOES[edicao].read_bytes())
            return httpx.Response(200, content=corpo, request=request)
        if caminho == "/xlsHPM.php":
            chave = tuple(request.url.params[k] for k in ("produto", "safraIni", "safraFim"))
            if "historico" in trocar:
                return httpx.Response(200, content=trocar["historico"], request=request)
            if chave in HISTORICOS:
                return httpx.Response(200, content=HISTORICOS[chave].read_bytes(), request=request)
        return httpx.Response(404, request=request)

    namespace = SimpleNamespace(**vars(httpx))
    namespace.AsyncClient = lambda **kwargs: httpx.AsyncClient(
        transport=httpx.MockTransport(responder), **kwargs
    )
    monkeypatch.setattr(client, "httpx", namespace)
    monkeypatch.setattr(client, "_pdf_cache", None)
    monkeypatch.setattr(api, "_parsed_cache", None)
    return pedidos


@functools.cache
def _textos(edicao: str) -> tuple[str, ...]:
    real = parser._check_pdfplumber()
    with real.open(io.BytesIO(oficial.EDICOES[edicao].read_bytes())) as pdf:
        return tuple(pagina.extract_text() or "" for pagina in pdf.pages)


class _Documento:
    def __init__(self, textos: list[str]) -> None:
        self.pages = [SimpleNamespace(extract_text=lambda texto=texto: texto) for texto in textos]

    def __enter__(self) -> _Documento:
        return self

    def __exit__(self, *args: object) -> None:
        return None


def pdf_alterado(monkeypatch: pytest.MonkeyPatch, antes: str, depois: str) -> None:
    textos = list(_textos("01/07/2026"))
    assert any(antes in texto for texto in textos), antes
    alterados = [texto.replace(antes, depois) for texto in textos]
    monkeypatch.setattr(
        parser,
        "_check_pdfplumber",
        lambda: SimpleNamespace(open=lambda _: _Documento(alterados)),
    )


async def test_resumo_de_cada_edicao_confere_com_o_pdf(monkeypatch):
    with helpers.collect_failures() as check:
        for edicao, ausente in (("01/07/2026", "quinzena"), ("01/05/2026", "mensal")):
            instalar(monkeypatch, edicao)
            for periodo in oficial.periodos(edicao):
                with check(f"{edicao} {periodo}"):
                    with helpers.sem_excecao():
                        df, meta = await unica.safra_resumo(periodo=periodo, return_meta=True)
                    assert list(df.columns) == oficial.COLUNAS_RESUMO
                    assert oficial.publicado(df) == oficial.esperado_resumo(edicao, periodo)
                    assert meta.source_url == oficial.PDF_URL
                    helpers.conferir_corpo(meta, oficial.EDICOES[edicao].read_bytes())
            disponiveis = str(oficial.periodos(edicao)).replace("[", r"\[").replace("]", r"\]")
            with (
                check(f"{edicao} sem {ausente}"),
                helpers.levanta_exatamente(
                    InvalidParameterError, match=f"{edicao}.*'{ausente}'.*{disponiveis}"
                ),
            ):
                await unica.safra_resumo(periodo=ausente)


async def test_series_de_cada_edicao_conferem_com_o_pdf(monkeypatch):
    with helpers.collect_failures() as check:
        for edicao in oficial.EDICOES:
            instalar(monkeypatch, edicao)
            for produto in PRODUTOS:
                with check(f"{edicao} {produto}"):
                    with helpers.sem_excecao():
                        df, meta = await unica.moagem_quinzenal(produto, return_meta=True)
                    assert list(df.columns) == oficial.COLUNAS_SERIES
                    assert oficial.publicado(df) == oficial.esperado_series(edicao, produto)
                    assert meta.source_url == oficial.PDF_URL
                    helpers.conferir_corpo(meta, oficial.EDICOES[edicao].read_bytes())
        with check("regiao"):
            instalar(monkeypatch, "01/07/2026")
            with helpers.sem_excecao():
                df = await unica.moagem_quinzenal("acucar", regiao="centro_sul")
            assert oficial.publicado(df) == [
                r
                for r in oficial.esperado_series("01/07/2026", "acucar")
                if r["regiao"] == "centro_sul"
            ]


async def test_historico_confere_com_a_planilha(monkeypatch):
    pedidos = instalar(monkeypatch)
    oficiais = dict(
        parse_qsl(urlsplit(oficial.manifest()["respostas"][2]["url"]).query, keep_blank_values=True)
    )

    with helpers.collect_failures() as check:
        for (produto, inicio, fim), arquivo in HISTORICOS.items():
            with check(produto):
                pedidos.clear()
                filtros = (
                    {} if inicio == "1980/1981" else {"safra_inicio": inicio, "safra_fim": fim}
                )
                with helpers.sem_excecao():
                    df, meta = await unica.producao_historica(produto, **filtros, return_meta=True)
                assert dict(pedidos[0].url.params) == {
                    **oficiais,
                    "produto": produto,
                    "safraIni": inicio,
                    "safraFim": fim,
                }
                assert oficial.publicado(df) == oficial.esperado_historico(arquivo, produto)
                helpers.conferir_corpo(meta, arquivo.read_bytes())
                assert meta.source_url.endswith(
                    f"produto={produto}&safraIni={inicio}&safraFim={fim}"
                )


async def test_historico_aceita_safra_curta_e_pede_o_formato_da_fonte(monkeypatch):
    pedidos = instalar(monkeypatch)

    curta = await unica.producao_historica("acucar", safra_inicio="2018/19", safra_fim="20/21")
    completa = await unica.producao_historica(
        "acucar", safra_inicio="2018/2019", safra_fim="2020/2021"
    )

    assert [(p.url.params["safraIni"], p.url.params["safraFim"]) for p in pedidos] == [
        ("2018/2019", "2020/2021")
    ] * 2
    assert curta.equals(completa)


async def test_parametros_invalidos_recusados_antes_da_rede(monkeypatch):
    pedidos = instalar(monkeypatch)

    with helpers.collect_failures() as check:
        for nome, chamada, erro in (
            ("produto quinzenal", lambda: unica.moagem_quinzenal("soja"), InvalidParameterError),
            ("regiao", lambda: unica.moagem_quinzenal(regiao="nordeste"), InvalidParameterError),
            ("periodo", lambda: unica.safra_resumo(periodo="semanal"), InvalidParameterError),
            ("produto historico", lambda: unica.producao_historica("milho"), InvalidParameterError),
            ("safra", lambda: unica.producao_historica(safra_inicio="2018"), InvalidParameterError),
            (
                "safra com anos não consecutivos",
                lambda: unica.producao_historica(safra_inicio="2018/2020"),
                InvalidParameterError,
            ),
            (
                "intervalo invertido",
                lambda: unica.producao_historica(safra_inicio="2019/2020", safra_fim="2018/2019"),
                InvalidParameterError,
            ),
            (
                "safra depois da última publicada",
                lambda: unica.producao_historica(safra_fim="2024/2025"),
                InvalidParameterError,
            ),
            (
                "safra que não é texto",
                lambda: unica.producao_historica(safra_inicio=2018),
                InvalidParameterError,
            ),
        ):
            with check(nome), helpers.levanta_exatamente(erro):
                await chamada()
    assert pedidos == []


async def test_download_invalido_nao_vira_dado():
    manutencao = b"<html><body>manutencao</body></html>" * 3000
    with helpers.collect_failures() as check:
        for nome, trocar, chamada, erro, mensagem in (
            (
                "listagem sem PDF",
                {"listagem": manutencao},
                lambda: unica.safra_resumo(),
                ParseError,
                "URL do PDF",
            ),
            (
                "PDF que é HTML",
                {"pdf": manutencao},
                lambda: unica.safra_resumo(),
                SourceUnavailableError,
                "PDF inválido",
            ),
            (
                "histórico que é HTML",
                {"historico": manutencao},
                lambda: unica.producao_historica("cana"),
                SourceUnavailableError,
                "não é XLSX",
            ),
        ):
            with check(nome), pytest.MonkeyPatch.context() as mp:
                instalar(mp, trocar=trocar)
                with helpers.levanta_exatamente(erro, match=mensagem):
                    await chamada()


async def test_layout_quebrado_vira_erro():
    with helpers.collect_failures() as check:
        for nome, antes, depois, mensagem in (
            ("título de período", "posição MENSAL", "posição SEMANAL", "SEMANAL"),
            ("resumo não numérico", "115.933 118.090", "115.933 xyz", "valor não numérico 'xyz'"),
            (
                "cana ausente do acumulado",
                "Cana-de-açúcar ¹ 206.573 214.470 3,82% 115.933 118.090 1,86% 90.640 96.380 6,33%",
                "",
                r"ausentes no resumo \(acumulado\): \['cana'\]",
            ),
            (
                "série com 8 valores",
                "16/04 8.253.154 11.007.246 33%",
                "16/04 8.253.154 11.007.246",
                "esperado 9",
            ),
            ("série não numérica", "11.007.246 33%", "x 33%", "não numérico"),
            ("ATR/t fora do plausível", "122,18 123,75", "122,18 323,75", "ATR/tonelada"),
            ("mix acima de 100", "51,04% 42,52%", "51,04% 142,52%", "Mix"),
            ("série negativa", "11.007.246 33%", "-11.007.246 33%", "negativos"),
            ("série acima do plausível", "214.470.277", "914.470.277", "acima do plausível"),
            ("tabela ausente", "Tabela 5. Histórico", "Quadro 5. Histórico", "Tabelas ausentes"),
            ("capa sem posição", "Posição até 01/07/2026", "Posição em julho", "Capa"),
        ):
            with check(nome), pytest.MonkeyPatch.context() as mp:
                instalar(mp)
                pdf_alterado(mp, antes, depois)
                with helpers.levanta_exatamente(ParseError, match=mensagem):
                    await unica.safra_resumo()


def planilha_alterada(linha: str, coluna: int, valor: object) -> bytes:
    livro = openpyxl.load_workbook(
        io.BytesIO(HISTORICOS[("cana", "1980/1981", "2020/2021")].read_bytes())
    )
    folha = livro.active
    alvo = next(c for c in folha["A"] if isinstance(c.value, str) and c.value.strip() == linha)
    folha.cell(row=alvo.row, column=coluna).value = valor
    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()


async def test_planilha_quebrada_vira_erro():
    with helpers.collect_failures() as check:
        for nome, linha, coluna, valor, mensagem in (
            ("valor negativo", "São Paulo", 2, -1, "negativos"),
            ("acima do plausível", "Brasil", 2, 900_000, "acima do plausível"),
            (
                "unidade desconhecida",
                "Unidade: Mil toneladas",
                1,
                "Unidade: Mil litros",
                "Unidade desconhecida",
            ),
            ("localidade desconhecida", "Acre", 1, "Atlântida", "Localidade não reconhecida"),
            (
                "sem linha de unidade",
                "Unidade: Mil toneladas",
                1,
                "Medida: Mil toneladas",
                "Unidade:' não encontrada",
            ),
            ("sem cabeçalho", "Estado/Safra", 1, "UF/Safra", "Estado/Safra' não encontrado"),
        ):
            with check(nome), pytest.MonkeyPatch.context() as mp:
                instalar(mp, trocar={"historico": planilha_alterada(linha, coluna, valor)})
                with helpers.levanta_exatamente(ParseError, match=mensagem):
                    await unica.producao_historica("cana")
        with check("sem a linha Source"), pytest.MonkeyPatch.context() as mp:
            fonte = next(
                c.value
                for c in openpyxl.load_workbook(
                    io.BytesIO(HISTORICOS[("cana", "1980/1981", "2020/2021")].read_bytes())
                ).active["A"]
                if isinstance(c.value, str) and c.value.startswith("Source")
            )
            instalar(mp, trocar={"historico": planilha_alterada(fonte.strip(), 1, None)})
            with helpers.sem_excecao():
                df = await unica.producao_historica("cana")
            assert oficial.publicado(df) == oficial.esperado_historico(
                HISTORICOS[("cana", "1980/1981", "2020/2021")], "cana"
            )


async def test_valor_ausente_no_resumo_vira_nulo():
    with helpers.collect_failures() as check:
        for nome, antes, depois, regiao, coluna in (
            ("atual", "115.933 118.090", "115.933 n/d", "sao_paulo", "valor"),
            ("anterior", "206.573 214.470", "- 214.470", "centro_sul", "valor_safra_anterior"),
        ):
            with check(nome), pytest.MonkeyPatch.context() as mp:
                instalar(mp)
                pdf_alterado(mp, antes, depois)
                with helpers.sem_excecao():
                    df = await unica.safra_resumo(periodo="acumulado")
                esperado = oficial.esperado_resumo("01/07/2026", "acumulado")
                for linha in esperado:
                    if linha["produto"] == "cana" and linha["regiao"] == regiao:
                        linha[coluna] = None
                assert oficial.publicado(df) == esperado


async def test_valor_ausente_na_serie_vira_nulo():
    with helpers.collect_failures() as check:
        for nome, antes, depois, regiao, coluna in (
            ("atual", "11.007.246 33%", "n/d 33%", "sao_paulo", "valor"),
            ("variação", "20.403.906 22%", "20.403.906 -", "centro_sul", "variacao_pct"),
            (
                "anterior",
                "8.422.586 9.396.660",
                "– 9.396.660",
                "demais_estados",
                "valor_safra_anterior",
            ),
        ):
            with check(nome), pytest.MonkeyPatch.context() as mp:
                instalar(mp)
                pdf_alterado(mp, antes, depois)
                with helpers.sem_excecao():
                    df = await unica.moagem_quinzenal("cana")
                esperado = oficial.esperado_series("01/07/2026", "cana")
                for linha in esperado:
                    if linha["quinzena"] == "16/04" and linha["regiao"] == regiao:
                        linha[coluna] = None
                assert oficial.publicado(df) == esperado


async def test_quinzena_de_janeiro_e_do_ano_seguinte(monkeypatch):
    instalar(monkeypatch)
    pdf_alterado(monkeypatch, "16/06 91.441.515", "16/01 91.441.515")

    with helpers.sem_excecao():
        df = await unica.moagem_quinzenal("cana", regiao="sao_paulo")

    assert df.loc[df["quinzena"] == "16/01", "data"].tolist() == [pd.Timestamp(2027, 1, 16)]


async def test_aviso_de_licenca(monkeypatch):
    instalar(monkeypatch)

    with warnings.catch_warnings(record=True) as avisos, helpers.sem_excecao():
        warnings.simplefilter("always")
        await unica.producao_historica("cana")

    assert [str(a.message) for a in avisos if "UNICA" in str(a.message)] == [api._LICENSE_WARNING]


async def test_as_polars_publica_os_mesmos_valores(monkeypatch):
    pl = pytest.importorskip("polars")
    instalar(monkeypatch)

    with helpers.sem_excecao():
        resumo = await unica.safra_resumo(periodo="mensal", as_polars=True)
        serie = await unica.moagem_quinzenal("cana", as_polars=True)
        historico = await unica.producao_historica("cana", as_polars=True)

    assert all(isinstance(frame, pl.DataFrame) for frame in (resumo, serie, historico))
    assert oficial.publicado(resumo.to_pandas()) == oficial.esperado_resumo("01/07/2026", "mensal")
    assert oficial.publicado(serie.to_pandas()) == oficial.esperado_series("01/07/2026", "cana")
    assert oficial.publicado(historico.to_pandas()) == oficial.esperado_historico(
        oficial.OFICIAL / "historico_cana_1980_2021.xlsx", "cana"
    )
