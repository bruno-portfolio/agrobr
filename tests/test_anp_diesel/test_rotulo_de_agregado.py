from __future__ import annotations

import warnings
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd

from agrobr import datasets
from agrobr.alt.anp_diesel import api, parser
from tests.helpers import sem_excecao
from tests.test_anp_diesel import test_reconciliacao as reconciliacao

CONSULTA = {
    "nivel": "municipio",
    "uf": "MT",
    "produto": "DIESEL",
    "inicio": "2025-12-01",
    "fim": "2026-01-31",
    "return_meta": True,
}
AVISO = "anp_diesel: 1 linha(s) com rótulo de agregado na coluna de município descartada(s): TOTAL"


def _com_total(destino: Path) -> dict[str, Path]:
    """Cópia de uma linha de MT no meio da planilha, com o município "TOTAL" e preço 99,99."""
    livro = openpyxl.load_workbook(reconciliacao.GOLDEN / "request_04.xlsx")
    folha = livro.worksheets[0]
    modelo = next(
        linha
        for linha in folha.iter_rows(min_row=13)
        if linha[3].value == "MATO GROSSO"
        and linha[5].value == "OLEO DIESEL"
        and linha[0].value >= datetime(2025, 12, 1)
    )
    valores = [celula.value for celula in modelo]
    valores[4], valores[8] = "TOTAL", 99.99
    meio = (13 + folha.max_row) // 2
    folha.insert_rows(meio)
    for coluna, valor in enumerate(valores, start=1):
        folha.cell(meio, coluna, valor)
    caminho = destino / "com_total.xlsx"
    livro.save(caminho)
    return {"request_04.xlsx": caminho}


def test_parser_descarta_a_linha_total_com_aviso(tmp_path):
    filtros = {"produto": "DIESEL", "uf": "MT", "nivel": "municipio"}
    original = parser.parse_precos(
        (reconciliacao.GOLDEN / "request_04.xlsx").read_bytes(), **filtros
    )

    with sem_excecao():
        frame = parser.parse_precos(_com_total(tmp_path)["request_04.xlsx"].read_bytes(), **filtros)

    assert frame.attrs.pop("avisos") == [AVISO]
    assert original.attrs.pop("avisos") == []
    pd.testing.assert_frame_equal(frame, original)


async def test_linha_total_sai_do_resultado_com_aviso(monkeypatch, tmp_path):
    reconciliacao.install_inputs(monkeypatch)
    original, _ = await api.precos_diesel(**CONSULTA)
    reconciliacao.install_inputs(monkeypatch, _com_total(tmp_path))

    with warnings.catch_warnings(record=True) as emitidos, sem_excecao():
        warnings.simplefilter("always")
        frame, meta = await api.precos_diesel(**CONSULTA)

    pd.testing.assert_frame_equal(frame, original)
    assert "TOTAL" not in set(frame["municipio"])
    assert meta.validation_warnings == [AVISO]
    assert [str(aviso.message) for aviso in emitidos if "agregado" in str(aviso.message)] == [AVISO]


async def test_aviso_da_linha_total_chega_ao_dataset(monkeypatch, tmp_path):
    reconciliacao.install_inputs(monkeypatch, _com_total(tmp_path))

    with warnings.catch_warnings(record=True), sem_excecao():
        warnings.simplefilter("always")
        frame, meta = await datasets.precos_diesel(**CONSULTA)

    assert "TOTAL" not in set(frame["municipio"])
    assert AVISO in meta.validation_warnings
