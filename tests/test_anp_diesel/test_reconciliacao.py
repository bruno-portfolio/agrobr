from __future__ import annotations

import hashlib
import json
import math
import os
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any
from unittest.mock import AsyncMock

import openpyxl
import pandas as pd
import pytest

from agrobr import datasets
from agrobr.alt.anp_diesel import api, client, models, parser
from agrobr.exceptions import ContractViolationError, InvalidParameterError, ParseError
from scripts import reconciliar_anp_precos as reconciliation
from tests import helpers

ROOT = Path(__file__).resolve().parents[2]
GOLDEN = ROOT / "tests/golden_data/reconciliacao_registros_precos_zoneamento_seguro_20260918/anp"
MANIFEST = json.loads((GOLDEN / "manifest.json").read_text(encoding="utf-8"))
ORIGINALS = os.environ.get("AGROBR_RECONCILIACAO_ANP_ORIGINALS", "")
KEYS = ["data", "nivel", "uf", "municipio", "produto"]
DATES = {"data", "periodo_inicio", "periodo_fim"}
NUMERIC = {"preco_venda", "preco_compra", "margem", "n_postos", "n_postos_media", "n_semanas"}


def install_inputs(
    monkeypatch: pytest.MonkeyPatch, replacements: dict[str, Path] | None = None
) -> dict[str, list[str]]:
    requests = []
    for resource in MANIFEST["resources"]:
        original = resource["original"]
        path = (
            Path(ORIGINALS) / original["file"] if ORIGINALS else GOLDEN / resource["derived_file"]
        )
        if replacements:
            path = replacements.get(resource["derived_file"], path)
        requests.append(
            {
                "match": {"path": original["requested_url"], "params": {}, "skip": 0},
                "file": str(path),
                "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            }
        )
    catalog = MANIFEST["catalog"]
    requests.append(
        {
            "match": {"path": catalog["requested_url"], "params": {}, "skip": 0},
            "file": catalog["golden_file"],
            "content_type": "text/html;charset=utf-8",
        }
    )
    return helpers.install_replay_http(monkeypatch, {"requests": requests}, GOLDEN)


def assert_expected(frame: pd.DataFrame, case: dict[str, Any]) -> None:
    expected = case["expected"]
    assert len(frame) == len(expected), (case["id"], len(frame), len(expected))
    assert not frame.duplicated(KEYS).any()
    assert set(frame.columns) == set(expected[0]) - {"locators"}
    actual = {
        (row.data.strftime("%Y-%m-%d"), row.nivel, row.uf, row.municipio, row.produto): row
        for row in frame.itertuples(index=False)
    }
    for row in expected:
        key = tuple(row[field] for field in KEYS)
        assert key in actual, (case["id"], key)
        found = actual[key]
        for field, value in row.items():
            if field == "locators":
                continue
            got = getattr(found, field)
            if value is None:
                assert pd.isna(got), (key, field, got)
            elif field in DATES:
                assert got.strftime("%Y-%m-%d") == value, (key, field, got, value)
            elif field in NUMERIC:
                assert math.isfinite(float(got)), (key, field, got)
                scale = max(1.0, abs(value))
                if field == "margem":
                    scale = max(scale, row["preco_venda"], row["preco_compra"] or 0)
                bound = 2 * (row["n_semanas"] + 2) * 2**-52 * scale
                assert abs(float(got) - value) <= bound, (key, field, got, value, bound)
            else:
                assert got == value, (key, field, got, value)


def overlapping_workbooks(target: Path, *, conflicting: bool) -> tuple[dict[str, Path], float]:
    before = openpyxl.load_workbook(GOLDEN / "request_03.xlsx")
    after = openpyxl.load_workbook(GOLDEN / "request_04.xlsx", read_only=True, data_only=True)
    sheet = before.worksheets[0]
    values = next(
        list(row)
        for row in after.worksheets[0].iter_rows(min_row=13, values_only=True)
        if isinstance(row[0], datetime)
        and row[0].date() == date(2023, 12, 31)
        and row[3] == "SAO PAULO"
        and row[4] == "SAO PAULO"
        and row[5] == "OLEO DIESEL S10"
    )
    published_price = float(values[8])
    if conflicting:
        values[8] += 0.01
    sheet.append(values)
    path = target / "overlapping.xlsx"
    before.save(path)
    before.close()
    after.close()
    return {"request_03.xlsx": path}, published_price


@pytest.mark.parametrize("layer", ["source", "dataset"])
@pytest.mark.parametrize("aggregation", ["semanal", "mensal"])
@pytest.mark.parametrize("conflicting", [False, True])
async def test_anp_distingue_repeticao_identica_de_conflito_entre_arquivos(
    monkeypatch, tmp_path, layer, aggregation, conflicting
):
    replacements, published_price = overlapping_workbooks(tmp_path, conflicting=conflicting)
    seen = install_inputs(monkeypatch, replacements)
    fetch = api.precos_diesel if layer == "source" else datasets.precos_diesel
    query = {
        "nivel": "municipio",
        "uf": "SP",
        "municipio": "São Paulo",
        "produto": "DIESEL S10",
        "inicio": "2023-12-31",
        "fim": "2023-12-31",
        "agregacao": aggregation,
    }
    if conflicting:
        with pytest.raises(ParseError, match="duplicadas") as caught:
            await fetch(**query)
        if layer == "dataset":
            assert caught.value.errors and all(
                kind == "parse" for _, kind, _ in caught.value.errors
            )
            assert isinstance(caught.value.__cause__, ParseError)
    else:
        with pytest.warns(UserWarning, match="linhas semanais idênticas"):
            frame, meta = await fetch(**query, return_meta=True)
        assert frame["preco_venda"].tolist() == [published_price]
        assert frame["municipio"].tolist() == ["SAO PAULO"]
        assert frame["periodo_inicio"].tolist() == [pd.Timestamp("2023-12-31")]
        assert frame["n_semanas"].tolist() == [1]
        assert any("idênticas" in aviso for aviso in meta.validation_warnings)
    helpers.assert_replay_served(seen)


@pytest.mark.parametrize("layer", ["source", "dataset"])
@pytest.mark.parametrize("case", MANIFEST["cases"], ids=lambda case: case["id"])
async def test_publicacao_anp_celulas_originais(
    layer: str, case: dict[str, Any], monkeypatch: pytest.MonkeyPatch
):
    seen = install_inputs(monkeypatch)
    function = api.precos_diesel if layer == "source" else datasets.precos_diesel
    frame, meta = await function(**case["query"], return_meta=True)
    assert_expected(frame, case)
    helpers.assert_replay_served(seen)
    assert meta.selected_source == "anp_diesel"
    assert meta.attempted_sources == ["anp_diesel"]
    assert meta.fetch_timestamp.utcoffset().total_seconds() == 0
    assert meta.validation_passed
    assert meta.source_details["unit"] == "BRL/litro"
    assert meta.source_details["filters"]["inicio"] == case["query"]["inicio"]
    assert meta.source_details["filters"]["fim"] == case["query"]["fim"]
    for receipt in meta.source_details["resources"]:
        resource = next(
            item
            for item in MANIFEST["resources"]
            if item["original"]["requested_url"] == receipt["requested_url"]
        )
        expected_hash = resource["original"]["sha256"] if ORIGINALS else resource["derived_sha256"]
        assert receipt["sha256"] == expected_hash


@pytest.mark.parametrize("resource", MANIFEST["resources"], ids=lambda item: item["derived_file"])
def test_recorte_anp_tem_hash_e_ultima_linha(resource: dict[str, Any]):
    path = GOLDEN / resource["derived_file"]
    assert hashlib.sha256(path.read_bytes()).hexdigest() == resource["derived_sha256"]
    original_positions = {row["original_row"] for row in resource["selected_rows"]}
    assert resource["last_real_row"] in original_positions
    assert set(resource["last_diesel_rows"].values()) <= original_positions


def mutated_workbook(resource: dict[str, Any], mutation: str, target: Path) -> Path:
    book = openpyxl.load_workbook(GOLDEN / resource["derived_file"])
    sheet = book.worksheets[0]
    header_row = resource["layout"]["header_row"]
    columns = {cell.value: cell.column for cell in sheet[header_row] if cell.value is not None}
    selected = [
        index
        for index in range(header_row + 1, sheet.max_row + 1)
        if sheet.cell(index, columns["PRODUTO"]).value == "OLEO DIESEL S10"
    ]
    last = selected[-1]
    if mutation == "value":
        cell = sheet.cell(last, columns["PREÇO MÉDIO REVENDA"])
        cell.value += 0.01
    elif mutation == "period":
        for name in ("DATA INICIAL", "DATA FINAL"):
            sheet.cell(last, columns[name]).value += timedelta(days=7)
    elif mutation == "remove_last":
        sheet.delete_rows(last)
    elif mutation == "key":
        sheet.cell(last, columns["PRODUTO"]).value = "OLEO DIESEL"
    elif mutation == "unit":
        sheet.cell(last, columns["UNIDADE DE MEDIDA"]).value = "R$/m3"
    elif mutation == "unused_header":
        sheet.cell(header_row, columns["DESVIO PADRÃO REVENDA"]).value = "CAMPO SEM DECISAO"
    elif mutation == "new_column":
        sheet.cell(header_row, sheet.max_column + 1).value = "COLUNA NOVA"
    elif mutation == "new_sheet":
        book.create_sheet("NOVA PUBLICACAO")
    elif mutation == "new_product":
        sheet.cell(last, columns["PRODUTO"]).value = "OLEO DIESEL NOVO"
    else:
        raise ValueError(mutation)
    book.save(target)
    book.close()
    return target


@pytest.mark.parametrize("layer", ["source", "dataset"])
@pytest.mark.parametrize("mutation", ["value", "period", "remove_last", "key", "unit"])
async def test_mutacao_corpo_anp_detectada(
    layer: str, mutation: str, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
):
    resource = MANIFEST["resources"][0]
    case = next(item for item in MANIFEST["cases"] if item["id"] == "brasil_atual_semanal")
    path = mutated_workbook(resource, mutation, tmp_path / "mutated.xlsx")
    seen = install_inputs(monkeypatch, {resource["derived_file"]: path})
    function = api.precos_diesel if layer == "source" else datasets.precos_diesel
    with pytest.raises((AssertionError, ParseError, ContractViolationError)):
        frame = await function(**case["query"])
        assert_expected(frame, case)
    helpers.assert_replay_served(seen)


@pytest.mark.parametrize("resource", MANIFEST["resources"], ids=lambda item: item["derived_file"])
def test_n1_anp_estrutura_conhecida(resource: dict[str, Any]):
    assert (
        reconciliation.compare_workbook((GOLDEN / resource["derived_file"]).read_bytes(), resource)[
            "status"
        ]
        == "ok"
    )


@pytest.mark.parametrize(
    "mutation", ["unit", "unused_header", "new_column", "new_sheet", "new_product"]
)
def test_n1_anp_estrutura_desconhecida(mutation: str, tmp_path: Path):
    resource = MANIFEST["resources"][0]
    path = mutated_workbook(resource, mutation, tmp_path / "mutated.xlsx")
    assert reconciliation.compare_workbook(path.read_bytes(), resource)["status"] == "mismatch"


def test_n1_anp_catalogo_conhecido():
    catalog = MANIFEST["catalog"]
    html = (GOLDEN / catalog["golden_file"]).read_text(encoding="utf-8")
    assert reconciliation.compare_catalog(html, catalog)["status"] == "ok"


@pytest.mark.parametrize("mutation", ["new_workbook", "missing_workbook"])
def test_n1_anp_catalogo_desconhecido(mutation: str):
    catalog = MANIFEST["catalog"]
    html = (GOLDEN / catalog["golden_file"]).read_text(encoding="utf-8")
    url = MANIFEST["resources"][-1]["original"]["url"]
    if mutation == "new_workbook":
        html += f'<a href="{url.replace("2026.xlsx", "2027.xlsx")}">Nova publicação</a>'
    else:
        html = html.replace(url, "https://example.test/removed")
    assert reconciliation.compare_catalog(html, catalog)["status"] == "mismatch"


async def test_limite_dezembro_descobre_arquivo_adjacente_publicado(
    monkeypatch: pytest.MonkeyPatch,
):
    monkeypatch.setattr(api.time_utils, "hoje", lambda: date(2027, 2, 1))
    catalog = dict(models.PRECOS_MUNICIPIOS_URLS)
    catalog.pop("2026")
    url = f"{models.SHLP_BASE}/semanal/semanal-municipios-2026-2027.xlsx"
    catalog["2026-2027"] = url
    fetch = AsyncMock(return_value=catalog)
    monkeypatch.setattr(client, "fetch_precos_catalog", fetch)
    query = {"nivel": "municipio", "inicio": date(2026, 12, 31), "fim": date(2026, 12, 31)}
    assert await api._resolve_price_urls(query) == [url]
    fetch.assert_awaited_once()


@pytest.mark.parametrize("municipio", ["São Paulo", "sao paulo", 3550308, "3550308"])
async def test_municipio_por_nome_ou_codigo_seleciona_as_mesmas_linhas(monkeypatch, municipio):
    case = next(case for case in MANIFEST["cases"] if case["id"] == "municipio_inicio_2022_semanal")
    install_inputs(monkeypatch)
    query = {**case["query"], "municipio": municipio}
    frame = await api.precos_diesel(**query)
    assert_expected(frame, case)


def test_nome_ibge_com_apostrofo_casa_com_o_nome_da_planilha():
    raw = (GOLDEN / "request_05.xlsx").read_bytes()
    alvo = api.normalize_price_query(
        "RS", 4317103, "DIESEL S10", None, None, "semanal", "municipio"
    )
    frame = parser.parse_precos(raw, uf=alvo["uf"], municipio=alvo["municipio"], nivel="municipio")
    assert alvo["municipio"] == "Sant'Ana do Livramento"
    assert frame["municipio"].unique().tolist() == ["SANTANA DO LIVRAMENTO"]


@pytest.mark.parametrize(
    ("municipio", "uf", "trecho"),
    [("Cuiab", "MT", "Cuiabá/MT (5103403)"), ("Bom Jesus", None, "informe a uf")],
)
def test_pedaco_ou_nome_ambiguo_recusado_antes_da_rede(municipio, uf, trecho):
    with pytest.raises(InvalidParameterError, match=re.escape(trecho)):
        api.normalize_price_query(uf, municipio, "DIESEL S10", None, None, "semanal", "municipio")
