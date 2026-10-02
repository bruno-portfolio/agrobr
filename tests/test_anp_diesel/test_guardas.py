from __future__ import annotations

import hashlib
import json
from datetime import date
from pathlib import Path
from unittest.mock import AsyncMock

import httpx
import openpyxl
import pandas as pd
import pytest

from agrobr import constants, contracts
from agrobr.alt.anp_diesel import _catalog, api, client, models, parser
from agrobr.exceptions import InvalidParameterError, ParseError, SourceUnavailableError
from scripts import reconciliar_anp_precos as reconciliacao
from tests.helpers import levanta_exatamente, sem_excecao
from tests.test_anp_diesel import test_reconciliacao_r11 as r11
from tests.test_anp_diesel.test_catalog import _html
from tests.test_anp_diesel.test_parser import _make_precos_xlsx, _make_vendas_csv

LINHA_UF = {
    "ESTADO - SIGLA": "SP",
    "MUNICÍPIO": "SAO PAULO",
    "PRODUTO": "DIESEL S10",
    "DATA INICIAL": "01/01/2024",
    "DATA FINAL": "07/01/2024",
    "PREÇO MÉDIO REVENDA": "6.45",
    "PREÇO MÉDIO DISTRIBUIÇÃO": "5.80",
    "NÚMERO DE POSTOS PESQUISADOS": "150",
}
VENDA = {
    "ANO": "2024",
    "MES": "JAN",
    "GRANDE REGIAO": "REGIAO SUDESTE",
    "UNIDADE DA FEDERACAO": "SAO PAULO",
    "PRODUTO": "OLEO DIESEL",
    "VENDAS": "800000",
}


@pytest.mark.parametrize(
    ("caminho", "causa"),
    [
        ("semanal/semanal-municipios-2026-2025.xlsx", "Período municipal invertido"),
        ("mensal/semanal-municipios-2027.xlsx", "Planilha fora do catálogo semanal oficial"),
    ],
    ids=["invertida", "fora_da_pasta"],
)
def test_catalogo_recusa_planilha_municipal_invalida(caminho, causa):
    html = f'<a href="{models.SHLP_BASE}/{caminho}">planilha</a>'
    with levanta_exatamente(ParseError, "Planilha municipal inválida no catálogo") as erro:
        _catalog.parse_municipal_catalog(html)
    assert causa in str(erro.value.__cause__)


def servir_catalogo(monkeypatch: pytest.MonkeyPatch, status: int) -> list[str]:
    pedidos: list[str] = []
    original = httpx.AsyncClient

    def responder(request: httpx.Request) -> httpx.Response:
        pedidos.append(str(request.url))
        return httpx.Response(status, text=_html())

    monkeypatch.setattr(
        httpx,
        "AsyncClient",
        lambda *args, **kwargs: original(
            *args, **{**kwargs, "transport": httpx.MockTransport(responder)}
        ),
    )
    return pedidos


async def test_catalogo_municipal_baixado_da_pagina_oficial(monkeypatch):
    pedidos = servir_catalogo(monkeypatch, 200)
    with sem_excecao():
        catalogo = await client.fetch_precos_catalog()
    assert catalogo == models.PRECOS_MUNICIPIOS_URLS
    assert pedidos == [constants.URLS[constants.Fonte.ANP_DIESEL]["precos_catalogo"]]


async def test_catalogo_municipal_indisponivel_fica_visivel(monkeypatch):
    servir_catalogo(monkeypatch, 404)
    with levanta_exatamente(SourceUnavailableError, "HTTP 404: o recurso não existe na URL"):
        await client.fetch_precos_catalog()


@pytest.mark.parametrize("periodo", [{"inicio": 2024}, {"fim": 2024}], ids=["inicio", "fim"])
async def test_vendas_recusa_periodo_que_nao_e_data(periodo, monkeypatch):
    monkeypatch.setattr(
        client, "fetch_vendas_m3", AsyncMock(side_effect=AssertionError("rede antes da validação"))
    )
    with levanta_exatamente(
        InvalidParameterError, "inicio e fim devem ser datas ou strings YYYY-MM-DD"
    ):
        await api.vendas_diesel(**periodo)


def test_vendas_recusa_csv_malformado():
    with levanta_exatamente(ParseError, "Erro ao ler CSV de vendas"):
        parser.parse_vendas(b"ANO;MES;VENDAS\n2024;JAN;1\n2024;JAN;1;2;3\n")


def test_vendas_recusa_uf_publicada_invalida():
    with levanta_exatamente(ParseError, "Venda inválida no registro 1") as erro:
        parser.parse_vendas(_make_vendas_csv([{**VENDA, "UNIDADE DA FEDERACAO": "ATLANTIDA"}]))
    assert "UF publicada inválida" in str(erro.value)


def test_vendas_filtro_sem_linha_da_uf():
    with levanta_exatamente(ParseError, "Nenhuma venda extraida do CSV"):
        parser.parse_vendas(_make_vendas_csv([VENDA]), uf="AC")


@pytest.mark.parametrize(
    ("linha", "colunas", "motivo"),
    [
        (
            {k: v for k, v in LINHA_UF.items() if k != "ESTADO - SIGLA"},
            [k for k in LINHA_UF if k != "ESTADO - SIGLA"],
            "XLSX municipal sem coluna de UF",
        ),
        ({**LINHA_UF, "DATA INICIAL": "01/01/2300"}, list(LINHA_UF), "Data fora do intervalo"),
        ({**LINHA_UF, "ESTADO - SIGLA": ""}, list(LINHA_UF), "Nivel estadual/municipal exige UF"),
        ({**LINHA_UF, "ESTADO - SIGLA": "XX"}, list(LINHA_UF), "UF publicada invalida"),
        ({**LINHA_UF, "DATA INICIAL": 45292}, list(LINHA_UF), "Data semanal invalida"),
    ],
    ids=["sem_coluna_de_uf", "data_2300", "uf_vazia", "uf_invalida", "data_numerica"],
)
def test_precos_recusa_planilha_incoerente(linha, colunas, motivo):
    with levanta_exatamente(ParseError, "Preco semanal invalido") as erro:
        parser.parse_precos(_make_precos_xlsx([linha], columns=colunas))
    assert motivo in str(erro.value)


@pytest.mark.parametrize(
    ("linha", "coluna", "valor", "erro"),
    [
        (0, "produto", "GASOLINA", "produto: valores fora do domínio ['DIESEL', 'DIESEL S10']"),
        (0, "n_postos", pd.NA, "n_postos é contagem publicada somente na frequência semanal"),
        (
            0,
            "n_postos_media",
            80.0,
            "n_postos_media é média de contagens somente na frequência mensal",
        ),
        (0, "n_semanas", 2, "n_semanas deve ser 1 nos registros semanais"),
        (
            0,
            "data",
            pd.Timestamp("2024-01-02"),
            "Período inválido ou início semanal diferente de data",
        ),
        (
            5,
            "data",
            pd.Timestamp("2024-01-08"),
            "data mensal deve ser o primeiro dia do mês de início das semanas",
        ),
        (0, "uf", "", "UF incompatível com o nível territorial"),
        (0, "municipio", "", "Município incompatível com o nível territorial"),
    ],
    ids=[
        "produto",
        "n_postos",
        "n_postos_media",
        "n_semanas",
        "data_semanal",
        "data_mensal",
        "uf",
        "municipio",
    ],
)
def test_contrato_recusa_saida_incoerente(linha, coluna, valor, erro):
    semanal = parser.parse_precos(_make_precos_xlsx())
    saida = pd.concat([semanal, parser.agregar_mensal(semanal)], ignore_index=True)
    contrato = contracts.get_contract("anp_diesel_precos")
    assert contrato.validate(saida) == (True, [])
    saida.loc[linha, coluna] = valor
    assert contrato.validate(saida) == (False, [erro])


@pytest.mark.parametrize(
    ("periodo", "mes"),
    [({"inicio": "2024-02-01"}, "2024-02-01"), ({"fim": "2024-01-01"}, "2024-01-01")],
    ids=["inicio", "fim"],
)
async def test_vendas_filtradas_por_periodo_saem_reindexadas(periodo, mes, monkeypatch):
    monkeypatch.setattr(client, "fetch_vendas_m3", AsyncMock(return_value=_make_vendas_csv()))
    with sem_excecao():
        vendas = await api.vendas_diesel(**periodo)
    assert vendas.index.equals(pd.RangeIndex(len(vendas)))
    assert sorted(vendas["uf"]) == ["MT", "SP"]
    assert set(vendas["data"]) == {pd.Timestamp(mes)}


def test_vendas_sem_diesel_diz_o_motivo():
    with levanta_exatamente(ParseError, "Nenhum registro de diesel encontrado em vendas"):
        parser.parse_vendas(_make_vendas_csv([{**VENDA, "PRODUTO": "GASOLINA C"}]))


async def test_ano_fora_do_catalogo_recusado_sem_baixar_o_catalogo(monkeypatch):
    monkeypatch.setattr(api.time_utils, "hoje", lambda: date(2027, 2, 1))
    monkeypatch.setattr(
        client, "fetch_precos_catalog", AsyncMock(side_effect=AssertionError("catálogo baixado"))
    )
    with levanta_exatamente(
        InvalidParameterError, "2030-01-01 fora do catálogo municipal da ANP: 2022 a 2027"
    ):
        await api.precos_diesel(nivel="municipio", fim="2030-01-01")


def test_vendas_cabecalho_com_espacos():
    limpo = parser.parse_vendas(_make_vendas_csv([VENDA]))
    with sem_excecao():
        lido = parser.parse_vendas(
            _make_vendas_csv([{f" {chave} ": valor for chave, valor in VENDA.items()}])
        )
    assert lido.equals(limpo)


def capturas(destino: Path, trocas: dict[str, Path], adulterado: str | None = None) -> Path:
    recibos = []
    corpos = [
        (
            recurso["original"],
            recurso["derived_file"],
            trocas.get(recurso["derived_file"], r11.GOLDEN / recurso["derived_file"]),
        )
        for recurso in r11.MANIFEST["resources"]
    ]
    catalogo = r11.MANIFEST["catalog"]
    corpos.append((catalogo, catalogo["golden_file"], r11.GOLDEN / catalogo["golden_file"]))
    for original, nome, origem in corpos:
        corpo = origem.read_bytes()
        (destino / nome).write_bytes(corpo)
        recibos.append(
            {
                "requested_url": original["requested_url"],
                "file": nome,
                "sha256": "0" * 64 if nome == adulterado else hashlib.sha256(corpo).hexdigest(),
                "fetched_at": original["fetched_at"],
            }
        )
    (destino / "receipts.json").write_text(json.dumps(recibos), encoding="utf-8")
    return destino


def test_n1_anp_run_sobre_capturas_conhecidas(tmp_path):
    with sem_excecao():
        codigo = reconciliacao.run(capturas(tmp_path, {}), tmp_path / "relatorio.json")
    relatorio = json.loads((tmp_path / "relatorio.json").read_text(encoding="utf-8"))
    assert codigo == 0
    assert [item["status"] for item in relatorio["checks"]] == ["ok"] * 6
    assert (
        relatorio["mode"] == "captured HTTP bodies; capture times retained; no new network request"
    )


def test_n1_anp_run_acusa_planilha_divergente(tmp_path):
    recurso = r11.MANIFEST["resources"][0]
    mutada = r11.mutated_workbook(recurso, "unit", tmp_path / "mutada.xlsx")
    destino = tmp_path / "capturas"
    destino.mkdir()
    with sem_excecao():
        codigo = reconciliacao.run(
            capturas(destino, {recurso["derived_file"]: mutada}), tmp_path / "relatorio.json"
        )
    relatorio = json.loads((tmp_path / "relatorio.json").read_text(encoding="utf-8"))
    assert codigo == 1
    assert [item["status"] for item in relatorio["checks"]] == ["mismatch"] + ["ok"] * 5
    assert relatorio["checks"][0]["url"] == recurso["original"]["requested_url"]


@pytest.mark.parametrize(
    "adulterado", ["request_01.xlsx", "catalog.html"], ids=["planilha", "catalogo"]
)
def test_n1_anp_run_recusa_captura_adulterada(adulterado, tmp_path):
    with levanta_exatamente(AssertionError):
        reconciliacao.run(capturas(tmp_path, {}, adulterado), tmp_path / "relatorio.json")


def planilha_derivada(recurso: dict, mutacao: str, destino: Path) -> bytes:
    livro = openpyxl.load_workbook(r11.GOLDEN / recurso["derived_file"])
    folha = livro.worksheets[0]
    cabecalho = recurso["layout"]["header_row"]
    colunas = {celula.value: celula.column for celula in folha[cabecalho] if celula.value}
    if mutacao == "cabecalho_deslocado":
        folha.insert_rows(1)
    elif mutacao == "sem_s10":
        for linha in range(cabecalho + 1, folha.max_row + 1):
            if folha.cell(linha, colunas["PRODUTO"]).value == "OLEO DIESEL S10":
                folha.cell(linha, colunas["PRODUTO"]).value = "OLEO DIESEL"
    else:
        folha.cell(folha.max_row, colunas["ESTADO"]).value = "ESTADO NOVO"
    livro.save(destino)
    return destino.read_bytes()


@pytest.mark.parametrize(
    ("mutacao", "indice", "problema"),
    [
        ("cabecalho_deslocado", 0, "posição do cabeçalho mudou ou cabeçalho ausente"),
        ("sem_s10", 0, "família de diesel ausente"),
        ("uf_nova", 1, "UF publicada sem decisão"),
    ],
    ids=["cabecalho_deslocado", "sem_s10", "uf_nova"],
)
def test_n1_anp_estrutura_acusa_deriva(mutacao, indice, problema, tmp_path):
    recurso = r11.MANIFEST["resources"][indice]
    resultado = reconciliacao.compare_workbook(
        planilha_derivada(recurso, mutacao, tmp_path / "derivada.xlsx"), recurso
    )
    assert resultado["status"] == "mismatch"
    assert problema in resultado["problems"]
