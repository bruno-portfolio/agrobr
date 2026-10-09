from __future__ import annotations

import io
import zipfile
from pathlib import Path
from types import SimpleNamespace

import httpx
import openpyxl
import pydantic
import pytest

from agrobr import conab, constants
from agrobr.conab import client
from agrobr.conab._custo_producao import (
    _context,
    _parse,
    _sociobio_context,
    _sociobio_parse,
    _workbook,
    models,
)
from agrobr.conab.ceasa import parser as ceasa_parser
from agrobr.conab.parsers import v1
from agrobr.conab.progresso import parser as progresso_parser
from agrobr.exceptions import InvalidParameterError, ParseError
from tests import helpers

GOLDEN = Path(__file__).parents[1] / "golden_data" / "conab"


def contexto_custo():
    book = _workbook.Workbook((GOLDEN / "custos_20260908" / "soja.xls").read_bytes())
    try:
        sheet = book.read("Barreiras-BA-2025")
        resource = models.RecursoCusto(
            cultura="soja",
            planilha="soja.xls",
            titulo="Soja",
            pagina_url="https://www.gov.br/conab/soja.xls/view",
        )
        return _context.context(sheet, resource, 0)
    finally:
        book.close()


def contexto_sociobio():
    book = _workbook.Workbook((helpers.SOCIOBIO_GOLDEN / "acai.xlsx").read_bytes())
    try:
        sheet = book.read("Codajás-AM-2008")
        resource = models.RecursoCusto(
            cultura="acai",
            planilha="acai.xlsx",
            titulo="Açaí",
            pagina_url="https://www.gov.br/conab/acai.xlsx/view",
        )
        return _sociobio_context.context(sheet, resource, 0)
    finally:
        book.close()


@pytest.mark.parametrize(
    "rows,formats,merges,message",
    [
        ([["Discriminação", "Moeda"], ["1 - Sementes", 1]], {}, {}, "Cabeçalho/unidades"),
        (
            [["Discriminação", "R$/ha", "%"], ["1 - Sementes", 1, 0.1]],
            {(1, 2): "0%;0.00"},
            {},
            "Formato numérico mistura seções com e sem %",
        ),
        ([["Discriminação", "R$/ha"], ["1 - Sementes", "inf"]], {}, {}, "Medida não finita"),
        (
            [["Discriminação", ""], ["", "R$/ha"], ["1 - Sementes", 1]],
            {},
            {(0, 0): (0, 2)},
            "Cabeçalhos sobrepostos",
        ),
        (
            [["Discriminação", "R$/ha", "Coluna desconhecida"], ["1 - Sementes", 1, ""]],
            {},
            {},
            "Cabeçalho não reconhecido",
        ),
        (
            [["Discriminação", "R$/ha", ""], ["", "", "Nota solta"], ["1 - Sementes", 1, ""]],
            {},
            {},
            "Conteúdo fora das colunas",
        ),
        (
            [["Discriminação", "R$/ha"], ["", 1], ["1 - Sementes", 1]],
            {},
            {},
            "Medida sem descrição",
        ),
        (
            [["Discriminação", "R$/ha", ""], ["1 - Sementes", 1, "Nota solta"]],
            {},
            {},
            "Colunas não reconhecidas com conteúdo",
        ),
        ([["Discriminação", "R$/ha"], ["Observação", ""]], {}, {}, "Nenhuma observação numérica"),
    ],
    ids=["1018", "1019", "1021", "1023", "1024", "1027", "1028", "1030", "1031"],
)
def test_custos_recusa_layout_ou_medida_inconsistente(rows, formats, merges, message):
    sheet = _workbook.Aba("Barreiras-BA-2025", rows, formats, merges)
    with helpers.levanta_exatamente(ParseError, message):
        _parse.parse_selected(sheet, contexto_custo())


def test_custos_totais_recusa_total_publicado_duplicado():
    sheet = _workbook.Aba(
        "Barreiras-BA-2025",
        [["Discriminação", "R$/ha"], ["CUSTO TOTAL", 1], ["CUSTO TOTAL", 2]],
        {},
    )
    result = _parse.parse_selected(sheet, contexto_custo())
    with helpers.levanta_exatamente(ParseError, "Total publicado ambíguo: ct_ha"):
        _parse.totals(result)


@pytest.mark.parametrize(
    "rows,formats,merges,message",
    [
        (
            [["Item", "R$/kg", "%"], ["1 - Coleta", 1, 0.1]],
            {},
            {},
            "Cabeçalho de descrição ausente",
        ),
        (
            [["Discriminação", "R$/kg", "%", "%"], ["1 - Coleta", 1, 0.1, 0.1]],
            {},
            {},
            "Cabeçalho duplicado",
        ),
        (
            [["Discriminação", "R$/kg", "%"], ["1 - Coleta", 1, 0.1]],
            {(1, 2): "0%;0.00"},
            {},
            "Formato numérico mistura seções com e sem %",
        ),
        (
            [["Discriminação", "R$/kg", "%"], ["1 - Coleta", 1, 1e308]],
            {(1, 2): "0%"},
            {},
            "Medida não finita em R2C3",
        ),
        (
            [["Discriminação", "R$/kg", "", "%"], ["1 - Coleta", 1, 2, 0.1]],
            {},
            {(0, 1): (1, 3)},
            "Mais de um valor na coluna valor",
        ),
        (
            [["Discriminação", "R$/kg", "%"], ["Categoria desconhecida", 1, 0.1]],
            {},
            {},
            "Linha com medida sem tipo reconhecido",
        ),
        (
            [["Discriminação", "R$/kg", "%"], ["Observação", "", ""]],
            {},
            {},
            "Aba sem observações reconhecidas",
        ),
    ],
    ids=["1058", "1059", "1063", "1064", "1065", "1067", "1070"],
)
def test_sociobio_recusa_layout_ou_medida_inconsistente(rows, formats, merges, message):
    sheet = _workbook.Aba("Codajás-AM-2008", rows, formats, merges)
    with helpers.levanta_exatamente(ParseError, message):
        _sociobio_parse.parse_selected(sheet, contexto_sociobio())


async def test_sociobio_recusa_seletor_vazio_antes_da_rede():
    with helpers.levanta_exatamente(InvalidParameterError, "Seletor vazio"):
        await conab.custo_sociobiodiversidade("acai", local=" ")


def test_ceasa_recusa_metadata_sem_coluna_de_preco():
    with helpers.levanta_exatamente(ParseError, "Cabeçalhos de preço"):
        ceasa_parser.parse_precos(
            {"resultset": [["Tomate (kg)", 5]], "metadata": [{"colName": "ProdutoUnid"}]}
        )


async def test_catalogo_recusa_levantamento_fora_do_intervalo(monkeypatch):
    html = (
        '<a href="https://www.gov.br/conab/13o-levantamento-safra-2025-26/tabela.xlsx">Tabela</a>'
    ).ljust(constants.MIN_HTML_PAGE_SIZE)
    chamadas = []

    def responder(request):
        chamadas.append(request)
        return httpx.Response(200, text=html, request=request)

    namespace = SimpleNamespace(**vars(httpx))
    namespace.AsyncClient = lambda **kwargs: httpx.AsyncClient(
        transport=httpx.MockTransport(responder), **kwargs
    )
    monkeypatch.setattr(client, "httpx", namespace)

    with helpers.levanta_exatamente(ParseError) as capturado:
        await conab.levantamentos()

    assert capturado.value.source == "conab"
    causa = capturado.value.__cause__
    assert isinstance(causa, pydantic.ValidationError)
    assert [(erro["loc"], erro["input"], erro.get("ctx")) for erro in causa.errors()] == [
        (("levantamento",), 13, {"le": 12})
    ]
    assert len(chamadas) == 1
    assert str(chamadas[0].url) == constants.URLS[constants.Fonte.CONAB]["boletim_graos"]


def planilha(rows, name):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = name
    for row in rows:
        sheet.append(row)
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    stream.seek(0)
    return stream


@pytest.mark.parametrize(
    "method,name,rows,message",
    [
        ("safra", "Soja", [["Sem cabeçalho", 1]], "Não encontrou header"),
        ("suprimento", "Suprimento", [["Sem cabeçalho", 1]], "Não encontrou header"),
        (
            "suprimento_soja",
            "Suprimento - Soja",
            [["PRODUTO", ""], ["", "sem safra"]],
            "Não encontrou safras",
        ),
        (
            "brasil",
            "Brasil - Total por Produto",
            [["PRODUTO", "ÁREA"]],
            "Não foi possível detectar colunas de safra",
        ),
        (
            "brasil",
            "Brasil - Total por Produto",
            [["PRODUTO", "ÁREA", "ÁREA"], ["", "2024/25", "2025/26"]],
            "Métrica de safra ambígua",
        ),
        (
            "brasil",
            "Brasil - Total por Produto",
            [["PRODUTO", ""], ["", "2024/25"]],
            "Safra sem métrica",
        ),
    ],
    ids=["1175", "1179", "1183", "1188", "1189", "1191"],
)
def test_safras_recusa_cabecalho_inconsistente(method, name, rows, message):
    stream = planilha(rows, name)
    parser = v1.ConabParserV1()
    with helpers.levanta_exatamente(ParseError, message):
        if method == "safra":
            parser.parse_safra_produto(stream, "soja")
        elif method == "brasil":
            parser.parse_brasil_total(stream)
        else:
            parser.parse_suprimento(stream, "soja" if method == "suprimento_soja" else "milho")


def test_suprimento_recusa_revisao_duplicada():
    rows = [
        [
            "PRODUTO",
            "SAFRA",
            "",
            "ESTOQUE INICIAL",
            "PRODUÇÃO",
            "IMPORTAÇÃO",
            "SUPRIMENTO",
            "CONSUMO",
            "EXPORTAÇÃO",
            "DEMANDA TOTAL",
            "ESTOQUE FINAL",
        ],
        ["MILHO", "2025/26", "jul/26", 10, 100, 1, 111, 50, 40, 90, 21],
        ["MILHO", "2025/26", "jul/26", 10, 120, 1, 131, 50, 40, 90, 41],
    ]
    with helpers.levanta_exatamente(ParseError, "Revisão ambígua"):
        v1.ConabParserV1().parse_suprimento(planilha(rows, "Suprimento"), "milho")


def test_progresso_avisa_estado_desconhecido():
    content = helpers.progresso_xlsx_cells({"B12": "Atlantida"})
    with helpers.capturar_logs() as logs:
        frame = progresso_parser.parse_progresso_xlsx(content)
    assert any(
        log.get("event") == "conab_progresso_estado_desconhecido"
        and log.get("estado") == "Atlantida"
        for log in logs
    )
    assert "Atlantida" not in frame["uf"].tolist()


def test_progresso_recusa_celula_numerica_ilegivel():
    original = planilha([[1]], "Progresso")
    output = io.BytesIO()
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(output, "w") as destination:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                assert b"<v>1</v>" in content
                content = content.replace(b"<v>1</v>", b"<v>ilegivel</v>")
            destination.writestr(entry, content)
    with helpers.levanta_exatamente(ParseError, "Erro ao ler sheet 'Progresso'"):
        progresso_parser.parse_progresso_xlsx(output.getvalue())
