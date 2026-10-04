from __future__ import annotations

import asyncio
import hashlib
import json
import warnings
from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Any
from unittest.mock import AsyncMock

import bs4
import httpx
import openpyxl
import pandas as pd
import pytest
import xlrd

from agrobr import constants, datasets
from agrobr.conab import api, client
from agrobr.conab._serie_historica import api as serie_api
from agrobr.conab._serie_historica import client as serie_client
from agrobr.conab._serie_historica import parser as serie_parser
from agrobr.conab.parsers import v1
from agrobr.datasets.producao_anual import _fetch_conab
from agrobr.exceptions import SourceUnavailableError
from tests import helpers

GOLDEN = Path(__file__).parents[1] / "golden_data"
R3 = GOLDEN / "reconciliacao_r3_20260918"
SET_2025 = GOLDEN / "conab/levantamento_12_2024_25_20260922"
SET_2025_URL = json.loads((SET_2025 / "PROVENANCE.json").read_text(encoding="utf-8"))["files"][0][
    "url"
]
SET_2026_URL = (
    "https://www.gov.br/conab/pt-br/atuacao/informacoes-agropecuarias/safras/safra-de-graos/"
    "boletim-da-safra-de-graos/12o-levantamento-safra-2025-26/"
    "site_previsao_de_safra-por_produto-set-2026.xlsx"
)
SET_2022_URL = (
    "https://www.gov.br/conab/pt-br/atuacao/informacoes-agropecuarias/safras/safra-de-graos/"
    "boletim-da-safra-de-graos/12o-levantamento-safra-2021-22/"
    "site-previsao_de_safra-por_produto-set-2022.xlsx"
)
SET_2021_URL = (
    "https://www.gov.br/conab/pt-br/atuacao/informacoes-agropecuarias/safras/safra-de-graos/"
    "boletim-da-safra-de-graos/12o-levantamento-safra-2020-21/"
    "site-previsao-de-safra-por-produto-set-2021.xls"
)
AGO_2020_URL = (
    "https://www.gov.br/conab/pt-br/atuacao/informacoes-agropecuarias/safras/safra-de-graos/"
    "boletim-da-safra-de-graos/11o-levantamento-safra-2019-20/"
    "site-previsao-de-safra-por-produto-ago-2020-copia.xls"
)
INVERNO = ("Aveia", "Canola", "Centeio", "Cevada", "Trigo", "Triticale")
TIPOS = ("cores", "preto", "caupi")
PROXIMA_EDICAO_URL = (
    "https://www.gov.br/conab/pt-br/atuacao/informacoes-agropecuarias/safras/safra-de-graos/"
    "boletim-da-safra-de-graos/1o-levantamento-safra-2026-27/nao-baixada.xlsx"
)
EDICOES = {
    SET_2025_URL: SET_2025 / "site_previsao_de_safra-por_produto-set-2025.xlsx",
    SET_2026_URL: R3 / "7cd4df7946e5c57f.xlsx",
}
SERIES = {
    **{
        case["product"]: helpers.SERIE_HISTORICA_GOLDEN / case["file"]
        for case in helpers.load_serie_historica_manifest()["cases"]
    },
    **{
        produto: GOLDEN / f"conab/serie_historica_zeros_20260923/{produto}.xls"
        for produto in ("amendoim_2", "feijao_3", "mamona")
    },
    **{
        caminho.stem: caminho
        for caminho in (GOLDEN / "conab/serie_historica_20260924").glob("*.xls")
    },
}
ABAS = {
    "algodao_caroco": {
        "area_plantada": "Área",
        "produtividade": "Produtividade Caroço de Algodão",
        "producao": "Produção de Caroço de Algodão",
    },
    "algodao_pluma": {
        "area_plantada": "Área",
        "produtividade": "Produtividade Pluma",
        "producao": "Produção de Pluma",
    },
}
GENERICAS = {"area_plantada": "Área", "produtividade": "Produtividade", "producao": "Produção"}


def _servir(monkeypatch: pytest.MonkeyPatch, fora_do_ar: tuple[str, ...] = ()) -> dict[str, Any]:
    catalogo = bs4.BeautifulSoup((R3 / "dc5524f326c16e26.html").read_bytes(), "lxml")
    for link in catalogo.select("a.proximo, a[rel~=next]"):
        link.decompose()
    series = {
        serie_client.get_xls_url(produto): caminho.read_bytes()
        for produto, caminho in SERIES.items()
        if produto not in fora_do_ar
    }
    estado: dict[str, Any] = {"pedidos": [], "abertos": 0, "pico": 0}

    async def responder(request: httpx.Request) -> httpx.Response:
        estado["pedidos"].append(str(request.url))
        estado["abertos"] += 1
        estado["pico"] = max(estado["pico"], estado["abertos"])
        await asyncio.sleep(0.01)
        estado["abertos"] -= 1
        conteudo = series.get(str(request.url))
        return httpx.Response(200 if conteudo else 404, content=conteudo or b"", request=request)

    async def baixar(url: str) -> bytes:
        return EDICOES[url].read_bytes()

    original = httpx.AsyncClient
    monkeypatch.setattr(client, "fetch_boletim_page", AsyncMock(return_value=str(catalogo)))
    monkeypatch.setattr(client, "_fetch_http", baixar)
    monkeypatch.setattr(
        serie_client.httpx,
        "AsyncClient",
        lambda **kwargs: original(transport=httpx.MockTransport(responder), **kwargs),
    )
    monkeypatch.setattr(api, "_hoje", lambda: date(2026, 9, 25))
    return estado


def _rotulo(valor: Any) -> str:
    return str(int(valor)) if isinstance(valor, float) else str(valor).strip()


def _serie(produto: str, periodo: str, rotulo: str) -> dict[str, float]:
    livro = xlrd.open_workbook(str(SERIES[produto]))
    resultado = {}
    for campo, nome in ABAS.get(produto, GENERICAS).items():
        aba = livro.sheet_by_name(nome)
        coluna = [_rotulo(valor) for valor in aba.row_values(5)].index(periodo)
        linha = [_rotulo(valor).upper() for valor in aba.col_values(0)].index(rotulo)
        valor = aba.cell_value(linha, coluna)
        resultado[campo] = valor if isinstance(valor, float) else None
    return resultado


def _legenda(produto: str) -> str:
    aba = xlrd.open_workbook(str(SERIES[produto])).sheet_by_index(0)
    texto = next(str(valor) for valor in aba.col_values(0) if "Estimativa em" in str(valor))
    return texto.split("Estimativa em ")[1].rstrip(".").replace(" de ", "/").strip().lower()


def _nulos(valores: dict[str, Any]) -> dict[str, Any]:
    return {campo: None if pd.isna(valor) else valor for campo, valor in valores.items()}


async def test_safra_antiga_sai_da_serie_historica_mais_recente(monkeypatch: pytest.MonkeyPatch):
    _servir(monkeypatch)
    frame, meta = await api.safras("soja", safra="2022/23", return_meta=True)
    trigo = await api.safras("trigo", safra="2021/22", uf="PR")
    estimativa = await datasets.estimativa_safra("soja", safra="2022/23", uf="MT", fonte="conab")
    anual, _ = await _fetch_conab("soja", ano=2023, nivel="uf", uf="MT")

    publicado = {uf: _serie("soja", "2022/23", uf) for uf in frame["uf"]}
    assert len(publicado) == 27
    assert {
        linha.uf: _nulos(
            {
                "area_plantada": linha.area_plantada,
                "produtividade": linha.produtividade,
                "producao": linha.producao,
            }
        )
        for linha in frame.itertuples()
    } == publicado
    assert publicado["MT"]["producao"] == 46905.8
    assert frame["producao"].sum() == pytest.approx(_serie("soja", "2022/23", "BRASIL")["producao"])
    assert frame["levantamento"].isna().all() and frame["data_publicacao"].isna().all()
    assert len(trigo) == 1
    assert trigo[["safra", "uf", "producao"]].to_dict("records") == [
        {"safra": "2021/22", "uf": "PR", "producao": _serie("trigo", "2022", "PR")["producao"]}
    ]
    assert estimativa["producao"].tolist() == [46905.8]
    assert anual["producao"].tolist() == [46905.8 * 1000]
    assert meta.cache_expires_at is None
    assert meta.source_details["publicacao"] == {
        "origem": "serie_historica",
        "series": [
            {
                "produto": "soja",
                "url": serie_client.get_xls_url("soja"),
                "sha256": hashlib.sha256(SERIES["soja"].read_bytes()).hexdigest(),
                "referencia": "setembro/2026",
            }
        ],
    }


async def test_serie_nao_posterior_ao_boletim_fica_o_boletim_com_aviso(
    monkeypatch: pytest.MonkeyPatch,
):
    _servir(monkeypatch)
    edicoes = [
        {
            "url": SET_2026_URL,
            "levantamento": 12,
            "safra": "2025/26",
            "ano_inicio": 2025,
            "ano_fim": 26,
            "data_publicacao": date(2026, 9, 15),
        },
        {
            "url": SET_2025_URL,
            "levantamento": 12,
            "safra": "2024/25",
            "ano_inicio": 2024,
            "ano_fim": 25,
            "data_publicacao": date(2026, 10, 1),
        },
    ]
    monkeypatch.setattr(client, "list_levantamentos", AsyncMock(return_value=edicoes))
    with warnings.catch_warnings(record=True) as avisos:
        warnings.simplefilter("always")
        frame = await api.safras("soja", safra="2023/24", uf="MT")

    livro = openpyxl.load_workbook(EDICOES[SET_2025_URL], read_only=True, data_only=True)
    linhas = list(livro["Soja"].iter_rows(values_only=True))
    colunas = [indice for indice, celula in enumerate(linhas[5]) if celula == "Safra 23/24"]
    mt = next(valores for valores in linhas if valores[0] == "MT")
    assert frame[["area_plantada", "produtividade", "producao", "levantamento"]].to_dict(
        "records"
    ) == [
        {
            "area_plantada": float(mt[colunas[0]]),
            "produtividade": float(mt[colunas[1]]),
            "producao": float(mt[colunas[2]]),
            "levantamento": 12,
        }
    ]
    assert [str(aviso.message) for aviso in avisos if "série histórica" in str(aviso.message)] == [
        "CONAB: a série histórica de soja (setembro/2026) não é posterior ao 12º levantamento de "
        "2024/25 (2026-10-01); a safra 2023/24 sai dessa publicação do boletim"
    ]


async def test_brasil_total_de_safra_antiga_vem_das_series(monkeypatch: pytest.MonkeyPatch):
    estado = _servir(monkeypatch)
    frame, meta = await api.brasil_total(safra="2022/23", return_meta=True)

    esperado = []
    for rotulo, grupo, produto, _ in constants.CONAB_BRASIL_TOTAL_SERIES:
        periodo = "2023" if grupo == constants.CONAB_INVERNO else "2022/23"
        esperado.append((rotulo, grupo, _serie(produto, periodo, "BRASIL")))
    produtos = frame[~frame["rotulo"].isin(["SUBTOTAL", "BRASIL (2)"])]
    assert [
        (
            linha.rotulo,
            None if pd.isna(linha.grupo) else linha.grupo,
            {
                "area_plantada": float(linha.area_plantada),
                "produtividade": float(linha.produtividade),
                "producao": float(linha.producao),
            },
        )
        for linha in produtos.itertuples()
    ] == esperado

    suprimento = openpyxl.load_workbook(EDICOES[SET_2026_URL], read_only=True, data_only=True)
    linhas = list(suprimento["Suprimento - Soja"].iter_rows(values_only=True))
    coluna = next(
        i for valores in linhas for i, celula in enumerate(valores) if celula == "2022/23"
    )
    producao = next(v for v in linhas if str(v[0]).strip().endswith("Produção"))[coluna]
    assert producao == 159154.3
    assert float(frame.loc[frame["rotulo"] == "SOJA", "producao"].iloc[0]) == producao

    def soma(linhas: pd.DataFrame, campo: str) -> float:
        return float(sum(linhas[campo]))

    verao = [
        rotulo for rotulo, _, _, secao in constants.CONAB_BRASIL_TOTAL_SERIES if secao == "verao"
    ]
    subtotais = frame[frame["rotulo"].isin(["SUBTOTAL", "BRASIL (2)"])]
    assert subtotais["rotulo"].tolist() == ["SUBTOTAL", "SUBTOTAL", "BRASIL (2)"]
    assert frame.index[frame["rotulo"] == "SUBTOTAL"].tolist() == [30, 37]
    for campo in ("area_plantada", "producao"):
        verao_frame = produtos[produtos["rotulo"].isin(verao) & produtos["grupo"].isna()]
        inverno_frame = produtos[produtos["grupo"] == constants.CONAB_INVERNO]
        assert float(subtotais[campo].iloc[0]) == pytest.approx(soma(verao_frame, campo))
        assert float(subtotais[campo].iloc[1]) == pytest.approx(soma(inverno_frame, campo))
        assert float(subtotais[campo].iloc[2]) == pytest.approx(
            soma(verao_frame, campo) + soma(inverno_frame, campo)
        )
    assert float(subtotais["produtividade"].iloc[2]) == pytest.approx(
        float(subtotais["producao"].iloc[2]) * 1000 / float(subtotais["area_plantada"].iloc[2])
    )
    urls = {
        serie_client.get_xls_url(produto)
        for _, _, produto, _ in constants.CONAB_BRASIL_TOTAL_SERIES
    }
    assert sorted(estado["pedidos"]) == sorted(urls)
    assert estado["pico"] <= 4
    publicacao = meta.source_details["publicacao"]
    assert publicacao["origem"] == "serie_historica"
    assert {serie["url"] for serie in publicacao["series"]} == urls
    assert {serie["url"]: serie["referencia"] for serie in publicacao["series"]} == {
        serie_client.get_xls_url(produto): _legenda(produto)
        for _, _, produto, _ in constants.CONAB_BRASIL_TOTAL_SERIES
    }
    assert _legenda("feijao_cores_2") == "agosto/2026"
    assert {serie["url"]: serie["sha256"] for serie in publicacao["series"]} == {
        serie_client.get_xls_url(produto): hashlib.sha256(caminho.read_bytes()).hexdigest()
        for produto, caminho in SERIES.items()
        if serie_client.get_xls_url(produto) in urls
    }


async def test_brasil_total_de_safra_antiga_falha_inteiro_sem_uma_serie(
    monkeypatch: pytest.MonkeyPatch,
):
    _servir(monkeypatch, fora_do_ar=("milho_2",))
    with pytest.raises(SourceUnavailableError) as erro:
        await api.brasil_total(safra="2022/23")
    assert "Milho 2ª Safra (MILHO TOTAL)" in erro.value.last_error
    assert "use levantamento= para a edição original" in erro.value.last_error


@pytest.mark.parametrize("arquivo", ["676a715368f4fb34.xlsx", "7cd4df7946e5c57f.xlsx"])
def test_subtotal_e_brasil_sao_a_soma_das_linhas_de_produto(arquivo: str):
    aba = pd.read_excel(
        R3 / arquivo, sheet_name="Brasil - Total por Produto", header=None, engine="calamine"
    )
    secoes = {rotulo: secao for rotulo, _, _, secao in constants.CONAB_BRASIL_TOTAL_SERIES}
    rotulos = aba[0].astype(str).str.strip()
    for area, producao in ((1, 7), (2, 8)):
        somas: dict[str | None, list[float]] = {"verao": [0.0, 0.0], "inverno": [0.0, 0.0]}
        publicados = []
        grupo_inverno = False
        for indice, rotulo in rotulos.items():
            if rotulo == constants.CONAB_INVERNO:
                grupo_inverno = True
            elif rotulo in ("SUBTOTAL", "BRASIL (2)"):
                publicados.append((rotulo, aba.iloc[indice, area], aba.iloc[indice, producao]))
            elif secoes.get(rotulo) == ("inverno" if grupo_inverno else "verao"):
                somas[secoes[rotulo]][0] += aba.iloc[indice, area]
                somas[secoes[rotulo]][1] += aba.iloc[indice, producao]
        calculados = [
            ("SUBTOTAL", *somas["verao"]),
            ("SUBTOTAL", *somas["inverno"]),
            ("BRASIL (2)", *(v + i for v, i in zip(somas["verao"], somas["inverno"], strict=True))),
        ]
        for (rotulo, area_pub, prod_pub), (_, area_calc, prod_calc) in zip(
            publicados, calculados, strict=True
        ):
            assert area_calc == pytest.approx(area_pub, abs=0.05), (arquivo, rotulo)
            assert prod_calc == pytest.approx(prod_pub, abs=0.05), (arquivo, rotulo)


async def test_safra_antiga_nao_repete_o_boletim_congelado(monkeypatch: pytest.MonkeyPatch):
    _servir(monkeypatch)
    monkeypatch.setitem(EDICOES, SET_2022_URL, R3 / "676a715368f4fb34.xlsx")
    edicoes = [
        {
            "url": SET_2026_URL,
            "levantamento": 12,
            "safra": "2025/26",
            "ano_inicio": 2025,
            "ano_fim": 26,
            "data_publicacao": date(2026, 9, 15),
        },
        {
            "url": SET_2022_URL,
            "levantamento": 12,
            "safra": "2021/22",
            "ano_inicio": 2021,
            "ano_fim": 22,
            "data_publicacao": date(2022, 9, 8),
        },
    ]
    monkeypatch.setattr(client, "list_levantamentos", AsyncMock(return_value=edicoes))
    frame = await api.safras("soja", safra="2020/21", uf="MS")

    soja = pd.read_excel(EDICOES[SET_2022_URL], sheet_name="Soja", header=None, engine="calamine")
    coluna = [i for i, celula in enumerate(soja.iloc[5]) if celula == "Safra 20/21"][2]
    congelado = float(soja.loc[soja[0] == "MS", coluna].iloc[0])
    revisado = _serie("soja", "2020/21", "MS")["producao"]
    assert congelado != revisado
    assert frame["producao"].tolist() == [revisado]


def test_mapa_das_linhas_bate_com_a_aba_brasil_na_safra_anterior():
    aba = pd.read_excel(
        R3 / "7cd4df7946e5c57f.xlsx", sheet_name="Brasil - Total por Produto", header=None
    )
    publicado = {}
    titulo = secao = None
    for _, linha in aba.iloc[7:].iterrows():
        rotulo = str(linha[0]).strip()
        if rotulo == constants.CONAB_INVERNO:
            secao = rotulo
            continue
        grupo = secao if rotulo == rotulo.upper() else titulo
        if rotulo == rotulo.upper():
            titulo = rotulo
        publicado[(rotulo, grupo)] = {"area_plantada": linha[1], "producao": linha[7]}
    for rotulo, grupo, produto, _ in constants.CONAB_BRASIL_TOTAL_SERIES:
        periodo = "2025" if grupo == constants.CONAB_INVERNO else "2024/25"
        serie = _serie(produto, periodo, "BRASIL")
        campos = {"area_plantada", "producao"}
        if rotulo == "GERGELIM":
            campos = set()
        elif rotulo == "FEIJÃO 3ª SAFRA":
            campos = {"area_plantada"}
        for campo in campos:
            assert serie[campo] == pytest.approx(publicado[(rotulo, grupo)][campo]), (rotulo, campo)


async def test_balanco_sem_edicao_que_traga_a_safra_falha(monkeypatch: pytest.MonkeyPatch):
    estado = _servir(monkeypatch)
    with pytest.raises(SourceUnavailableError, match="No levantamento found for safra=2030/31"):
        await api.balanco("soja", safra="2030/31")
    assert estado["pedidos"] == []


async def test_safra_fora_do_alcance_da_serie_sai_vazia(monkeypatch: pytest.MonkeyPatch):
    _servir(monkeypatch)
    assert _serie("soja", "1976/77", "BRASIL")["area_plantada"] is not None
    with helpers.sem_excecao():
        frame = await api.safras("soja", safra="1975/76")
    assert frame.empty


async def test_brasil_total_com_serie_defasada_fica_no_boletim(monkeypatch: pytest.MonkeyPatch):
    _servir(monkeypatch)
    monkeypatch.setitem(EDICOES, SET_2022_URL, R3 / "676a715368f4fb34.xlsx")
    edicoes = [
        {
            "url": SET_2026_URL,
            "levantamento": 12,
            "safra": "2025/26",
            "ano_inicio": 2025,
            "ano_fim": 26,
            "data_publicacao": date(2026, 9, 15),
        },
        {
            "url": SET_2022_URL,
            "levantamento": 12,
            "safra": "2021/22",
            "ano_inicio": 2021,
            "ano_fim": 22,
            "data_publicacao": date(2027, 1, 1),
        },
    ]
    monkeypatch.setattr(client, "list_levantamentos", AsyncMock(return_value=edicoes))
    with warnings.catch_warnings(record=True) as avisos:
        warnings.simplefilter("always")
        frame = await api.brasil_total(safra="2020/21")

    aba = pd.read_excel(
        R3 / "676a715368f4fb34.xlsx",
        sheet_name="Brasil - Total por Produto",
        header=None,
        engine="calamine",
    )
    soja = float(aba.loc[aba[0].astype(str).str.strip() == "SOJA", 7].iloc[0])
    assert soja != _serie("soja", "2020/21", "BRASIL")["producao"]
    assert float(frame.loc[frame["rotulo"] == "SOJA", "producao"].iloc[0]) == soja
    assert any("não é posterior ao 12º levantamento de 2021/22" in str(a.message) for a in avisos)


async def test_erro_inesperado_no_download_da_serie_nao_e_engolido(monkeypatch: pytest.MonkeyPatch):
    _servir(monkeypatch)

    async def quebrar(produto: str) -> Any:
        raise RuntimeError(f"falha inesperada em {produto}")

    monkeypatch.setattr(serie_client, "download_xls", quebrar)
    with helpers.levanta_exatamente(RuntimeError, match="falha inesperada"):
        await api.brasil_total(safra="2022/23")


def test_serie_sem_legenda_nao_tem_referencia():
    livro = openpyxl.Workbook()
    livro.active.append(["SOJA - BRASIL"])
    livro.active.append(["Fonte: Conab"])
    saida = BytesIO()
    livro.save(saida)
    assert serie_parser.referencia_publicacao(saida.getvalue()) is None
    assert serie_parser.referencia_publicacao(SERIES["cana"].read_bytes()) == (
        "agosto/2026",
        date(2026, 8, 1),
    )


async def test_balanco_sem_produto_traz_a_soja_da_aba_propria(monkeypatch: pytest.MonkeyPatch):
    _servir(monkeypatch)
    todos = await api.balanco(safra="2025/26")
    soja = await api.balanco("soja", safra="2025/26")

    livro = openpyxl.load_workbook(EDICOES[SET_2026_URL], read_only=True, data_only=True)
    linhas = [
        valores
        for valores in livro["Suprimento - Soja"].iter_rows(values_only=True)
        if any(celula is not None for celula in valores)
    ]
    safras = next(
        valores for indice, valores in enumerate(linhas) if linhas[indice - 1][1] == "SAFRA"
    )
    producao = next(valores for valores in linhas if str(valores[0]).endswith("Produção"))
    publicado = {
        safra: valor for safra, valor in zip(safras[1:], producao[1:], strict=True) if safra
    }
    assert len(publicado) == 6
    assert publicado["2025/26"] == 180406.6

    da_soja = todos[todos["produto"] == "SOJA"].reset_index(drop=True)
    assert dict(zip(da_soja["safra"], da_soja["producao"], strict=True)) == publicado
    pd.testing.assert_frame_equal(da_soja, soja.reset_index(drop=True), check_like=True)
    assert sorted(set(todos["produto"]) - {"SOJA"}) == [
        "ALGODÃO",
        "ARROZ EM CASCA",
        "FEIJÃO",
        "MILHO",
        "TRIGO",
    ]


async def test_balanco_sem_produto_em_edicao_sem_aba_da_soja_le_so_a_aba_longa(
    monkeypatch: pytest.MonkeyPatch,
):
    _servir(monkeypatch)
    monkeypatch.setitem(EDICOES, SET_2021_URL, R3 / "1569ad0cf189f4d5.xls")
    edicao = {
        "url": SET_2021_URL,
        "levantamento": 12,
        "safra": "2020/21",
        "ano_inicio": 2020,
        "ano_fim": 21,
        "data_publicacao": date(2021, 9, 9),
    }
    monkeypatch.setattr(client, "list_levantamentos", AsyncMock(return_value=[edicao]))
    with helpers.sem_excecao():
        frame = await api.balanco(safra="2020/21", levantamento=12)

    livro = xlrd.open_workbook(str(R3 / "1569ad0cf189f4d5.xls"))
    assert "Suprimento - Soja" not in livro.sheet_names()
    rotulos = [str(valor).strip() for valor in livro.sheet_by_name("Suprimento").col_values(0)]
    fim = next(indice for indice, rotulo in enumerate(rotulos) if rotulo.startswith("Nota"))
    produtos = [rotulo for rotulo in rotulos[rotulos.index("PRODUTO") + 1 : fim] if rotulo]
    assert produtos == ["ALGODÃO", "ARROZ EM CASCA", "FEIJÃO", "MILHO", "TRIGO"]
    assert sorted(set(frame["produto"])) == produtos
    milho = frame[(frame["produto"] == "MILHO") & (frame["safra"] == "2020/21")]
    assert milho["producao"].tolist() == [85749.0]


async def test_balanco_incoerente_avisa_sem_mudar_o_dado(monkeypatch: pytest.MonkeyPatch):
    _servir(monkeypatch)
    with warnings.catch_warnings(record=True) as avisos:
        warnings.simplefilter("always")
        arroz = await api.balanco("arroz", safra="2025/26")
        await api.balanco("milho", safra="2025/26")

    livro = openpyxl.load_workbook(EDICOES[SET_2026_URL], read_only=True, data_only=True)
    blocos = []
    produto = None
    for valores in livro["Suprimento"].iter_rows(values_only=True):
        produto = str(valores[0]).strip() if valores[0] else produto
        blocos.append((produto, valores))
    publicado = next(v for p, v in blocos if p == "ARROZ EM CASCA" and v[1] == "2024/25")
    suprimento, consumo, exportacao, estoque_final = (publicado[i] for i in (6, 7, 8, 10))
    assert round(suprimento - consumo - exportacao - estoque_final, 1) == -363.8

    linha = arroz.loc[arroz["safra"] == "2024/25"].iloc[0]
    assert linha["estoque_final"] == pytest.approx(estoque_final)
    assert linha["suprimento"] == pytest.approx(suprimento)
    assert [str(a.message) for a in avisos if "balanço" in str(a.message)] == [
        "CONAB: no balanço de ARROZ EM CASCA 2024/25 (12º levantamento de 2025/26), "
        "suprimento − consumo − exportação − estoque final = -363.8 mil t, e não 0; "
        "o agrobr repassa os números publicados"
    ]


def _aviso_de_soma(rotulo: str, coluna: str, ufs: dict[str, float], brasil: float) -> str:
    soma = sum(ufs.values())
    return (
        f"CONAB: em {rotulo}, a soma das {len(ufs)} UFs em {coluna} dá {soma:.1f}, e o BRASIL "
        f"publicado, {brasil:.1f} (diferença de {soma - brasil:.1f}); o agrobr repassa os "
        "números publicados"
    )


async def test_soma_das_ufs_do_boletim_diferente_do_brasil_avisa_sem_mudar_o_dado(
    monkeypatch: pytest.MonkeyPatch,
):
    _servir(monkeypatch)
    monkeypatch.setitem(EDICOES, SET_2022_URL, R3 / "676a715368f4fb34.xlsx")
    edicao = {
        "url": SET_2022_URL,
        "levantamento": 12,
        "safra": "2021/22",
        "ano_inicio": 2021,
        "ano_fim": 22,
        "data_publicacao": date(2022, 9, 8),
    }
    monkeypatch.setattr(client, "list_levantamentos", AsyncMock(return_value=[edicao]))
    with warnings.catch_warnings(record=True) as avisos:
        warnings.simplefilter("always")
        frames = {
            produto: await api.safras(produto, safra="2020/21")
            for produto in ("milho_2", "aveia", "soja")
        }

    publicados = {}
    for produto, aba, rotulo in (
        ("milho_2", "Milho 2a", "Safra 20/21"),
        ("aveia", "Aveia", "Safra 2021"),
        ("soja", "Soja", "Safra 20/21"),
    ):
        planilha = pd.read_excel(
            EDICOES[SET_2022_URL], sheet_name=aba, header=None, engine="calamine"
        )
        colunas = [i for i, celula in enumerate(planilha.iloc[5]) if celula == rotulo]
        rotulos = planilha[0].astype(str).str.strip()
        brasil = planilha[rotulos == "BRASIL"].iloc[0]
        for coluna, posicao in (("area_plantada", colunas[0]), ("producao", colunas[2])):
            ufs = {
                uf: float(planilha.loc[rotulos == uf, posicao].iloc[0])
                for uf in rotulos[rotulos.isin(constants.CONAB_UFS)]
            }
            assert dict(zip(frames[produto]["uf"], frames[produto][coluna], strict=True)) == ufs
            publicados[(produto, coluna)] = (ufs, float(brasil[posicao]))

    milho, brasil_milho = publicados[("milho_2", "producao")]
    aveia, brasil_aveia = publicados[("aveia", "area_plantada")]
    assert (round(sum(milho.values()), 1), brasil_milho) == (59981.5, 60741.6)
    assert (round(sum(aveia.values()), 1), brasil_aveia) == (508.1, 503.4)
    assert [str(a.message) for a in avisos if "soma das" in str(a.message)] == [
        _aviso_de_soma(
            "milho_2 2020/21 (12º levantamento de 2021/22)", "producao", milho, brasil_milho
        ),
        _aviso_de_soma(
            "aveia 2020/21 (12º levantamento de 2021/22)", "area_plantada", aveia, brasil_aveia
        ),
    ]


def _celulas_da_serie(
    produto: str, aba: str, periodo: str, escala: float = 1.0
) -> tuple[dict[str, float], float]:
    planilha = xlrd.open_workbook(str(SERIES[produto])).sheet_by_name(aba)
    coluna = [_rotulo(valor) for valor in planilha.row_values(5)].index(periodo)
    ufs = {}
    brasil = None
    for linha in range(6, planilha.nrows):
        rotulo = _rotulo(planilha.cell_value(linha, 0)).upper()
        valor = planilha.cell_value(linha, coluna)
        if rotulo in constants.CONAB_UFS and isinstance(valor, float):
            ufs[rotulo] = valor * escala
        elif rotulo == "BRASIL":
            brasil = valor * escala
    assert brasil is not None
    return ufs, brasil


async def test_soma_das_ufs_da_serie_diferente_do_brasil_avisa_sem_mudar_o_dado(
    monkeypatch: pytest.MonkeyPatch,
):
    _servir(monkeypatch)
    with warnings.catch_warnings(record=True) as avisos:
        warnings.simplefilter("always")
        trigo = await serie_api.serie_historica("trigo")
        cafe = await serie_api.serie_historica("cafe")
        safra_antiga = await api.safras("trigo", safra="2003/04")

    esperados = []
    for periodo in ("2003", "2004"):
        for aba, coluna in (("Área", "area_plantada_mil_ha"), ("Produção", "producao_mil_ton")):
            ufs, brasil = _celulas_da_serie("trigo", aba, periodo)
            linhas = trigo[trigo["safra"] == periodo]
            assert dict(zip(linhas["uf"], linhas[coluna], strict=True)) == ufs
            esperados.append(
                _aviso_de_soma(f"trigo {periodo} (série histórica)", coluna, ufs, brasil)
            )
    area_cafe, brasil_cafe = _celulas_da_serie("cafe", "Área em produção", "2007", 0.001)
    linhas = cafe[cafe["safra"] == "2007"]
    assert dict(zip(linhas["uf"], linhas["area_em_producao_mil_ha"], strict=True)) == area_cafe
    esperados.append(
        _aviso_de_soma(
            "cafe 2007 (série histórica)", "area_em_producao_mil_ha", area_cafe, brasil_cafe
        )
    )
    for aba, coluna in (("Área", "area_plantada"), ("Produção", "producao")):
        ufs, brasil = _celulas_da_serie("trigo", aba, "2004")
        assert dict(zip(safra_antiga["uf"], safra_antiga[coluna], strict=True)) == ufs
        esperados.append(_aviso_de_soma("trigo 2003/04 (série histórica)", coluna, ufs, brasil))

    producao_cafe, brasil_producao_cafe = _celulas_da_serie("cafe", "Produção", "2007", 0.06)
    assert len(producao_cafe) == 11
    assert round(brasil_producao_cafe - sum(producao_cafe.values()), 2) == 24.24
    assert [str(a.message) for a in avisos if "soma das" in str(a.message)] == esperados


def _inverno(arquivo: str, aba: str, rotulo: str) -> dict[str, dict[str, float | None]]:
    planilha = xlrd.open_workbook(str(R3 / arquivo)).sheet_by_name(aba)
    colunas = [i for i, valor in enumerate(planilha.row_values(5)) if valor == rotulo]
    assert len(colunas) == 3
    publicado = {}
    for linha in range(7, planilha.nrows):
        uf = str(planilha.cell_value(linha, 0)).strip()
        if uf in constants.CONAB_UFS:
            valores = [planilha.cell_value(linha, coluna) for coluna in colunas]
            publicado[uf] = {
                campo: valor if isinstance(valor, float) else None
                for campo, valor in zip(GENERICAS, valores, strict=True)
            }
    return publicado


async def test_cereais_de_inverno_saem_da_aba_com_ano_que_publica_a_safra(
    monkeypatch: pytest.MonkeyPatch,
):
    _servir(monkeypatch)
    monkeypatch.setitem(EDICOES, SET_2021_URL, R3 / "1569ad0cf189f4d5.xls")
    monkeypatch.setitem(EDICOES, AGO_2020_URL, R3 / "c3c31afbe3e52d92.xls")
    edicoes = [
        {
            "url": SET_2021_URL,
            "levantamento": 12,
            "safra": "2020/21",
            "ano_inicio": 2020,
            "ano_fim": 21,
            "data_publicacao": date(2021, 9, 9),
        },
        {
            "url": AGO_2020_URL,
            "levantamento": 11,
            "safra": "2019/20",
            "ano_inicio": 2019,
            "ano_fim": 20,
            "data_publicacao": date(2020, 8, 11),
        },
    ]
    monkeypatch.setattr(client, "list_levantamentos", AsyncMock(return_value=edicoes))
    assert xlrd.open_workbook(str(R3 / "1569ad0cf189f4d5.xls")).sheet_by_name(
        "Trigo 2020"
    ).row_values(5)[1:3] == [23, 23]

    for cultura in INVERNO:
        for safra, levantamento, arquivo, aba, rotulo in (
            ("2019/20", None, "1569ad0cf189f4d5.xls", f"{cultura} 2021", "Safra 2020"),
            ("2020/21", 12, "1569ad0cf189f4d5.xls", f"{cultura} 2021", "Safra 2021"),
            ("2019/20", 11, "c3c31afbe3e52d92.xls", f"{cultura} 2020", "Safra 2020"),
        ):
            with helpers.sem_excecao():
                frame = await api.safras(cultura.lower(), safra=safra, levantamento=levantamento)
            publicado = _inverno(arquivo, aba, rotulo)
            assert len(publicado) == 27
            assert {
                linha.uf: _nulos(
                    {
                        "area_plantada": linha.area_plantada,
                        "produtividade": linha.produtividade,
                        "producao": linha.producao,
                    }
                )
                for linha in frame.itertuples()
            } == publicado, (cultura, safra, levantamento)
            assert set(frame["safra"]) == {safra}


def test_edicao_sem_o_inverno_da_safra_nao_traz_estimativa():
    parser = v1.ConabParserV1()
    for arquivo, safra in (
        ("c3c31afbe3e52d92.xls", "2020/21"),
        ("1569ad0cf189f4d5.xls", "2021/22"),
    ):
        bruto = (R3 / arquivo).read_bytes()
        for cultura in INVERNO:
            with helpers.sem_excecao():
                registros = parser.parse_safra_produto(BytesIO(bruto), cultura.lower(), safra)
            assert registros == [], (arquivo, cultura)


async def test_brasil_total_pela_serie_avisa_partes_que_nao_fecham_com_o_total(
    monkeypatch: pytest.MonkeyPatch,
):
    _servir(monkeypatch)
    with warnings.catch_warnings(record=True) as avisos:
        warnings.simplefilter("always")
        frame = await api.brasil_total(safra="2021/22")

    esperados = []
    for safra in "123":
        total = _serie(f"feijao_{safra}", "2021/22", "BRASIL")
        tipos = [_serie(f"feijao_{tipo}_{safra}", "2021/22", "BRASIL") for tipo in TIPOS]
        for campo in ("area_plantada", "producao"):
            soma = sum(tipo[campo] for tipo in tipos)
            esperados.append(
                f"CONAB: no brasil_total de 2021/22 (série histórica), Cores + Preto + Caupi em "
                f"{campo} somam {soma:.1f}, e FEIJÃO {safra}ª SAFRA publicado, {total[campo]:.1f} "
                f"(diferença de {soma - total[campo]:.1f}); o agrobr repassa os números publicados"
            )
    terceira = [_serie(f"feijao_{tipo}_3", "2021/22", "BRASIL")["producao"] for tipo in TIPOS]
    assert (terceira, _serie("feijao_3", "2021/22", "BRASIL")["producao"]) == (
        [700.6, 9.7, 37.7],
        707.2,
    )
    cores = frame[(frame["rotulo"] == "Cores") & (frame["grupo"] == "FEIJÃO 3ª SAFRA")]
    assert float(cores["producao"].iloc[0]) == 700.6
    assert [str(a.message) for a in avisos if "brasil_total de" in str(a.message)] == esperados

    safras = [_serie(f"feijao_{safra}", "2015/16", "BRASIL")["producao"] for safra in "123"]
    assert round(sum(safras) - _serie("feijao", "2015/16", "BRASIL")["producao"], 1) == 0.6
    with warnings.catch_warnings(record=True) as avisos:
        warnings.simplefilter("always")
        await api.brasil_total(safra="2015/16")
    assert [str(a.message) for a in avisos if "brasil_total de" in str(a.message)] == [
        "CONAB: no brasil_total de 2015/16 (série histórica), subtotal de verão: fora da soma, "
        "ainda não levantados: GERGELIM (a série começa em 2018/19)"
    ]


@pytest.mark.parametrize("safra", ["2018/19", "2017/18", "1975/76"])
async def test_brasil_total_pela_serie_nao_soma_ausencia_como_zero(
    monkeypatch: pytest.MonkeyPatch, safra: str
):
    _servir(monkeypatch)
    with warnings.catch_warnings(record=True) as avisos:
        warnings.simplefilter("always")
        frame, meta = await api.brasil_total(safra=safra, return_meta=True)

    prefixo = f"CONAB: no brasil_total de {safra} (série histórica), subtotal de verão: "
    emitidos = [str(a.message) for a in avisos if str(a.message).startswith(prefixo)]
    assert [aviso for aviso in meta.validation_warnings if aviso.startswith(prefixo)] == emitidos
    verao = [
        rotulo for rotulo, _, _, secao in constants.CONAB_BRASIL_TOTAL_SERIES if secao == "verao"
    ]
    produtos = frame[~frame["rotulo"].isin(["SUBTOTAL", "BRASIL (2)"])]
    partes = produtos[produtos["rotulo"].isin(verao) & produtos["grupo"].isna()]
    inverno = produtos[produtos["grupo"] == constants.CONAB_INVERNO]
    subtotal, subtotal_inverno, brasil = frame[
        frame["rotulo"].isin(["SUBTOTAL", "BRASIL (2)"])
    ].to_dict("records")
    for campo in ("area_plantada", "producao"):
        assert subtotal_inverno[campo] == pytest.approx(float(inverno[campo].sum()))

    if safra == "1975/76":
        assert emitidos[0].startswith(f"{prefixo}fora da soma, ainda não levantados: ")
        assert "SOJA (a série começa em 1976/77)" in emitidos[0]
        assert "GERGELIM (a série começa em 2018/19)" in emitidos[0]
        assert emitidos[1:] == [f"{prefixo}nulo, e o BRASIL também (nenhuma cultura levantada)"]
        assert partes[["area_plantada", "producao"]].isna().all().all()
        assert subtotal_inverno["area_plantada"] == pytest.approx(142.5)
        for linha in (subtotal, brasil):
            assert pd.isna(linha["area_plantada"]) and pd.isna(linha["producao"])
            assert pd.isna(linha["produtividade"])
        return

    if safra == "2017/18":
        assert emitidos == [
            f"{prefixo}fora da soma, ainda não levantados: GERGELIM (a série começa em 2018/19)"
        ]
        gergelim = partes[partes["rotulo"] == "GERGELIM"]
        assert gergelim[["area_plantada", "producao"]].isna().all().all()
        partes = partes[partes["rotulo"] != "GERGELIM"]
    else:
        assert emitidos == []
    assert partes[["area_plantada", "producao"]].notna().all().all()
    for campo in ("area_plantada", "producao"):
        assert subtotal[campo] == pytest.approx(float(partes[campo].sum()))
        assert brasil[campo] == pytest.approx(subtotal[campo] + subtotal_inverno[campo])


async def test_serie_com_brasil_ambiguo_nao_confere_a_soma_das_ufs(
    monkeypatch: pytest.MonkeyPatch,
):
    _servir(monkeypatch)
    ler = serie_parser._read_selected_sheet

    def duplicar_brasil(*args: Any) -> pd.DataFrame:
        frame = ler(*args)
        brasil = frame[frame[0].astype(str).str.strip() == "BRASIL"]
        return pd.concat([frame, brasil], ignore_index=True)

    monkeypatch.setattr(serie_parser, "_read_selected_sheet", duplicar_brasil)
    with warnings.catch_warnings(record=True) as avisos:
        warnings.simplefilter("always")
        with helpers.sem_excecao():
            trigo = await serie_api.serie_historica("trigo")
    assert [str(a.message) for a in avisos if "soma das" in str(a.message)] == []
    assert sorted(trigo.loc[trigo["safra"] == "2003", "uf"]) == sorted(constants.CONAB_UFS)


async def test_safra_que_acabou_de_sair_do_boletim_confere_a_serie_com_a_ultima_edicao(
    monkeypatch: pytest.MonkeyPatch,
):
    _servir(monkeypatch)
    edicoes = [
        {
            "url": PROXIMA_EDICAO_URL,
            "levantamento": 1,
            "safra": "2026/27",
            "ano_inicio": 2026,
            "ano_fim": 27,
            "data_publicacao": date(2026, 10, 8),
        },
        {
            "url": SET_2026_URL,
            "levantamento": 12,
            "safra": "2025/26",
            "ano_inicio": 2025,
            "ano_fim": 26,
            "data_publicacao": date(2026, 9, 15),
        },
    ]
    monkeypatch.setattr(client, "list_levantamentos", AsyncMock(return_value=edicoes))
    monkeypatch.setattr(
        serie_parser, "referencia_publicacao", lambda _: ("outubro/2026", date(2026, 10, 1))
    )
    with warnings.catch_warnings(record=True) as avisos:
        warnings.simplefilter("always")
        gergelim, meta = await api.safras("gergelim", safra="2024/25", return_meta=True)
        _, meta_soja = await api.safras("soja", safra="2024/25", return_meta=True)
        _, meta_total = await api.brasil_total(safra="2024/25", return_meta=True)

    livro = openpyxl.load_workbook(EDICOES[SET_2026_URL], read_only=True, data_only=True)
    linhas = list(livro["Gergelim"].iter_rows(values_only=True))
    colunas = [i for i, celula in enumerate(linhas[5]) if celula == "Safra 24/25"]
    brasil = next(valores for valores in linhas if valores[0] == "BRASIL")
    aba = livro["Brasil - Total por Produto"].iter_rows(values_only=True)
    totais = {str(valores[0]).strip(): valores for valores in aba if valores[0]}
    serie = _serie("gergelim", "2024/25", "BRASIL")
    publicacao = {
        "levantamento": 12,
        "safra": "2025/26",
        "data_publicacao": "2026-09-15",
        "url": SET_2026_URL,
    }
    divergencias = [
        {"linha": linha, "coluna": coluna, "serie": serie[coluna], "boletim": brasil[posicao]}
        for linha, (coluna, posicao) in (
            ("gergelim", ("area_plantada", colunas[0])),
            ("gergelim", ("producao", colunas[2])),
        )
    ]
    assert [(d["serie"], d["boletim"]) for d in divergencias] == [(608.0, 901.8), (399.4, 610.9)]
    assert meta.source_details["publicacao"].get("conferencia") == {
        **publicacao,
        "divergencias": divergencias,
    }
    assert meta_soja.source_details["publicacao"].get("conferencia") == {
        **publicacao,
        "divergencias": [],
    }
    assert gergelim["producao"].tolist() == [
        _serie("gergelim", "2024/25", uf)["producao"] for uf in gergelim["uf"]
    ]

    feijao = {
        "linha": "FEIJÃO 3ª SAFRA",
        "coluna": "producao",
        "serie": _serie("feijao_3", "2024/25", "BRASIL")["producao"],
        "boletim": totais["FEIJÃO 3ª SAFRA"][7],
    }
    no_total = [
        feijao,
        *(
            {**divergencia, "linha": "GERGELIM", "boletim": totais["GERGELIM"][posicao]}
            for divergencia, posicao in zip(divergencias, (1, 7), strict=True)
        ),
    ]
    assert meta_total.source_details["publicacao"].get("conferencia") == {
        **publicacao,
        "divergencias": no_total,
    }

    def aviso(itens: list[dict[str, Any]]) -> str:
        return (
            "CONAB: 2024/25 sai da série histórica, que diverge do 12º levantamento de 2025/26, a "
            "última edição do boletim que publicou a safra: "
            + "; ".join(
                f"{d['linha']} {d['coluna']} {d['serie']:.1f} na série × {d['boletim']:.1f} no "
                "boletim"
                for d in itens
            )
            + ". A legenda da série é posterior, mas a coluna pode não ter sido revista"
        )

    assert [str(a.message) for a in avisos if "sai da série" in str(a.message)] == [
        aviso(divergencias),
        aviso(no_total),
    ]


async def test_balanco_aceita_o_produto_com_acento_e_caixa(monkeypatch: pytest.MonkeyPatch):
    _servir(monkeypatch)
    with helpers.sem_excecao(), warnings.catch_warnings():
        warnings.simplefilter("ignore")
        feijao = await api.balanco("feijao", safra="2025/26")
        acentuado = await api.balanco("Feijão", safra="2025/26")
        algodao = await api.balanco("ALGODÃO", safra="2025/26")
    assert set(feijao["produto"]) == {"FEIJÃO"}
    pd.testing.assert_frame_equal(acentuado, feijao)
    assert set(algodao["produto"]) == {"ALGODÃO"}
