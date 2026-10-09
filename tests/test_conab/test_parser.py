from __future__ import annotations

from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import Any

import openpyxl
import pytest

from agrobr.conab.parsers.v1 import ConabParserV1
from agrobr.exceptions import ParseError
from tests import helpers

SAMPLE_FILE = (
    Path(__file__).parent.parent
    / "golden_data/conab/levantamento_12_2024_25_20260922"
    / "site_previsao_de_safra-por_produto-set-2025.xlsx"
)


@pytest.fixture
def sample_xlsx():
    if not SAMPLE_FILE.exists():
        pytest.skip(f"Arquivo de amostra não encontrado: {SAMPLE_FILE}")

    with open(SAMPLE_FILE, "rb") as f:
        return BytesIO(f.read())


@pytest.fixture
def parser():
    return ConabParserV1()


class TestConabParserEdgeCases:
    @pytest.mark.parametrize(
        "scenario,parameters",
        [("test_parse_decimal_milhar_br", {}), ("test_parse_decimal_traco", {})],
        ids=["parse_decimal_milhar_br-0", "parse_decimal_traco-0"],
    )
    def test_numeros_textuais_e_nulos(self, scenario: str, parameters: dict[str, Any], parser: Any):
        with (
            helpers.collect_failures() as check,
            check((scenario, parameters)),
            helpers.isolated_dataset_case((scenario, parameters)),
        ):
            if scenario == "test_parse_decimal_milhar_br":
                assert parser._parse_decimal("1.234,5") == Decimal("1234.5")
            elif scenario == "test_parse_decimal_traco":
                assert parser._parse_decimal("-") is None

    def test_parse_brasil_total_no_header(self, parser, sample_xlsx):
        wb = openpyxl.load_workbook(sample_xlsx)
        ws = wb["Brasil - Total por Produto"]
        assert ws["A5"].value == "PRODUTO"
        ws["A5"] = "CABEÇALHO ALTERADO"
        buf = BytesIO()
        wb.save(buf)
        wb.close()
        buf.seek(0)
        with pytest.raises(ParseError, match="header.*Brasil - Total por Produto"):
            parser.parse_brasil_total(buf)


class TestConabParserV1:
    def test_parse_produto_invalido(self, parser, sample_xlsx):
        sample_xlsx.seek(0)

        with pytest.raises(ParseError, match="Produto não suportado: produto_inexistente"):
            parser.parse_safra_produto(
                xlsx=sample_xlsx,
                produto="produto_inexistente",
            )
