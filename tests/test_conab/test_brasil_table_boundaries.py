from __future__ import annotations

from collections import Counter
from io import BytesIO
from pathlib import Path
from unittest.mock import AsyncMock

import openpyxl
import pytest

from agrobr.conab import api, client
from tests.helpers import conferir_corpo

FIXTURE = Path(__file__).parents[1] / "golden_data/conab/safra_2025_26_agosto/response.xlsx"
SET_2026 = Path(__file__).parents[1] / "golden_data/reconciliacao_r3_20260918/7cd4df7946e5c57f.xlsx"


@pytest.mark.asyncio
@pytest.mark.parametrize("missing_metrics", [False, True])
async def test_brasil_total_reads_product_table_and_preserves_null_rows(
    monkeypatch: pytest.MonkeyPatch, missing_metrics: bool
):
    content = FIXTURE.read_bytes()
    workbook = openpyxl.load_workbook(BytesIO(content), data_only=True)
    try:
        sheet = workbook["Brasil - Total por Produto"]
        assert sheet["A39"].value == "CULTURAS DE INVERNO"
        assert sheet["E39"].value == "PRODUTIVIDADE (Em kg/ha)"
        assert sheet["A50"].value.startswith("Legenda:")
        expected_labels = [sheet.cell(row, 1).value for row in (*range(8, 39), *range(42, 50))]
        if missing_metrics:
            for col in (2, 3, 5, 6, 8, 9):
                sheet.cell(13, col).value = None
            buffer = BytesIO()
            workbook.save(buffer)
            content = buffer.getvalue()
    finally:
        workbook.close()
    monkeypatch.setattr(
        client,
        "fetch_safra_xlsx",
        AsyncMock(return_value=(BytesIO(content), {"source_method": "httpx"})),
    )

    frame, meta = await api.brasil_total(return_meta=True)

    conferir_corpo(meta, content)
    assert len(frame) == meta.records_count == 78
    assert frame["rotulo"].tolist() == [label for label in expected_labels for _ in range(2)]
    assert frame["safra"].tolist() == ["2024/25", "2025/26"] * 39
    assert len(frame[frame["produto"] == "subtotal"]) == 4
    assert len(frame[frame["produto"] == "brasil"]) == 2
    assert frame.loc[frame["rotulo"] == "ALGODÃO - CAROÇO (1)", "produto"].tolist() == [
        "algodao_caroco",
        "algodao_caroco",
    ]
    assert not frame["produto"].str.contains(r"[()]").any()
    assert meta.schema_version == "2.0"
    assert frame[["area_plantada", "produtividade", "producao"]].dtypes.eq("float64").all()
    if missing_metrics:
        arroz = frame[frame["produto"] == "arroz"]
        assert len(arroz) == 2
        assert arroz[["area_plantada", "produtividade", "producao"]].isna().all().all()


@pytest.mark.asyncio
@pytest.mark.parametrize("as_polars", [False, True])
async def test_brasil_total_identifica_a_epoca_do_feijao_e_o_bloco_do_subtotal(
    monkeypatch: pytest.MonkeyPatch,
    as_polars: bool,
):
    if as_polars:
        polars = pytest.importorskip("polars")
    content = SET_2026.read_bytes()
    workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        rows = list(workbook["Brasil - Total por Produto"].iter_rows(values_only=True))
    finally:
        workbook.close()
    assert rows[5][8] == "Safra 25/26"
    assert [rows[title - 1][0] for title in (17, 21, 25)] == [
        "FEIJÃO 1ª SAFRA",
        "FEIJÃO 2ª SAFRA",
        "FEIJÃO 3ª SAFRA",
    ]
    assert rows[38][0] == "CULTURAS DE INVERNO"
    expected = [
        (rows[title - 1][0], row[0], float(row[8]))
        for title in (17, 21, 25)
        for row in rows[title : title + 3]
    ]
    expected += [
        (None, rows[37][0], float(rows[37][8])),
        (rows[38][0], rows[47][0], float(rows[47][8])),
    ]
    assert [value for _, label, value in expected if label == "Cores"] == pytest.approx(
        [601.7, 456.9, 623.4]
    )
    monkeypatch.setattr(
        client,
        "fetch_safra_xlsx",
        AsyncMock(return_value=(BytesIO(content), {"source_method": "httpx"})),
    )

    frame = await api.brasil_total(as_polars=as_polars)
    if as_polars:
        assert isinstance(frame, polars.DataFrame)
        assert frame.schema["producao"] == polars.Float64
        assert frame.schema["produto"] == polars.String
        frame = frame.to_pandas()

    current = frame[frame["safra"] == "2025/26"]
    observed = [
        (group if isinstance(group, str) else None, label, float(value))
        for group, label, value in current[["grupo", "rotulo", "producao"]].itertuples(index=False)
        if label in {"Cores", "Preto", "Caupi", "SUBTOTAL"}
    ]
    assert Counter(observed) == Counter(expected)
    assert not frame.duplicated(["produto", "grupo", "safra"]).any()
    assert frame.loc[frame["produto"] == "brasil", "grupo"].isna().all()
    cores = current.loc[current["rotulo"] == "Cores", ["produto", "producao"]]
    assert cores["produto"].tolist() == ["feijao_cores_1", "feijao_cores_2", "feijao_cores_3"]
    assert cores["producao"].tolist() == pytest.approx([601.7, 456.9, 623.4])
    assert current["unidade_area"].eq("mil_ha").all()
    assert current["unidade_producao"].eq("mil_ton").all()
