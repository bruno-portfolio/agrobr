from __future__ import annotations

import json
import warnings
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

import bs4
import openpyxl
import pytest

from agrobr import constants, datasets
from agrobr.conab import api, client
from agrobr.exceptions import InvalidParameterError, ParseError, SourceUnavailableError
from agrobr.models import Safra
from tests import helpers

_JANUARY_PATH = "4o-levantamento-safra-2025-26/site_previsao_de_safra-por_produto-jan-2026"
R3 = Path(__file__).parents[1] / "golden_data/reconciliacao_r3_20260918"
SET_2025 = Path(__file__).parents[1] / "golden_data/conab/levantamento_12_2024_25_20260922"
SET_2025_URL = json.loads((SET_2025 / "PROVENANCE.json").read_text(encoding="utf-8"))["files"][0][
    "url"
]
SET_2026_URL = (
    "https://www.gov.br/conab/pt-br/atuacao/informacoes-agropecuarias/safras/safra-de-graos/"
    "boletim-da-safra-de-graos/12o-levantamento-safra-2025-26/"
    "site_previsao_de_safra-por_produto-set-2026.xlsx"
)
EDICOES = {
    SET_2025_URL: SET_2025 / SET_2025_URL.rsplit("/", 1)[1],
    SET_2026_URL: R3 / "7cd4df7946e5c57f.xlsx",
}


_CATALOG_URL = constants.URLS[constants.Fonte.CONAB]["boletim_graos"]
_OCTOBER_PATH = (
    "1o-levantamento-safra-2024-25/tabela-de-dados-producao-e-balanco-de-oferta-e-demanda-de-graos"
)


def _table(path: str, label: str = "Tabela de dados") -> str:
    return f'<a href="{_CATALOG_URL}/{path}">{label}</a>'


@pytest.mark.parametrize("entrypoint", ["api", "client"])
@pytest.mark.parametrize(
    "kwargs",
    [
        {"levantamento": True},
        {"levantamento": False},
        {"levantamento": 1.0},
        {"levantamento": "1"},
        {"levantamento": 0},
        {"levantamento": 13},
        {"safra": True},
        {"safra": 2025},
        {"safra": ""},
        {"safra": "2025"},
        {"safra": "2025/27"},
        {"safra": "2025/2024"},
        {"safra": "2025/2126"},
        {"safra": "2025-26"},
    ],
)
async def test_selection_rejects_invalid_before_io(entrypoint: str, kwargs: dict):
    with (
        patch.object(client, "list_levantamentos", new_callable=AsyncMock) as listing,
        patch.object(client, "download_xlsx", new_callable=AsyncMock) as download,
        helpers.collect_failures() as check,
    ):
        with check("erro"), pytest.raises(InvalidParameterError):
            if entrypoint == "api":
                await api.safras("soja", **kwargs)
            else:
                await client.fetch_safra_xlsx(**kwargs)
        with check("antes_da_rede"):
            assert listing.await_count == 0
            assert download.await_count == 0


async def test_missing_edition_fails_after_catalog_exhausted_without_download():
    next_url = f"{_CATALOG_URL}?b_start:int=30"
    first = (
        _table("11o-levantamento-safra-2025-26/dados.xlsx")
        + f'<a class="proximo" href="{next_url}">Próximo</a>'
    )
    with (
        patch.object(client, "fetch_boletim_page", new_callable=AsyncMock, return_value=first),
        patch.object(
            client,
            "_fetch_http",
            new_callable=AsyncMock,
            return_value=_table(_OCTOBER_PATH).encode(),
        ) as fetch,
        patch.object(client, "download_xlsx", new_callable=AsyncMock) as download,
        pytest.raises(SourceUnavailableError, match="levantamento=12"),
    ):
        await client.fetch_safra_xlsx(safra="2025/26", levantamento=12)
    fetch.assert_awaited_once_with(next_url)
    download.assert_not_awaited()


async def _metadata_for(levantamento: int, uf: str | None, safra: str | None = "2025/26"):
    row = Safra(
        fonte=constants.Fonte.CONAB,
        produto="soja",
        safra="2025/26",
        uf="MT",
        area_plantada=Decimal("1"),
        producao=Decimal("3"),
        produtividade=Decimal("3000"),
        levantamento=levantamento,
        data_publicacao=date(2026, 1, 15),
    )
    metadata = {
        "safra": "2025/26",
        "levantamento": levantamento,
        "url": f"https://www.gov.br/conab/lev{levantamento}.xlsx",
        "data_publicacao": date(2026, 1, 15),
        "source_method": "httpx",
    }
    parsed = MagicMock(version=1)
    parsed.parse_safra_produto.return_value = [row]
    with (
        patch.object(
            client,
            "fetch_safra_xlsx",
            new_callable=AsyncMock,
            return_value=(BytesIO(b"data"), metadata),
        ),
        patch.object(api, "ConabParserV1", return_value=parsed),
    ):
        frame, meta = await api.safras(
            "soja", safra=safra, levantamento=levantamento, uf=uf, return_meta=True
        )
    return frame, meta


async def test_metadata_cache_key_identifies_effective_edition_and_uf():
    first, meta_first = await _metadata_for(1, "MT")
    _, meta_latest = await _metadata_for(11, "MT")
    _, meta_other_uf = await _metadata_for(1, "GO")
    _, meta_alias = await _metadata_for(1, " mt ", "2025/2026")
    _, meta_implicit_year = await _metadata_for(1, "MT", None)
    assert meta_first.cache_key != meta_latest.cache_key
    assert meta_first.cache_key != meta_other_uf.cache_key
    assert meta_first.cache_key == meta_alias.cache_key == meta_implicit_year.cache_key
    assert meta_first.cache_expires_at is None
    assert meta_first.source_url.endswith("lev1.xlsx")
    assert meta_latest.source_url.endswith("lev11.xlsx")
    assert not meta_first.from_cache
    assert first.loc[0, "levantamento"] == 1
    assert first.loc[0, "data_publicacao"].date() == date(2026, 1, 15)


async def test_catalog_does_not_admit_unrelated_or_non_table_links():
    html = "".join(
        [
            _table("4o-levantamento-safra-2025-26/apresentacao.pdf"),
            _table("4o-levantamento-safra-2025-26/boletim", "Boletim da Safra de Grãos"),
            '<a href="https://example.com/4o-levantamento-safra-2025-26/dados">Tabela</a>',
            '<a href="javascript:alert(1)">Tabela</a>',
            '<a href="/tabela-de-precos">Tabela de preços</a>',
        ]
    )
    assert await client.list_levantamentos(html=html) == []


@pytest.mark.parametrize(
    ("origem", "destino"),
    [("https://", "http://"), ("www.gov.br/", "www.gov.br:8443/"), ("https://", "https://u:p@")],
)
async def test_catalog_does_not_admit_plain_http_port_or_userinfo(origem: str, destino: str):
    url = f"{_CATALOG_URL}/{_JANUARY_PATH}".replace(origem, destino)
    assert await client.list_levantamentos(html=f'<a href="{url}">Tabela de dados</a>') == []


@pytest.mark.parametrize(
    "next_url", [_CATALOG_URL, "https://example.com/", "https://www.gov.br/other"]
)
async def test_catalog_rejects_pagination_cycle_or_escape(next_url: str):
    html = _table(_JANUARY_PATH) + f'<a class="proximo" href="{next_url}">Próximo</a>'
    with (
        patch.object(client, "fetch_boletim_page", new_callable=AsyncMock, return_value=html),
        patch.object(client, "_fetch_http", new_callable=AsyncMock) as fetch,
        pytest.raises(
            ParseError, match="Ciclo na paginação do catálogo|Paginação fora do catálogo CONAB"
        ),
    ):
        await client.list_levantamentos()
    fetch.assert_not_awaited()


async def test_catalog_broken_second_page_is_not_reported_as_complete():
    html = (
        _table(_JANUARY_PATH)
        + f'<a class="proximo" href="{_CATALOG_URL}?b_start:int=30">Próximo</a>'
    )
    with (
        patch.object(client, "fetch_boletim_page", new_callable=AsyncMock, return_value=html),
        patch.object(
            client, "_fetch_http", new_callable=AsyncMock, return_value=b"<html>maintenance</html>"
        ),
        pytest.raises(ParseError, match="sem tabelas"),
    ):
        await client.list_levantamentos()


async def test_catalog_and_download_accept_official_legacy_xls():
    path = "7o-levantamento-safra-2021-22/previsao-de-safra-por-produto-abr-2022-site.xls"
    found = await client.list_levantamentos(html=_table(path))
    assert len(found) == 1
    assert found[0]["safra"] == "2021/22"
    data = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"0" * 2000
    with patch.object(client, "_fetch_http", new_callable=AsyncMock, return_value=data):
        result = await client.download_xlsx(found[0]["url"])
    assert result.getvalue() == data


def _servir_edicoes_oficiais(monkeypatch):
    catalogo = bs4.BeautifulSoup((R3 / "dc5524f326c16e26.html").read_bytes(), "lxml")
    for link in catalogo.select("a.proximo, a[rel~=next]"):
        link.decompose()
    planilhas = {url: caminho.read_bytes() for url, caminho in EDICOES.items()}

    async def baixar(url):
        assert url in planilhas, url
        return planilhas[url]

    monkeypatch.setattr(client, "fetch_boletim_page", AsyncMock(return_value=str(catalogo)))
    monkeypatch.setattr(client, "_fetch_http", baixar)
    return planilhas


def _oficial(conteudo, aba, linha, safra):
    livro = openpyxl.load_workbook(BytesIO(conteudo), read_only=True, data_only=True)
    try:
        linhas = list(livro[aba].iter_rows(values_only=True))
    finally:
        livro.close()
    colunas = [indice for indice, celula in enumerate(linhas[5]) if celula == f"Safra {safra[2:]}"]
    valores = next(valores for valores in linhas if valores[0] == linha)
    assert len(colunas) == 3
    return tuple(float(valores[coluna]) for coluna in colunas)


async def test_safra_anterior_sem_levantamento_vem_da_revisao_mais_recente(monkeypatch):
    planilhas = _servir_edicoes_oficiais(monkeypatch)
    frame, meta = await api.safras("gergelim", safra="2024/25", uf="MT", return_meta=True)
    _, meta_corrente = await api.safras("gergelim", safra="2025/26", uf="MT", return_meta=True)
    total, meta_total = await api.brasil_total(safra="2024/25", return_meta=True)

    revisao = planilhas[SET_2026_URL]
    area, produtividade, producao = _oficial(revisao, "Gergelim", "MT", "2024/25")
    assert area == 695
    colunas = ["safra", "area_plantada", "produtividade", "producao", "levantamento"]
    assert frame[colunas].to_dict("records") == [
        {
            "safra": "2024/25",
            "area_plantada": area,
            "produtividade": produtividade,
            "producao": producao,
            "levantamento": 12,
        }
    ]
    assert frame["data_publicacao"].dt.date.tolist() == [date(2026, 9, 15)]
    publicacao = {
        "levantamento": 12,
        "safra": "2025/26",
        "data_publicacao": "2026-09-15",
        "url": SET_2026_URL,
    }
    assert meta.source_details["publicacao"] == publicacao
    assert meta_corrente.source_details["publicacao"] == publicacao
    assert meta.cache_key != meta_corrente.cache_key
    gergelim = total[total["produto"] == "gergelim"]
    assert gergelim["rotulo"].tolist() == ["GERGELIM"]
    assert gergelim["safra"].tolist() == ["2024/25"]
    assert [float(valor) for valor in gergelim["area_plantada"]] == [
        _oficial(revisao, "Brasil - Total por Produto", "GERGELIM", "2024/25")[0]
    ]
    assert meta_total.source_details["publicacao"] == publicacao


async def test_levantamento_fixa_a_publicacao_original_e_avisa_da_revisao(monkeypatch):
    planilhas = _servir_edicoes_oficiais(monkeypatch)
    with warnings.catch_warnings(record=True) as avisos:
        warnings.simplefilter("always")
        frame, meta = await api.safras(
            "gergelim", safra="2024/25", uf="MT", levantamento=12, return_meta=True
        )
        await api.safras("gergelim", safra="2025/26", uf="MT", levantamento=12)

    area, produtividade, producao = _oficial(planilhas[SET_2025_URL], "Gergelim", "MT", "2024/25")
    assert area == 401.2
    colunas = ["safra", "area_plantada", "produtividade", "producao", "levantamento"]
    assert frame[colunas].to_dict("records") == [
        {
            "safra": "2024/25",
            "area_plantada": area,
            "produtividade": produtividade,
            "producao": producao,
            "levantamento": 12,
        }
    ]
    assert meta.source_details["publicacao"] == {
        "levantamento": 12,
        "safra": "2024/25",
        "data_publicacao": "2025-09-11",
        "url": SET_2025_URL,
    }
    assert [str(aviso.message) for aviso in avisos if "CONAB" in str(aviso.message)] == [
        "CONAB: depois do 12º levantamento de 2024/25, a safra 2024/25 foi republicada no 12º "
        "levantamento de 2025/26 (2026-09-15), que pode trazer números revisados; sem "
        "`levantamento`, o agrobr entrega essa publicação"
    ]


async def test_estimativa_safra_segue_a_mesma_publicacao(monkeypatch):
    planilhas = _servir_edicoes_oficiais(monkeypatch)
    revisada = await datasets.estimativa_safra("milho", safra="2024/25", uf="MT", fonte="conab")
    original = await datasets.estimativa_safra(
        "milho", safra="2024/25", uf="MT", fonte="conab", levantamento=12
    )

    esperado = {
        "revisada": _oficial(planilhas[SET_2026_URL], "Milho Total", "MT", "2024/25"),
        "original": _oficial(planilhas[SET_2025_URL], "Milho Total", "MT", "2024/25"),
    }
    assert esperado["revisada"][2] != esperado["original"][2]
    colunas = ["area_plantada", "produtividade", "producao"]
    observado = {
        nome: tuple(float(valor) for valor in frame.loc[0, colunas])
        for nome, frame in (("revisada", revisada), ("original", original))
    }
    assert observado == esperado


def _celulas(conteudo, aba, *enderecos):
    livro = openpyxl.load_workbook(BytesIO(conteudo), read_only=True, data_only=True)
    try:
        return tuple(livro[aba][endereco].value for endereco in enderecos)
    finally:
        livro.close()


async def test_balanco_da_safra_vem_da_revisao_mais_recente(monkeypatch):
    planilhas = _servir_edicoes_oficiais(monkeypatch)
    revisado, meta = await api.balanco("trigo", safra="2024/25", return_meta=True)
    assert meta.source_url == SET_2026_URL
    soja = await api.balanco("soja", safra="2024/25")
    antigo, meta_antigo = await api.balanco("milho", safra="2018/19", return_meta=True)
    with warnings.catch_warnings(record=True) as avisos:
        warnings.simplefilter("always")
        original, meta_original = await api.balanco(
            "trigo", safra="2024/25", levantamento=12, return_meta=True
        )

    def producao(frame, periodo):
        return float(frame.loc[frame["safra"] == periodo, "producao"].iloc[0])

    assert (meta_original.source_url, meta_antigo.source_url) == (SET_2025_URL, SET_2025_URL)
    helpers.conferir_corpo(meta, planilhas[SET_2026_URL])
    helpers.conferir_corpo(meta_original, planilhas[SET_2025_URL])
    novo, velho = planilhas[SET_2026_URL], planilhas[SET_2025_URL]
    assert _celulas(novo, "Suprimento", "A38", "B43") == ("TRIGO", "2025*")
    assert _celulas(velho, "Suprimento", "B44", "C45") == ("2025**", "set/25")
    assert _celulas(velho, "Suprimento", "B30") == ("2018/19",)
    assert _celulas(novo, "Suprimento - Soja", "F6", "A9") == ("2024/25", "1.2. Produção")
    esperado = {
        "revisado": _celulas(novo, "Suprimento", "E43")[0],
        "original": _celulas(velho, "Suprimento", "E45")[0],
        "soja": _celulas(novo, "Suprimento - Soja", "F9")[0],
        "milho_2018": _celulas(velho, "Suprimento", "E30")[0],
    }
    assert (esperado["revisado"], esperado["original"]) == (7873.4, pytest.approx(7536.1))
    assert {
        "revisado": producao(revisado, "2025"),
        "original": producao(original, "2025"),
        "soja": producao(soja, "2024/25"),
        "milho_2018": producao(antigo, "2018/19"),
    } == esperado
    assert meta.source_details["publicacao"] == {
        "levantamento": 12,
        "safra": "2025/26",
        "data_publicacao": "2026-09-15",
        "url": SET_2026_URL,
    }
    assert any("republicada no 12º levantamento de 2025/26" in str(a.message) for a in avisos)
