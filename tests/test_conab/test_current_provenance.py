from __future__ import annotations

import hashlib
import json
from io import BytesIO
from pathlib import Path
from unittest.mock import AsyncMock

import bs4
import httpx
import openpyxl
import pytest

from agrobr import conab
from agrobr.conab import client
from tests.helpers import conferir_corpo

FIXTURE = Path(__file__).parents[1] / "golden_data/conab/safra_2025_26_agosto"


@pytest.mark.asyncio
@pytest.mark.parametrize("page_method", ["httpx", "playwright"])
@pytest.mark.parametrize("download_method", ["httpx", "playwright"])
async def test_current_wheat_matches_workbook_and_reports_download_method(
    monkeypatch: pytest.MonkeyPatch, page_method: str, download_method: str
):
    catalog = bs4.BeautifulSoup((FIXTURE / "response.html").read_bytes(), "lxml")
    for link in catalog.select("a.proximo, a[rel~=next]"):
        link.decompose()
    html = str(catalog).encode()
    content = (FIXTURE / "response.xlsx").read_bytes()
    provenance = json.loads((FIXTURE / "provenance.json").read_text(encoding="utf-8"))
    assert hashlib.sha256(content).hexdigest() == provenance["sha256"]

    async def fetch(url: str) -> bytes:
        is_xlsx = url.endswith(".xlsx")
        method = download_method if is_xlsx else page_method
        if method == "playwright":
            raise httpx.ConnectError("synthetic HTTP failure")
        return content if is_xlsx else html

    page = AsyncMock(return_value=html.decode())
    download = AsyncMock(return_value=BytesIO(content))
    monkeypatch.setattr(client, "_fetch_http", fetch)
    monkeypatch.setattr(client, "_fetch_boletim_page_browser", page)
    monkeypatch.setattr(client, "_download_xlsx_browser", download)

    frame, meta = await conab.safras("trigo", safra="2025/26", return_meta=True)

    workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = workbook["Trigo"]
        assert [sheet.cell(6, col).value for col in (3, 6, 9)] == ["Safra 2026"] * 3
        expected = {
            row[0]: (float(row[2]), float(row[5]), float(row[8]))
            for row in sheet.iter_rows(min_row=8, max_col=9, values_only=True)
            if isinstance(row[0], str) and len(row[0]) == 2
        }
    finally:
        workbook.close()
    actual = {
        row.uf: (row.area_plantada, row.produtividade, row.producao) for row in frame.itertuples()
    }
    assert len(expected) == 27
    assert actual == expected
    assert frame["producao"].sum() == pytest.approx(5813.1)
    assert page.await_count == (page_method == "playwright")
    assert download.await_count == (download_method == "playwright")
    assert meta.source_method == download_method
    assert meta.source_url == provenance["url"]
    conferir_corpo(meta, content)
