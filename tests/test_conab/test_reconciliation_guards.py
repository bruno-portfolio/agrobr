from __future__ import annotations

from io import BytesIO

import openpyxl
import pandas as pd
import pytest

from agrobr.conab import structure
from agrobr.conab.parsers.v1 import ConabParserV1
from agrobr.conab.progresso import parser as progress_parser
from agrobr.exceptions import ParseError
from tests.helpers import levanta_exatamente, progresso_xlsx_cells


@pytest.fixture
def survey_book():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Soja"
    for coordinate, value in {
        "A5": "REGIÃO/UF",
        "B5": "ÁREA (Em mil ha)",
        "E5": "PRODUTIVIDADE (Em kg/ha)",
        "H5": "PRODUÇÃO (Em mil t)",
        "B6": "Safra 24/25",
        "C6": "Safra 25/26",
        "D6": "VAR. %",
        "E6": "Safra 24/25",
        "F6": "Safra 25/26",
        "G6": "VAR. %",
        "H6": "Safra 24/25",
        "I6": "Safra 25/26",
        "J6": "VAR. %",
        "A8": "MT",
        "B8": 10,
        "C8": 11,
        "E8": 3000,
        "F8": 3000,
        "H8": 30,
        "I8": 33,
        "A9": "RS",
        "B9": 5,
        "C9": 6,
        "E9": 3000,
        "F9": 3000,
        "H9": 15,
        "I9": 18,
    }.items():
        sheet[coordinate] = value
    yield book
    book.close()


@pytest.mark.parametrize(
    "damage, message",
    [
        ("duplicate_sheet", "ambíguas"),
        ("missing_sheet", "obrigatória ausente"),
        ("duplicate_period", "Safra ambígua"),
        ("missing_metric", "detectar colunas de safra completas"),
        ("wrong_unit", "Unidade inesperada"),
        ("negative_selected_row", "Linha inválida"),
        ("missing_state_rows", "Nenhuma linha de UF"),
    ],
)
def test_survey_refuses_incomplete_or_ambiguous_selected_structure(survey_book, damage, message):
    sheet = survey_book["Soja"]
    if damage == "duplicate_sheet":
        survey_book.copy_worksheet(sheet).title = "Soja "
    elif damage == "missing_sheet":
        sheet.title = "Outra cultura"
    elif damage == "duplicate_period":
        sheet["C6"] = "Safra 24/25"
    elif damage == "missing_metric":
        sheet["F6"] = None
    elif damage == "wrong_unit":
        sheet["B5"] = "ÁREA (Em km²)"
    elif damage == "negative_selected_row":
        sheet["C9"] = -1
    elif damage == "missing_state_rows":
        sheet["A8"], sheet["A9"] = "BRASIL", "SUL"
    buffer = BytesIO()
    survey_book.save(buffer)
    with pytest.raises(ParseError, match=message):
        ConabParserV1().parse_safra_produto(buffer, "soja")


@pytest.mark.parametrize(
    "headers, message",
    [
        (
            [
                "PRODUTO",
                "SAFRA",
                "",
                "ESTOQUE INICIAL",
                "PRODUÇÃO",
                "IMPORTAÇÃO",
                "SUPRIMENTO",
                "CONSUMO",
                "EXPORTAÇÃO",
            ],
            "ausentes",
        ),
        (
            [
                "PRODUTO",
                "SAFRA",
                "",
                "ESTOQUE INICIAL",
                "PRODUÇÃO",
                "IMPORTAÇÃO",
                "SUPRIMENTO",
                "CONSUMO",
                "EXPORTAÇÃO",
                "ESTOQUE FINAL",
                "ESTOQUE FINAL",
            ],
            "ambígua",
        ),
        (
            [
                "PRODUTO",
                "SAFRA",
                "",
                "ESTOQUE INICIAL",
                "PRODUÇÃO",
                "IMPORTAÇÃO",
                "SUPRIMENTO",
                "CONSUMO",
                "EXPORTAÇÃO",
                "ESTOQUE FINAL",
                "NOVA MÉTRICA",
            ],
            "mudou",
        ),
    ],
)
def test_supply_header_refuses_missing_duplicate_and_unknown_fields(headers, message):
    with pytest.raises(ParseError, match=message):
        structure.supply_columns(pd.Series(headers), 2)


@pytest.fixture
def soy_book():
    book = openpyxl.Workbook()
    wide = book.active
    wide.title = "Suprimento - Soja"
    for row in [
        ["PRODUTO", "SAFRA"],
        [None, "2024/25", None, "2025/26"],
        ["1. Soja em grão"],
        ["1.1. Estoque Inicial", 10, None, 20],
        ["1.2. Produção", 100, None, 200],
        ["1.3. Importação", 1, None, 2],
        ["1.4. Sementes/Outros", 5, None, 10],
        ["1.5. Exportação", 40, None, 80],
        ["1.6. Processamento", 50, None, 100],
        ["1.7. Estoque Final", 16, None, 32],
    ]:
        wide.append(row)
    long = book.create_sheet("Suprimento")
    long.append(
        [
            "PRODUTO",
            "SAFRA",
            "",
            "ESTOQUE INICIAL",
            "PRODUÇÃO",
            "IMPORTAÇÃO",
            "SUPRIMENTO",
            "CONSUMO",
            "EXPORTAÇÃO",
            "ESTOQUE FINAL",
        ]
    )
    long.append(["SOJA", "2025/26", None, 10, 9999, 0, 10009, 1, 0, 10008])
    yield book
    book.close()


def test_soy_supply_uses_actual_period_columns_and_derived_components(soy_book):
    buffer = BytesIO()
    soy_book.save(buffer)
    result = ConabParserV1().parse_suprimento(buffer, "soja")
    latest = next(item for item in result if item["safra"] == "2025/26")
    assert latest["producao"] == 200
    assert latest["suprimento_total"] == 222
    assert latest["consumo"] == 110


@pytest.mark.parametrize(
    "damage, message",
    [
        ("duplicate_period", "Safra duplicada"),
        ("missing_component", "Itens ausentes"),
        ("renamed_component", "Item desconhecido"),
        ("unknown_extra_component", "Item desconhecido"),
        ("duplicate_component", "Item duplicado"),
        ("unreadable_selected", "header"),
    ],
)
def test_selected_soy_sheet_failure_does_not_fall_back_to_other_table(soy_book, damage, message):
    sheet = soy_book["Suprimento - Soja"]
    if damage == "duplicate_period":
        sheet["D2"] = "2024/25"
    elif damage == "missing_component":
        sheet.delete_rows(9)
    elif damage == "renamed_component":
        sheet["A9"] = "1.6. Componente desconhecido"
    elif damage == "unknown_extra_component":
        sheet.append(["1.8. Exportação especial", 1, None, 2])
    elif damage == "duplicate_component":
        sheet.append(["1.8. Estoque Inicial", 20, None, 40])
    else:
        sheet["A1"] = "Cabeçalho alterado"
    buffer = BytesIO()
    soy_book.save(buffer)
    with pytest.raises(ParseError, match=message):
        ConabParserV1().parse_suprimento(buffer, "soja")


@pytest.mark.parametrize(
    "damage, message", [("duplicate_long", "ambíguas"), ("missing_long", "obrigatória ausente")]
)
def test_supply_without_product_refuses_ambiguous_or_missing_long_sheet(soy_book, damage, message):
    if damage == "duplicate_long":
        soy_book.copy_worksheet(soy_book["Suprimento"]).title = "Suprimento "
    else:
        soy_book["Suprimento"].title = "Outra tabela"
    buffer = BytesIO()
    soy_book.save(buffer)
    with levanta_exatamente(ParseError, match=message):
        ConabParserV1().parse_suprimento(buffer)


@pytest.mark.parametrize(
    "raw, expected",
    [("0%", 0), ("0,5%", 0.005), ("1%", 0.01), (" 0.5%* ", 0.005), ("10%*", 0.1), (0.5, 0.5)],
)
def test_progress_explicit_percent_unit_always_divides_by_100(raw, expected):
    content = progresso_xlsx_cells({"D113": raw})
    frame = progress_parser.parse_progresso_xlsx(content)
    row = frame.loc[frame.cultura.eq("Trigo") & frame.uf.eq("BA")].iloc[0]
    assert row.pct_semana_anterior == pytest.approx(expected)
    assert bool(row.revisado) == (isinstance(raw, str) and raw.strip().endswith("*"))


def test_progress_duplicate_geography_is_not_silently_returned():
    with pytest.raises(ParseError, match="duplicadas"):
        progress_parser.parse_progresso_xlsx(progresso_xlsx_cells({"B13": "Maranhão"}))


def test_progress_ambiguous_sheets_are_not_selected_by_order():
    book = openpyxl.load_workbook(BytesIO(progresso_xlsx_cells({})))
    book.copy_worksheet(book.active).title = "Progresso de safra revisado"
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    with pytest.raises(ParseError, match="ambíguas"):
        progress_parser.parse_progresso_xlsx(buffer.getvalue())


@pytest.mark.parametrize("cells", [{"E11": None}, {"C11": None, "D11": None, "E11": None}])
def test_progress_missing_week_cannot_silently_drop_first_crop(cells):
    with pytest.raises(ParseError, match="semanais incompletas|Bloco sem operação"):
        progress_parser.parse_progresso_xlsx(progresso_xlsx_cells(cells))
