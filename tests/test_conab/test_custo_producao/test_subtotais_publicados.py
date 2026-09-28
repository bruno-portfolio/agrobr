from __future__ import annotations

import hashlib
import json
import re
import warnings
from pathlib import Path

import openpyxl
import pytest
import xlrd

from agrobr.conab.custo_producao import _context, _parse, _workbook, api, models
from agrobr.utils.warnings import warn_once_reset
from tests import helpers
from tests.helpers import sem_excecao

GOLDEN = Path(__file__).resolve().parents[2] / "golden_data"
R4 = GOLDEN / "reconciliacao_r4_20260918"
ARROZ = GOLDEN / "conab" / "custos_c17_20260925"
SERIES = {
    "cafe_arabica": ("seriehistoricacustoscafearabica2003a2025.xls", R4 / "feb4999ec69b7274.xls"),
    "cafe_conilon": ("seriehistoricacustoscafeconilon2007a2025.xls", R4 / "aceab784e7d4a6f2.xls"),
    "algodao": (
        "serie-historica-custos-algodao-em-pluma-1998-a-2026.xlsx",
        R4 / "a6cdcd1877333e68.xlsx",
    ),
    "soja": (
        "serie-historica-custos-soja-1997-a-2025.xls",
        GOLDEN / "conab" / "custos_20260908" / "soja.xls",
    ),
}
ROMANO = re.compile(r"^[IVX]+\s*[-–]")


def _celulas(caminho: Path, aba: str) -> list[tuple[int, str, float | None]]:
    if caminho.suffix == ".xlsx":
        livro = openpyxl.load_workbook(caminho, data_only=True)
        folha = livro[aba]
        linhas = [list(linha) for linha in folha.iter_rows(values_only=True)]
        mesclas = [
            (faixa.min_row - 1, faixa.max_row, faixa.min_col - 1, faixa.max_col)
            for faixa in folha.merged_cells.ranges
        ]
        livro.close()
    else:
        planilha = xlrd.open_workbook(caminho, on_demand=True, formatting_info=True)
        folha = planilha.sheet_by_name(aba)
        linhas = [folha.row_values(indice) for indice in range(folha.nrows)]
        mesclas = list(folha.merged_cells)
        planilha.release_resources()
    rotulo = valores = None
    for numero, linha in enumerate(linhas[:30]):
        for coluna, celula in enumerate(linha):
            texto = " ".join(str(celula or "").upper().split())
            if texto.startswith("DISCRIMINA"):
                rotulo = coluna
            if texto in {"(R$/HA)", "R$/HA", "CUSTO POR HA", "CUSTO/HA"}:
                valores = next(
                    (
                        range(inicio, fim)
                        for linha_inicio, linha_fim, inicio, fim in mesclas
                        if linha_inicio <= numero < linha_fim and inicio <= coluna < fim
                    ),
                    range(coluna, coluna + 1),
                )
    assert rotulo is not None and valores is not None
    saida = []
    for numero, linha in enumerate(linhas, start=1):
        texto = str(linha[rotulo] or "").strip() if rotulo < len(linha) else ""
        numeros = [
            float(linha[coluna])
            for coluna in valores
            if coluna < len(linha)
            and isinstance(linha[coluna], int | float)
            and not isinstance(linha[coluna], bool)
        ]
        assert len(numeros) <= 1, (aba, numero)
        if texto:
            saida.append((numero, texto, numeros[0] if numeros else None))
    return saida


def _linha(celulas: list[tuple[int, str, float | None]], inicio: str) -> tuple[int, str, float]:
    numero, rotulo, valor = next(
        celula for celula in celulas if celula[1].upper().startswith(inicio.upper())
    )
    assert valor is not None, rotulo
    return numero, rotulo, valor


async def _publicado(monkeypatch, cultura: str, aba: str):
    helpers.install_reconciliacao_r4_http(monkeypatch)
    warn_once_reset()
    with warnings.catch_warnings(record=True) as capturados:
        warnings.simplefilter("always")
        frame, meta = await api.custo_producao(
            cultura, planilha=SERIES[cultura][0], aba=aba, return_meta=True, use_cache=False
        )
    avisos = [
        str(aviso.message)
        for aviso in capturados
        if issubclass(aviso.category, UserWarning)
        and str(aviso.message).startswith("CONAB custos:")
    ]
    return frame, meta, avisos


@pytest.mark.parametrize(
    "cultura,aba,sai,ficam",
    [
        ("cafe_arabica", "L.Eduardo-BA-2005", ["2 - Operação com máquinas"], ["2.1 -", "2.2 -"]),
        (
            "cafe_arabica",
            "L.Eduardo-BA-2009",
            ["2 - Operação com máquinas próprias"],
            ["' 2.1 -", "' 2.2 -"],
        ),
        ("soja", "OGM-P. do Leste-MT-2012", [], ["8 - Sementes", "8.1 -"]),
    ],
)
async def test_grupo_valorado_sai_pela_leitura_que_fecha_com_o_subtotal(
    monkeypatch, cultura, aba, sai, ficam
):
    frame, meta, avisos = await _publicado(monkeypatch, cultura, aba)

    celulas = _celulas(SERIES[cultura][1], aba)
    _, _, custeio_publicado = _linha(celulas, "TOTAL DAS DESPESAS DE CUSTEIO")
    custeio = frame[(frame["tipo_linha"] == "item") & frame["secao"].str.startswith("I ")]
    assert custeio["valor_ha"].sum() == pytest.approx(custeio_publicado, abs=0.01)
    linhas = set(frame["linha"])
    notas = {nota["linha"]: nota for nota in meta.source_details["parser"]["notes"]}
    for inicio in sai:
        numero, rotulo, valor = _linha(celulas, inicio)
        assert rotulo == inicio and numero not in linhas
        assert notas[numero]["tipo"] == "group_header"
        assert notas[numero]["valores"]["valor_ha"] == pytest.approx(valor)
    for inicio in ficam:
        assert _linha(celulas, inicio)[0] in linhas
    assert avisos == [] and meta.validation_warnings == []


async def test_cabecalho_romano_com_zeros_abre_a_secao(monkeypatch):
    frame, meta, avisos = await _publicado(monkeypatch, "cafe_conilon", "Rio Bananal-ES-2019")

    celulas = _celulas(SERIES["cafe_conilon"][1], "Rio Bananal-ES-2019")
    cabecalho, rotulo, zeros = _linha(celulas, "V - OUTROS CUSTOS FIXOS")
    total, _, publicado = _linha(celulas, "TOTAL DE OUTROS CUSTOS FIXOS")
    assert zeros == 0
    secao = frame[(frame["linha"] > cabecalho) & (frame["linha"] < total)]
    assert len(secao) == total - cabecalho - 1
    assert secao["secao"].eq(rotulo).all()
    assert secao["categoria"].eq("custos_fixos").all()
    assert cabecalho not in set(frame["linha"])
    assert secao["valor_ha"].sum() == pytest.approx(publicado, abs=0.01)
    assert avisos == []


async def test_bloco_depois_do_total_vira_nota(monkeypatch):
    frame, meta, avisos = await _publicado(monkeypatch, "cafe_conilon", "R.Moura-RO-2014")

    celulas = _celulas(SERIES["cafe_conilon"][1], "R.Moura-RO-2014")
    operacional, _, _ = _linha(celulas, "CUSTO OPERACIONAL")
    bloco = [
        (numero, valor)
        for numero, rotulo, valor in celulas
        if numero > operacional and not rotulo.startswith("Elaboração")
    ]
    assert len(bloco) == 5 and bloco[0][1] == pytest.approx(sum(valor for _, valor in bloco[1:]))
    assert frame["linha"].max() == operacional
    notas = {nota["linha"]: nota for nota in meta.source_details["parser"]["notes"]}
    for numero, valor in bloco:
        assert notas[numero]["tipo"] == "memo_after_total"
        assert notas[numero]["valores"]["valor_ha"] == pytest.approx(valor)
    assert avisos == []


def _arroz(aba: str) -> models.ResultadoCusto:
    manifesto = json.loads((ARROZ / "manifest.json").read_text(encoding="utf-8"))["arquivos"][0]
    bruto = (ARROZ / manifesto["arquivo"]).read_bytes()
    assert hashlib.sha256(bruto).hexdigest() == manifesto["sha256"]
    recurso = models.RecursoCusto(
        cultura=manifesto["cultura"],
        planilha=manifesto["planilha"],
        titulo=manifesto["titulo"],
        pagina_url=manifesto["pagina_url"],
    )
    livro = _workbook.Workbook(bruto)
    try:
        contexto = _context.context(livro.read(aba, head=True), recurso, livro.names.index(aba))
        return _parse.parse_selected(livro.read(aba), contexto)
    finally:
        livro.close()


@pytest.mark.parametrize("aba", ["Massaranduba-SC-2007", "Meleiro-SC-2007"])
def test_agregado_de_gestao_no_custeio_sai_quando_os_componentes_fecham(aba):
    resultado = _arroz(aba)

    celulas = _celulas(ARROZ / "arroz_irrigado.xls", aba)
    gestao, _, agregado = _linha(celulas, "3 - Gestão da propriedade familiar")
    subtotal, _, publicado = _linha(celulas, "TOTAL DAS DESPESAS DE CUSTEIO")
    itens = [o for o in resultado.observacoes if o.tipo_linha == "item" and o.linha < subtotal]
    linhas = {item.linha for item in itens}
    assert gestao not in linhas
    assert sum(item.valor_ha or 0.0 for item in itens) == pytest.approx(publicado, abs=0.01)
    custeio = next(n for n, rotulo, _ in celulas if rotulo.startswith("I - DESPESAS DE CUSTEIO"))
    publicados = [n for n, _, valor in celulas if custeio < n < subtotal and valor is not None]
    assert linhas == set(publicados) - {gestao}
    nota = next(n for n in resultado.detalhes["notes"] if n.get("tipo") == "aggregate_row")
    assert nota["linha"] == gestao
    assert nota["valores"]["valor_ha"] == pytest.approx(agregado)
    assert all(check["fecha"] for check in resultado.detalhes["subtotal_checks"])


def _soma_publicada(celulas: list[tuple[int, str, float | None]], rotulo: str) -> float:
    numero = next(n for n, r, _ in celulas if r == rotulo)
    formula = re.search(r"\(([A-Z](?:\s*\+\s*[A-Z])+)\s*=\s*[A-Z]\)", rotulo)
    if formula:
        letras = [letra.strip() for letra in formula.group(1).split("+")]
        return sum(
            next(
                valor
                for _, r, valor in celulas
                if valor is not None
                and re.search(rf"(?:\({letra}\)|=\s*{letra}\))\s*$", r.replace(" ", ""))
            )
            for letra in letras
        )
    cabecalho = max(n for n, r, _ in celulas if n < numero and ROMANO.match(r))
    return sum(valor for n, _, valor in celulas if cabecalho < n < numero and valor is not None)


@pytest.mark.parametrize(
    "cultura,aba,rotulos",
    [
        ("cafe_arabica", "Patrocínio-MG-2022", ["TOTAL DE DEPRECIAÇÕES (E)"]),
        ("algodao", "Barreiras-BA-2011", ["TOTAL DAS OUTRAS DESPESAS (B)"]),
        ("soja", "OGM-P. do Leste-MT-2011", ["TOTAL DAS DESPESAS DE CUSTEIO DA LAVOURA (A)"]),
        ("soja", "OGM-Londrina-PR-2010", ["TOTAL DE RENDA DE FATORES (I)", "CUSTO TOTAL (H+I=J)"]),
    ],
)
async def test_subtotal_que_nao_fecha_avisa_com_os_numeros_publicados(
    monkeypatch, cultura, aba, rotulos
):
    frame, meta, avisos = await _publicado(monkeypatch, cultura, aba)

    celulas = _celulas(SERIES[cultura][1], aba)
    assert len(avisos) == len(rotulos)
    assert meta.validation_warnings == avisos
    publicadas = {n for n, _, valor in celulas if valor is not None}
    assert set(frame["linha"]) <= publicadas
    for rotulo in rotulos:
        numero, publicado_rotulo, publicado = _linha(celulas, rotulo)
        soma = _soma_publicada(celulas, publicado_rotulo)
        assert abs(publicado - soma) > 0.02
        (aviso,) = [aviso for aviso in avisos if f"{publicado_rotulo} publicado" in aviso]
        assert f"publicado {publicado:.2f}" in aviso
        assert f"{soma:.2f} (diferença de {publicado - soma:.2f})" in aviso
        assert "o agrobr repassa os números publicados" in aviso
        assert numero in set(frame["linha"])


CONCEITOS = {
    "mao_de_obra": ("mão-de-obra", "mão de obra"),
    "insumos": ("defensivos", "agrotóxicos", "mudas", "sementes"),
    "operacoes": ("tratores", "máquinas próprias"),
}


async def test_categoria_do_custeio_nao_muda_com_o_rotulo_do_ano(monkeypatch):
    vistos: dict[str, set[tuple[str, str, str]]] = {}
    for aba in ("L.Eduardo-BA-2005", "L.Eduardo-BA-2010", "L.Eduardo-BA-2020"):
        frame, _, _ = await _publicado(monkeypatch, "cafe_arabica", aba)
        custeio = frame[(frame["tipo_linha"] == "item") & frame["secao"].str.startswith("I ")]
        for item, categoria in zip(custeio["item"], custeio["categoria"], strict=True):
            rotulo = item.strip().lower()
            for esperada, trechos in CONCEITOS.items():
                if any(trecho in rotulo for trecho in trechos):
                    vistos.setdefault(esperada, set()).add((aba, item.strip(), categoria))

    assert set(vistos) == set(CONCEITOS)
    for esperada, linhas in vistos.items():
        assert {categoria for _, _, categoria in linhas} == {esperada}, linhas
        assert len({rotulo for _, rotulo, _ in linhas}) > 1, linhas


async def test_categoria_sai_da_secao_publicada(monkeypatch):
    frame, _, _ = await _publicado(monkeypatch, "cafe_arabica", "L.Eduardo-BA-2020")

    itens = frame[frame["tipo_linha"] == "item"]
    por_secao = {
        secao.split(" - ")[0]: set(grupo["categoria"]) for secao, grupo in itens.groupby("secao")
    }
    assert por_secao["IV"] == por_secao["V"] == {"custos_fixos"}
    assert por_secao["II"] == por_secao["III"] == por_secao["VI"] == {"outros"}
    totais = frame[frame["tipo_linha"] == "total"]
    assert not totais.empty and "custos_fixos" not in set(totais["categoria"])


def test_subtotal_sem_valor_por_hectare_fica_fora_da_conferencia():
    folha, contexto = helpers.load_r4_cost_sheet(
        "feb4999ec69b7274.xls", "cafe_arabica", "L.Eduardo-BA-2005"
    )
    celulas = _celulas(SERIES["cafe_arabica"][1], "L.Eduardo-BA-2005")
    numero, _, publicado = _linha(celulas, "TOTAL DAS DESPESAS DE CUSTEIO")
    linha = folha.linhas[numero - 1]
    coluna = next(c for c, valor in enumerate(linha) if valor == publicado)
    linha[coluna] = ""

    with sem_excecao():
        resultado = _parse.parse_selected(folha, contexto)

    subtotal = next(o for o in resultado.observacoes if o.linha == numero)
    assert subtotal.tipo_linha == "subtotal" and subtotal.valor_ha is None
    assert numero not in {check["linha"] for check in resultado.detalhes["subtotal_checks"]}
