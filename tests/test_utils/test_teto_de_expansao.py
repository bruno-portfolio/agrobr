from __future__ import annotations

import ast
import io
import struct
import tracemalloc
import zipfile
from collections.abc import Callable
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
import pytest

from agrobr import constants
from agrobr.acervo_fundiario import parser as acervo_parser
from agrobr.alt.anp_diesel import parser as anp_parser
from agrobr.antaq import client as antaq_client
from agrobr.b3 import parser as b3_parser
from agrobr.cepea import serie as cepea_serie
from agrobr.exceptions import ResourceLimitError, SourceUnavailableError
from agrobr.ibge import ftp_client
from agrobr.mapbiomas import client as mapbiomas_client
from agrobr.mapbiomas import municipal_parser
from agrobr.unica import parser as unica_parser
from agrobr.utils import io as io_utils
from tests.helpers import levanta_exatamente, sem_excecao

PACOTE = Path(__file__).resolve().parents[2] / "agrobr"
TETO = 1024**2
EXPANSAO = 8 * 1024**2
FORA_DO_HELPER = {
    ("utils/io.py", "open_zip_member"): "o próprio helper",
    ("inmet/client.py", "historico_membros"): "teto próprio (INMET_HISTORICO_MAX_*)",
    ("conab/_custo_producao/_merged.py", "xlsx_header_merges"): (
        "roda depois do _workbook, que confere a soma declarada e a expansão real do XLSX"
    ),
    ("acervo_fundiario/parser.py", "_shapefile"): (
        "lê só os 4 bytes do cabeçalho do .shp, depois do check_zip_expansion"
    ),
    ("defensivos/snapshot.py", "_read_bundle"): "cache local gravado pelo próprio agrobr",
    ("rnc/snapshot.py", "_read_bundle"): "cache local gravado pelo próprio agrobr",
}


def _eh_zipfile(no: ast.expr | None) -> bool:
    return isinstance(no, ast.Attribute | ast.Name) and ast.unparse(no) in {
        "zipfile.ZipFile",
        "ZipFile",
    }


class _Leituras(ast.NodeVisitor):
    """Cada leitura de membro de ZIP (``read``, ``open`` ou ``extract``), pela função mais interna."""

    def __init__(self, arquivo: str) -> None:
        self.arquivo = arquivo
        self.funcoes: list[tuple[str, set[str]]] = []
        self.achadas: set[tuple[str, str]] = set()

    def _funcao(self, no: ast.FunctionDef | ast.AsyncFunctionDef) -> None:
        nomes = {
            arg.arg for arg in no.args.args + no.args.kwonlyargs if _eh_zipfile(arg.annotation)
        }
        self.funcoes.append((no.name, nomes))
        self.generic_visit(no)
        self.funcoes.pop()

    visit_FunctionDef = _funcao
    visit_AsyncFunctionDef = _funcao

    def _nomes(self) -> set[str]:
        return self.funcoes[-1][1] if self.funcoes else set()

    def visit_withitem(self, no: ast.withitem) -> None:
        if (
            isinstance(no.context_expr, ast.Call)
            and _eh_zipfile(no.context_expr.func)
            and isinstance(no.optional_vars, ast.Name)
        ):
            self._nomes().add(no.optional_vars.id)
        self.generic_visit(no)

    def visit_Assign(self, no: ast.Assign) -> None:
        if isinstance(no.value, ast.Call) and _eh_zipfile(no.value.func):
            self._nomes().update(alvo.id for alvo in no.targets if isinstance(alvo, ast.Name))
        self.generic_visit(no)

    def visit_Call(self, no: ast.Call) -> None:
        escrita = any(
            isinstance(valor, ast.Constant) and valor.value == "w"
            for valor in [*no.args, *(kw.value for kw in no.keywords)]
        )
        if (
            isinstance(no.func, ast.Attribute)
            and no.func.attr in {"read", "open", "extract", "extractall"}
            and isinstance(no.func.value, ast.Name)
            and no.func.value.id in self._nomes()
            and not escrita
        ):
            self.achadas.add((self.arquivo, self.funcoes[-1][0] if self.funcoes else "<modulo>"))
        self.generic_visit(no)


def test_leitura_de_zip_passa_pelo_teto():
    achadas: set[tuple[str, str]] = set()
    for arquivo in PACOTE.rglob("*.py"):
        leituras = _Leituras(arquivo.relative_to(PACOTE).as_posix())
        leituras.visit(ast.parse(arquivo.read_text(encoding="utf-8")))
        achadas |= leituras.achadas

    assert sorted(achadas - set(FORA_DO_HELPER)) == []
    assert sorted(set(FORA_DO_HELPER) - achadas) == []


def _zip(membros: dict[str, bytes]) -> bytes:
    saida = io.BytesIO()
    with zipfile.ZipFile(saida, "w", zipfile.ZIP_DEFLATED) as arquivo:
        for nome, dados in membros.items():
            arquivo.writestr(nome, dados)
    return saida.getvalue()


def _com_membros(trocas: dict[str, bytes]) -> bytes:
    livro = openpyxl.Workbook()
    livro.active["A1"] = "x"
    saida = io.BytesIO()
    livro.save(saida)
    with zipfile.ZipFile(io.BytesIO(saida.getvalue())) as origem:
        membros = {info.filename: origem.read(info) for info in origem.infolist()}
    return _zip({**membros, **trocas})


def _planilha(linhas: bytes, *, antes: bytes = b"") -> bytes:
    return (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        + antes
        + b"<sheetData>"
        + linhas
        + b"</sheetData></worksheet>"
    )


def _xlsx(celula: bytes) -> bytes:
    return _com_membros(
        {
            "xl/worksheets/sheet1.xml": _planilha(
                b'<row r="1"><c r="A1" t="inlineStr"><is><t>' + celula + b"</t></is></c></row>"
            )
        }
    )


def _forjar(conteudo: bytes, membro: str, declarado: int) -> bytes:
    """Troca o tamanho expandido declarado do membro, no cabeçalho local e no diretório central."""
    bruto = bytearray(conteudo)
    nome = membro.encode()
    for assinatura, deslocamento, campo, cabecalho in (
        (b"PK\x03\x04", 26, 22, 30),
        (b"PK\x01\x02", 28, 24, 46),
    ):
        posicao = bruto.find(assinatura)
        while posicao != -1:
            tamanho_nome = struct.unpack_from("<H", bruto, posicao + deslocamento)[0]
            if bytes(bruto[posicao + cabecalho : posicao + cabecalho + tamanho_nome]) == nome:
                struct.pack_into("<I", bruto, posicao + campo, declarado)
            posicao = bruto.find(assinatura, posicao + 4)
    return bytes(bruto)


BOMBA_XLSX = _xlsx(b"x" * EXPANSAO)
LEITORES: dict[str, tuple[str, Callable[[], bytes], Callable[[bytes], Any]]] = {
    "queimadas": (
        "queimadas",
        lambda: _zip({"focos.csv": b"0" * EXPANSAO}),
        lambda bomba: io_utils.extract_csv_from_zip(
            bomba, source="queimadas", url="https://exemplo"
        ),
    ),
    "b3_zip_interno": (
        "b3",
        lambda: _zip({"PR260925.zip": b"0" * EXPANSAO}),
        b3_parser.parse_ajustes_zip,
    ),
    "b3_xml": (
        "b3",
        lambda: _zip({"PR260925.zip": _zip({"BVBG.086.01.xml": b"<a>" + b" " * EXPANSAO})}),
        b3_parser.parse_ajustes_zip,
    ),
    "mapbiomas_zip": (
        "mapbiomas",
        lambda: _zip({constants.MAPBIOMAS_MUNICIPAL_MEMBER_11: b"0" * EXPANSAO}),
        lambda bomba: mapbiomas_client._municipal_member(bomba, "https://exemplo"),
    ),
    "mapbiomas_xlsx": (
        "mapbiomas",
        lambda: BOMBA_XLSX,
        lambda bomba: municipal_parser.parse_cobertura_municipal(bomba, colecao=11),
    ),
    "antaq": (
        "antaq",
        lambda: _zip({"2024Carga.txt": b"0" * EXPANSAO}),
        lambda bomba: antaq_client._extract_txt_from_zip(bomba, "2024Carga.txt"),
    ),
    "ibge_ftp": (
        "ibge",
        lambda: _zip({"Tab_3Mn.xls": b"0" * EXPANSAO}),
        ftp_client.extract_tables_from_zip,
    ),
    "anp_xlsx_calamine": ("anp_diesel", lambda: BOMBA_XLSX, anp_parser._read_precos_xlsx),
    "abiove_xlsx": (
        "abiove",
        lambda: BOMBA_XLSX,
        lambda bomba: io_utils.open_excel_safe(bomba, source="abiove"),
    ),
    "cepea_serie_xlsx": ("cepea", lambda: BOMBA_XLSX, cepea_serie.ler_planilha),
    "unica_xlsx": (
        "unica",
        lambda: BOMBA_XLSX,
        lambda bomba: unica_parser.parse_historico_xlsx(bomba, "cana"),
    ),
}


@pytest.mark.parametrize("leitor", list(LEITORES))
def test_bomba_recusada_antes_de_expandir(monkeypatch, leitor):
    fonte, montar, ler = LEITORES[leitor]
    bomba = montar()
    monkeypatch.setitem(constants.MAX_EXPANDED_BYTES, fonte, TETO)

    tracemalloc.start()
    try:
        with levanta_exatamente(ResourceLimitError, f"acima do teto de {TETO}"):
            ler(bomba)
        _atual, pico = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()

    assert len(bomba) < 64 * 1024
    assert pico < EXPANSAO // 2


def test_xlsx_com_tamanho_declarado_forjado_e_recusado(monkeypatch):
    monkeypatch.setitem(constants.MAX_EXPANDED_BYTES, "anp_diesel", TETO)
    forjado = _forjar(BOMBA_XLSX, "xl/worksheets/sheet1.xml", 1024)

    tracemalloc.start()
    try:
        with levanta_exatamente(ResourceLimitError, f"expande de fato para mais de {TETO} bytes"):
            anp_parser._read_precos_xlsx(forjado)
        _atual, pico = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()

    assert pico < EXPANSAO // 2


def test_censo_legado_soma_dos_membros_tem_teto(monkeypatch):
    monkeypatch.setitem(constants.MAX_EXPANDED_BYTES, "ibge", TETO)
    bomba = _zip({f"Tab_{indice}Mn.xls": b"0" * (TETO // 2) for indice in range(4)})

    tracemalloc.start()
    try:
        with levanta_exatamente(
            ResourceLimitError,
            f"os membros do ZIP expandem para {2 * TETO} bytes, acima do teto de {TETO}",
        ):
            ftp_client.extract_tables_from_zip(bomba)
        _atual, pico = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()

    assert pico < TETO // 2


def test_zip_com_tamanho_declarado_forjado_para_no_crc():
    forjado = _forjar(_zip({"focos.csv": b"0" * EXPANSAO}), "focos.csv", 1024)

    with levanta_exatamente(SourceUnavailableError, "Resposta não é um ZIP válido: Bad CRC-32"):
        io_utils.extract_csv_from_zip(forjado, source="queimadas", url="https://exemplo")


def test_xlsx_dentro_do_teto_segue_para_o_leitor():
    pequeno = _xlsx(b"x" * 10)

    with sem_excecao():
        io_utils.check_xlsx_expansion(pequeno, source="anp_diesel")
        io_utils.check_xlsx_expansion(b"\xd0\xcf\x11\xe0 xls antigo", source="cepea")
        io_utils.check_xlsx_expansion(b"PK\x03\x04 corrompido", source="cepea")


def test_fonte_sem_teto_proprio_usa_o_padrao():
    assert "deral" not in constants.MAX_EXPANDED_BYTES
    assert io_utils._expansion_limit("deral") == constants.MAX_EXPANDED_BYTES_DEFAULT


ESPARSA = _planilha(
    b'<row r="1"><c r="A1"/><c r="XFD1"/></row><row r="1048576"><c r="A1048576"/></row>'
)
RETANGULO_ESPARSO = 1_048_576 * 16_384
FORMAS_ESPARSAS = {
    "referencias": (ESPARSA, RETANGULO_ESPARSO),
    "dimensao_declarada": (
        _planilha(b'<row r="1"><c r="A1"/></row>', antes=b'<dimension ref="A1:XFD1048576"/>'),
        RETANGULO_ESPARSO,
    ),
    "celulas_sem_r": (
        _planilha(
            b'<row r="1">' + b"<c/>" * 20 + b'</row><row r="1048576"><c r="A1048576"/></row>'
        ),
        1_048_576 * 20,
    ),
    "prefixo": (
        b'<x:worksheet xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b'<x:sheetData><x:row r="1"><x:c r="A1"/><x:c r="XFD1"/></x:row>'
        b'<x:row r="1048576"><x:c r="A1048576"/></x:row></x:sheetData></x:worksheet>',
        RETANGULO_ESPARSO,
    ),
    "linha_em_ponto_flutuante": (
        _planilha(b'<row r="1"><c r="XFD1"/></row><row r="1048576.0"/>'),
        RETANGULO_ESPARSO,
    ),
    "referencia_com_entidade": (
        _planilha(b'<row r="1"><c r="&#88;FD1"/><c r="A1048576"/></row>'),
        RETANGULO_ESPARSO,
    ),
    "utf16": (
        ESPARSA.replace(b'encoding="UTF-8"', b'encoding="UTF-16"').decode().encode("utf-16"),
        RETANGULO_ESPARSO,
    ),
}


@pytest.mark.parametrize("forma", list(FORMAS_ESPARSAS))
def test_xlsx_esparso_recusado_pelo_retangulo(forma):
    planilha, celulas = FORMAS_ESPARSAS[forma]
    xlsx = _com_membros({"xl/worksheets/sheet1.xml": planilha})

    with levanta_exatamente(
        ResourceLimitError,
        f"sheet1.xml do XLSX monta {celulas} células, acima do teto de {constants.MAX_XLSX_CELLS}",
    ):
        io_utils.check_xlsx_expansion(xlsx, source="conab")

    assert len(xlsx) < 16 * 1024


ILEGIVEIS = {
    "xml_mal_formado": (
        _planilha(b'<row><c/></x><row r="1048576"><c r="XFD1048576"/></row>'),
        "não é XML bem formado",
    ),
    "doctype": (
        b'<!DOCTYPE worksheet [<!ENTITY e "1">]>' + _planilha(b'<row r="1"><c r="A1"/></row>'),
        "DOCTYPE",
    ),
}


@pytest.mark.parametrize("forma", list(ILEGIVEIS))
def test_planilha_ilegivel_recusada(forma):
    planilha, motivo = ILEGIVEIS[forma]
    xlsx = _com_membros({"xl/worksheets/sheet1.xml": planilha})

    with levanta_exatamente(ResourceLimitError, f"sheet1.xml do XLSX está ilegível [(]{motivo}"):
        io_utils.check_xlsx_expansion(xlsx, source="conab")


def test_membro_sem_planilha_mal_formado_segue_para_o_leitor():
    xlsx = _com_membros(
        {
            "xl/media/image1.png": b"\x89PNG\r\n\x1a\n" + bytes(range(256)) * 4,
            "xl/drawings/vmlDrawing1.vml": b"<xml><x:ClientData><x:Row>9</x:Row><br></xml>",
        }
    )

    with sem_excecao():
        io_utils.check_xlsx_expansion(xlsx, source="conab")


def test_contagem_rapida_e_exata_concordam():
    canonica = _planilha(
        b'<row r="2"><c r="C2"/></row><row r="5" spans="2:2"><c r="B5" t="s"><v>0</v></c></row>',
        antes=b'<dimension ref="B2:C5"/>',
    )
    fora_do_padrao = canonica.replace(b'<row r="5" spans="2:2">', b'<row spans="2:2" r="5">')
    contagens = []
    for planilha in (canonica, fora_do_padrao):
        bruto = _zip({"xl/worksheets/sheet1.xml": planilha})
        with zipfile.ZipFile(io.BytesIO(bruto)) as arquivo:
            info = arquivo.getinfo("xl/worksheets/sheet1.xml")
        contagens.append(
            (
                io_utils._fast_cells(io.BytesIO(bruto), info, constants.MAX_XLSX_CELLS),
                io_utils._exact_cells(io.BytesIO(bruto), info, constants.MAX_XLSX_CELLS),
            )
        )

    assert contagens == [(15, 15), (None, 15)]


@pytest.mark.parametrize(
    "leitor",
    ["mapbiomas_xlsx", "anp_xlsx_calamine", "abiove_xlsx", "cepea_serie_xlsx", "unica_xlsx"],
)
def test_xlsx_esparso_recusado_antes_do_leitor(monkeypatch, leitor):
    def proibido(*_args: Any, **_kwargs: Any) -> Any:
        raise AssertionError("o leitor de planilha abriu o XLSX esparso")

    esparso = _com_membros({"xl/worksheets/sheet1.xml": ESPARSA})
    for modulo, nome in ((pd, "read_excel"), (pd, "ExcelFile"), (openpyxl, "load_workbook")):
        monkeypatch.setattr(modulo, nome, proibido)
    _fonte, _montar, ler = LEITORES[leitor]

    with levanta_exatamente(ResourceLimitError, f"monta {RETANGULO_ESPARSO} células"):
        ler(esparso)


NAO_XLSX = {
    "sem_workbook_xml": (
        _zip({"xl/worksheets/sheet1.xml": _planilha(b'<row r="1"><c r="A1"/></row>')}),
        "sem xl/workbook.xml",
    ),
    "content_xml": (
        _com_membros({"content.xml": b"<office:document-content/>"}),
        "com content.xml",
    ),
    "workbook_bin": (_com_membros({"XL\\Workbook.bin": b"\x00"}), "com xl/workbook.bin"),
}


@pytest.mark.parametrize("forma", list(NAO_XLSX))
def test_zip_que_nao_e_xlsx_recusado_antes_do_leitor(monkeypatch, forma):
    def proibido(*_args: Any, **_kwargs: Any) -> Any:
        raise AssertionError("o leitor de planilha abriu o ZIP que não é XLSX")

    arquivo, motivo = NAO_XLSX[forma]
    for modulo, nome in ((pd, "read_excel"), (pd, "ExcelFile"), (openpyxl, "load_workbook")):
        monkeypatch.setattr(modulo, nome, proibido)

    for leitor in (
        "mapbiomas_xlsx",
        "anp_xlsx_calamine",
        "abiove_xlsx",
        "cepea_serie_xlsx",
        "unica_xlsx",
    ):
        _fonte, _montar, ler = LEITORES[leitor]
        with levanta_exatamente(
            ResourceLimitError,
            f"o arquivo não é XLSX [(]{motivo}[)]; o tamanho da planilha não é conhecido",
        ):
            ler(arquivo)


def test_xls_com_pacote_embutido_segue_para_o_leitor():
    golden = Path(__file__).resolve().parents[1] / "golden_data"
    xls = (golden / "conab" / "serie_historica_20260917" / "cafe.xls").read_bytes()
    assert zipfile.is_zipfile(io.BytesIO(xls))

    with sem_excecao():
        io_utils.check_xlsx_expansion(xls, source="conab_serie_historica")
    with levanta_exatamente(ResourceLimitError, "o arquivo não é XLSX [(]com content.xml[)]"):
        io_utils.check_xlsx_expansion(
            xls[:8] + _com_membros({"content.xml": b"<office:document-content/>"}),
            source="conab_serie_historica",
        )


def _shapefile_zip(pasta: Path, *, linhas: int = 4000) -> Path:
    gpd = pytest.importorskip("geopandas")
    geometria = pytest.importorskip("shapely.geometry")
    camada = gpd.GeoDataFrame(
        {"txt": ["x" * 250] * linhas},
        geometry=[geometria.Point(-50, -15)] * linhas,
        crs="EPSG:4674",
    )
    camada.to_file(pasta / "a.shp", engine="pyogrio")
    destino = pasta / "a.zip"
    partes = sorted(pasta.glob("a.*"))
    with zipfile.ZipFile(destino, "w", zipfile.ZIP_DEFLATED) as arquivo:
        for parte in partes:
            arquivo.write(parte, parte.name)
    return destino


LEITORES_ACERVO = {
    "tabular": acervo_parser.parse_assentamentos,
    "geo": acervo_parser.parse_assentamentos_geo,
}


@pytest.mark.parametrize("leitor", list(LEITORES_ACERVO))
def test_acervo_zip_com_tamanho_forjado_recusado_antes_do_gdal(monkeypatch, tmp_path, leitor):
    honesto = _shapefile_zip(tmp_path)
    with zipfile.ZipFile(honesto) as arquivo:
        declarado = sum(info.file_size for info in arquivo.infolist())
        declarado += 1024 - arquivo.getinfo("a.dbf").file_size
    forjado = tmp_path / "forjado.zip"
    forjado.write_bytes(_forjar(honesto.read_bytes(), "a.dbf", 1024))
    teto = declarado + 4096
    monkeypatch.setattr(constants, "ACERVO_MAX_DOWNLOAD_BYTES", teto)

    with levanta_exatamente(
        ResourceLimitError, f"o shapefile expande de fato para mais de {teto} bytes"
    ):
        LEITORES_ACERVO[leitor](forjado)


def test_acervo_zip_acima_do_teto_declarado_recusado(monkeypatch, tmp_path):
    honesto = _shapefile_zip(tmp_path)
    with zipfile.ZipFile(honesto) as arquivo:
        declarado = sum(info.file_size for info in arquivo.infolist())
    monkeypatch.setattr(constants, "ACERVO_MAX_DOWNLOAD_BYTES", declarado - 1)

    with levanta_exatamente(
        ResourceLimitError,
        f"o shapefile expande para {declarado} bytes, acima do teto de {declarado - 1}",
    ):
        acervo_parser.parse_snci(honesto)


def test_acervo_zip_dentro_do_teto_segue_para_o_gdal(tmp_path):
    honesto = _shapefile_zip(tmp_path)

    with sem_excecao():
        tabular = acervo_parser._read_tabular(honesto)
        recorte = acervo_parser._read_tabular(honesto, bbox=(-51.0, -16.0, -49.0, -14.0))

    assert len(tabular) == len(recorte) == 4000
